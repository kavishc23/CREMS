from __future__ import annotations

import json
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path("/Users/kavishchandra/Documents/CS400")
DOWNLOADS = Path("/Users/kavishchandra/Downloads")
OUTPUT = ROOT / "tmp/client_documents/prepared"
OUTPUT.mkdir(parents=True, exist_ok=True)


def text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value).strip()


def branch_code(location: str) -> str:
    value = location.upper()
    if "NAKASI" in value:
        return "NAK"
    if "NADI" in value:
        return "NAD"
    if "LAUTOKA" in value:
        return "LAU"
    if "LABASA" in value:
        return "LAB"
    return "SUV"


def vehicle_category(description: str) -> str:
    value = description.upper()
    if any(token in value for token in ("TRUCK", "C/CHS", "NPR", "FSR", "FTR", "NQR")):
        return "Commercial Truck"
    if any(token in value for token in ("PICK UP", "PICKUP", "NAVARA", "D-MAX", "BT50", "DOUBLE CAB")):
        return "Pickup"
    if any(token in value for token in ("HIACE", "URVAN", "SERENA", "NV150", "STARIA")):
        return "Van / People Mover"
    if any(token in value for token in ("PALISADE", "SANTA FE", "TUCSON", "X-TRAIL", "CR-V", "ZR-V", "X55")):
        return "SUV"
    return "Passenger Vehicle"


def equipment_category(description: str, sort_code: str) -> str:
    value = f"{description} {sort_code}".upper()
    rules = [
        ("FORK", "Forklift"), ("GEN", "Generator"), ("PALLET", "Pallet Truck"),
        ("EXCAV", "Excavator"), ("GRAD", "Motor Grader"), ("COMPACTOR", "Compactor"),
        ("LOADER", "Loader"), ("TELEHAND", "Telehandler"), ("TRAILER", "Trailer"),
        ("TRACTOR", "Track-Type Tractor"), ("TOW MOTOR", "Tow Motor"),
        ("WHEEL DOZER", "Wheel Dozer"), ("SKIDDER", "Wheel Skidder"),
        ("VEHICLE", "Support Vehicle"),
    ]
    for token, category in rules:
        if token in value:
            return category
    return "Equipment"


def model_year(value) -> str:
    if isinstance(value, (date, datetime)):
        return str(value.year)
    match = re.search(r"(19|20)\d{2}", text(value))
    return match.group(0) if match else ""


def read_rows(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    headers = [text(value) for value in rows[0]]
    return headers, [dict(zip(headers, row)) for row in rows[1:] if any(value is not None for value in row)]


vehicle_headers, vehicle_rows = read_rows(DOWNLOADS / "Rental Vehicle Stock List (1).xlsx")
equipment_headers, equipment_rows = read_rows(DOWNLOADS / "Rental Stock List - CAT (1).xlsx")

vehicles = []
branch_sequences = {}
for source_index, row in enumerate(vehicle_rows, start=2):
    location = text(row.get("Location_Description"))
    code = branch_code(location)
    branch_sequences[code] = branch_sequences.get(code, 0) + 1
    sequence = branch_sequences[code]
    description = text(row.get("DESCRIPTION"))
    make = description.split("-")[0].strip().title() if description else ""
    model = description.split("-", 1)[1].strip() if "-" in description else description
    odometer = row.get("Odometer")
    quality = []
    if not text(row.get("Rego_No")): quality.append("Missing registration number")
    if not text(row.get("COLOUR")): quality.append("Missing colour")
    if not text(row.get("DATE_FIRST_REGISTERED")): quality.append("Missing first-registration date")
    if text(row.get("Rego_No")).upper() == "TUCSON": quality.append("Registration appears to contain a model name")
    if isinstance(odometer, (int, float)) and odometer > 500000: quality.append("Odometer is unusually high; verify units/value")
    if isinstance(odometer, (int, float)) and odometer <= 25: quality.append("Very low odometer; confirm new vehicle or placeholder")
    if text(row.get("VIN")).endswith("."): quality.append("VIN has trailing punctuation")
    vehicles.append({
        "source_row": source_index,
        "asset_number": f"MOT-{code}-{sequence:04d}",
        "division_code": "MOTORS",
        "service_code": "VEHICLE_RENTAL",
        "branch_code": code,
        "source_location": location,
        "asset_type": "Vehicle",
        "category": vehicle_category(description),
        "name": description.title(),
        "manufacturer": make,
        "model": model,
        "model_year": model_year(row.get("Manufacture_Year")),
        "registration_number": f"REG-{code}-{sequence:04d}",
        "vin_or_chassis_number": f"VIN-ANON-{source_index - 1:05d}",
        "serial_number": "",
        "colour": text(row.get("COLOUR")).title(),
        "registration_expiry": text(row.get("Registration_Expiry_Date")),
        "acquisition_date": text(row.get("Stocked_Date")),
        "first_registered_date": text(row.get("DATE_FIRST_REGISTERED")),
        "meter_unit": "km",
        "current_meter_reading": odometer if isinstance(odometer, (int, float)) else "",
        "status": "Available",
        "ownership_type": "",
        "daily_rate": "",
        "source_identifier_policy": "Registration, VIN and stock number replaced with non-production identifiers",
        "quality_flags": "; ".join(quality),
    })

equipment = []
for source_index, row in enumerate(equipment_rows, start=2):
    location = text(row.get("Branch"))
    code = branch_code(location)
    description = text(row.get("Desc."))
    make_code = text(row.get("Make"))
    sequence = source_index - 1
    equipment.append({
        "source_row": source_index,
        "asset_number": f"CTR-{code}-{sequence:04d}",
        "division_code": "CARPTRAC",
        "service_code": "EQUIPMENT_HIRE",
        "branch_code": code,
        "source_location": location,
        "asset_type": "Equipment",
        "category": equipment_category(description, text(row.get("Sort"))),
        "name": description.title(),
        "manufacturer": "Caterpillar" if make_code.upper() == "CAT" else make_code,
        "model": text(row.get("Type")),
        "model_year": "",
        "registration_number": "",
        "vin_or_chassis_number": "",
        "serial_number": "",
        "colour": "",
        "registration_expiry": "",
        "acquisition_date": "",
        "first_registered_date": "",
        "meter_unit": "hours" if equipment_category(description, text(row.get("Sort"))) not in ("Trailer", "Pallet Truck") else "units",
        "current_meter_reading": "",
        "status": "Available",
        "ownership_type": "",
        "daily_rate": "",
        "source_identifier_policy": "Internal equipment number replaced with a non-production asset number",
        "quality_flags": "Confirm make/sort code definitions; confirm whether Type is model, capacity or subtype",
        "source_make_code": make_code,
        "source_sort_code": text(row.get("Sort")),
    })

payload = {
    "metadata": {
        "classification": "ANONYMISED CLIENT VALIDATION DATA - NOT FOR PRODUCTION IMPORT",
        "vehicle_source_columns": vehicle_headers,
        "equipment_source_columns": equipment_headers,
        "vehicle_count": len(vehicles),
        "equipment_count": len(equipment),
    },
    "assets": vehicles + equipment,
}

(OUTPUT / "client_asset_import_anonymised.json").write_text(json.dumps(payload, indent=2), encoding="utf-8")
print(json.dumps(payload["metadata"], indent=2))
