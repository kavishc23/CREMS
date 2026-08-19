from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path('/Users/kavishchandra/Documents/CS400')
OUT = ROOT / 'output/documents/CREMS_Functional_Specification_Document_v1.0.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)
SKILL = Path('/Users/kavishchandra/.codex/plugins/cache/openai-primary-runtime/documents/26.813.12317/skills/documents')
sys.path.insert(0, str(SKILL / 'scripts'))
from table_geometry import apply_table_geometry, column_widths_from_weights

BLACK='151515'; YELLOW='FFEA00'; GREEN='08783E'; PALE_GREEN='EAF4EE'; PALE_YELLOW='FFF9CC'
LIGHT_GRAY='F2F4F3'; MID_GRAY='686D69'; WHITE='FFFFFF'; RED='9B1C1C'; INK='202421'; BLUE='155E75'
CONTENT_WIDTH=9360

def font(run,size=10.2,bold=None,italic=None,color=INK,name='Aptos'):
    run.font.name=name; run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),name); run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),name)
    run.font.size=Pt(size); run.font.color.rgb=RGBColor.from_string(color)
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic

def shade(cell,fill):
    pr=cell._tc.get_or_add_tcPr(); sh=pr.find(qn('w:shd'))
    if sh is None: sh=OxmlElement('w:shd'); pr.append(sh)
    sh.set(qn('w:fill'),fill)

def repeat(row):
    pr=row._tr.get_or_add_trPr(); x=OxmlElement('w:tblHeader'); x.set(qn('w:val'),'true'); pr.append(x)

def no_split(row): row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))

def cell_text(cell,text,bold=False,color=INK,size=8.4,align=None):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.02
    if align is not None: p.alignment=align
    font(p.add_run(str(text)),size,bold,color=color); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc,headers,rows,weights,font_size=8.2,fill=GREEN):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.autofit=False
    for i,h in enumerate(headers): cell_text(t.rows[0].cells[i],h,True,WHITE,8.5,WD_ALIGN_PARAGRAPH.CENTER); shade(t.rows[0].cells[i],fill)
    repeat(t.rows[0])
    for ri,values in enumerate(rows):
        row=t.add_row(); no_split(row)
        for i,v in enumerate(values):
            cell_text(row.cells[i],v,size=font_size)
            if ri%2: shade(row.cells[i],LIGHT_GRAY)
    widths=column_widths_from_weights(weights,CONTENT_WIDTH)
    apply_table_geometry(t,widths,table_width_dxa=CONTENT_WIDTH,indent_dxa=120,cell_margins_dxa={'top':80,'bottom':80,'start':100,'end':100})
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def para(doc,text='',bold=False,italic=False,color=INK,size=10.2,after=5,before=0,align=None,keep=False):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.08; p.paragraph_format.keep_with_next=keep
    if align is not None: p.alignment=align
    font(p.add_run(text),size,bold,italic,color); return p

def bullet(doc,text,level=0):
    p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.06; font(p.add_run(text),9.8); return p

def callout(doc,label,text,fill=PALE_YELLOW):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(7); p.paragraph_format.left_indent=Inches(.08); p.paragraph_format.right_indent=Inches(.08)
    pr=p._p.get_or_add_pPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),fill); pr.append(sh)
    r=p.add_run(label+': '); font(r,9.7,True,color=BLACK); font(p.add_run(text),9.7,color=BLACK)

def heading(doc,text,level=1): doc.add_heading(text,level=level)
def page(doc): doc.add_page_break()

def add_page_field(p):
    p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(p.add_run('Page '),8,color=MID_GRAY); r=p.add_run()
    b=OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'),'begin'); i=OxmlElement('w:instrText'); i.set(qn('xml:space'),'preserve'); i.text=' PAGE '
    s=OxmlElement('w:fldChar'); s.set(qn('w:fldCharType'),'separate'); tx=OxmlElement('w:t'); tx.text='1'; s.append(tx); e=OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'),'end'); r._r.extend([b,i,s,e])

def module_header(doc,number,title,objective,actors,permissions):
    heading(doc,f'{number} {title}',1); para(doc,objective,after=7)
    table(doc,['Specification item','Definition'],[
        ('Primary actors',actors),('Required permissions / scope',permissions),('Entry condition','Authenticated actor has an active account and access to the selected division/branch unless the function is public.'),('Common exit condition','Data is validated, saved atomically, audit history is recorded and the next permitted action is shown.')
    ],[1.7,5.3],font_size=8.5)

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.72); sec.left_margin=Inches(.9); sec.right_margin=Inches(.9); sec.header_distance=Inches(.35); sec.footer_distance=Inches(.35)
styles=doc.styles; n=styles['Normal']; n.font.name='Aptos'; n._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); n._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos'); n.font.size=Pt(10.2); n.font.color.rgb=RGBColor.from_string(INK); n.paragraph_format.space_after=Pt(5); n.paragraph_format.line_spacing=1.08
for name,size,color,before,after in [('Heading 1',15,GREEN,13,7),('Heading 2',12.2,GREEN,10,5),('Heading 3',10.8,BLACK,7,3)]:
    st=styles[name]; st.font.name='Aptos Display'; st._element.rPr.rFonts.set(qn('w:ascii'),'Aptos Display'); st._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos Display'); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color); st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True

hp=sec.header.paragraphs[0]; font(hp.add_run('CREMS | Functional Specification Document'),8.2,True,color=MID_GRAY)
fp=sec.footer.paragraphs[0]; font(fp.add_run('Carpenters Fiji Pte Limited | Version 1.0 | For client validation'),8,color=MID_GRAY); add_page_field(sec.footer.add_paragraph())

# Cover — customer_pack template using compact_reference_guide density.
para(doc,'CARPENTERS FIJI PTE LIMITED',True,color=GREEN,size=11,after=35)
para(doc,'FUNCTIONAL SPECIFICATION\nDOCUMENT',True,color=BLACK,size=27,after=8)
para(doc,'Car Rental and Equipment Management System (CREMS)',True,color=GREEN,size=15.5,after=8)
para(doc,'Motors • Carptrac • Limited Shipping Hire Services',size=12.5,color=MID_GRAY,after=24)
callout(doc,'Purpose','Defines exactly how CREMS screens, workflows, validations, permissions, calculations, documents and integrations shall operate for implementation and client acceptance.')
table(doc,['Document metadata','Value'],[
    ('Document ID','CREMS-FSD-001'),('Version / status','1.0 — Client validation draft'),('Prepared for','Carpenters Fiji Pte Limited'),('Prepared by','CS400 IEP Project Team'),('Prepared date','19 August 2026'),('Initial functional scope','Carpenters Motors; Carptrac; Carpenters Shipping limited to portable toilets, large bins and scaffolding hire'),('Related baseline','CREMS Software Requirements Specification v1.0 and validated client meeting decisions')
],[1.8,5.2],font_size=8.7,fill=BLACK)
para(doc,'CONFIDENTIAL • FOR PROJECT REVIEW AND SIGN-OFF',True,color=GREEN,size=9.2,after=0,align=WD_ALIGN_PARAGRAPH.CENTER)

page(doc); heading(doc,'Document control',1)
table(doc,['Version','Date','Author / owner','Change','Approval'],[
    ('0.1','19 Aug 2026','CS400 IEP Project Team','Initial FSD derived from requirements and implemented CREMS baseline','Internal review'),('1.0','19 Aug 2026','CS400 IEP Project Team','Issued for Carpenters functional validation','Pending client sign-off')
],[.8,1.0,1.6,2.8,1.0])
heading(doc,'Review and sign-off',2)
table(doc,['Reviewer','Role','Decision','Date','Signature / comments'],[
    ('','Carpenters business representative','Approve / Approve with changes / Reject','',''),('','Carpenters IT / security representative','Approve / Approve with changes / Reject','',''),('','Project supervisor','Reviewed','','')
],[1.3,1.8,1.7,.8,1.7])
callout(doc,'Sign-off meaning','Approval confirms that the described functional behaviour is an agreed delivery baseline. Later changes are controlled through the change-request process; sign-off does not waive security, data-quality or acceptance-test obligations.',PALE_GREEN)
heading(doc,'Distribution',2); para(doc,'Controlled copies: Carpenters project stakeholders, CS400 project team and academic supervisor. Credentials, production secrets and personal customer data are excluded from this document.')

page(doc); heading(doc,'Contents',1)
contents=['1. Introduction and boundaries','2. System context and functional architecture','3. Roles, permissions and operating scope','4. Common user-interface behaviour','5. Authentication, sessions and access administration','6. Organization, divisions, branches and services','7. Customer and customer-portal functions','8. Asset register, QR, lifecycle and availability','9. Booking requests, quotations and approvals','10. Pickup, agreements, active hire and returns','11. Inspections and maintenance','12. Personnel, transport and field operations','13. Finance, profitability, dashboards and reports','14. Notifications, documents and scheduled processing','15. API, validation, errors and audit behaviour','16. Functional acceptance and traceability','17. Assumptions, open decisions and change control','18. Client sign-off','Appendix A. Status models','Appendix B. Output document catalogue','Appendix C. Functional requirement index']
for x in contents: para(doc,x,True,color=GREEN,size=9.2,after=2)

page(doc); heading(doc,'1. Introduction and boundaries',1)
heading(doc,'1.1 Purpose',2); para(doc,'This FSD translates the approved CREMS requirements into observable system behaviour. It is intended for client validation, UI and API implementation, database design, test-case preparation, user training and acceptance testing. “Shall” statements are mandatory for the signed baseline.')
heading(doc,'1.2 In-scope divisions and services',2)
table(doc,['Division / area','Assets and services','Division-specific behaviour'],[
    ('Carpenters Motors','Passenger and commercial rental vehicles','Registration, odometer, fuel, driver/licence checks, vehicle inspections and daily/weekly rates.'),
    ('Carptrac','Forklifts, cranes, excavators, generators and related hire equipment','Engine hours, capacity/specification fields, operator qualification and assignment, safety checks, hourly/daily rates and delivery where required.'),
    ('Carpenters Shipping — limited hire scope','Portable toilets (portaloos), large metal bins and scaffolding','Quantity or unit-based hire, delivery/collection, site instructions, condition/cleaning checks and transport/labour/setup charges. Freight, cargo, vessel and port operations are excluded.')
],[1.45,2.65,3.05],font_size=8.4)
heading(doc,'1.3 Functional boundaries',2)
for x in ['Public catalogue and availability search; authenticated customer booking and tracking.','Staff role- and permission-based portals scoped by division and branch.','Customer, asset, quotation, booking, approval, agreement, inspection, return, maintenance, operator, finance and reporting workflows.','Flexible service catalogue, category attributes, rate and charge definitions so future divisions can be configured with minimal code change.','QR-based asset identification, email/PDF outputs, scheduled alerts and immutable audit events.']: bullet(doc,x)
heading(doc,'1.4 Out of scope for this baseline',2)
for x in ['Shipping freight forwarding, vessel, cargo, port and container-terminal operations.','Live online payment-gateway settlement, GPS route optimization and offline mobile synchronization unless approved through change control.','Replacement of Carpenters’ accounting, payroll or HR systems; CREMS will remain integration-ready.','Production deployment, data migration and third-party commercial fees unless separately authorized.']: bullet(doc,x)

page(doc); heading(doc,'2. System context and functional architecture',1)
table(doc,['Layer / actor','Responsibility','Information exchanged'],[
    ('Public and customer web','Browse services, register/sign in, request a booking, track status, view quotations, agreements and invoices.','HTTPS JSON requests; authenticated customer session; downloadable documents.'),
    ('Staff web portal','Role-specific work queues for bookings, assets, maintenance, finance, reporting and administration.','HTTPS JSON requests; secure staff session; QR input; document actions.'),
    ('ASP.NET Core API','Validates identity, scope and input; applies business rules; coordinates persistence, documents, email and audit.','REST endpoints under /api; consistent status codes and correlation identifiers.'),
    ('SQL Server','Stores operational records, configuration, security history and audit data.','Transactional reads/writes through Entity Framework Core.'),
    ('Email service','Delivers verification, reset, quotation, confirmation, agreement, invoice and alert messages.','Queued message, approved recipient, attachments/links and delivery result.'),
    ('PDF / QR services','Generate immutable customer documents and asset identification labels.','Versioned document snapshots and non-sensitive QR token/code.')
],[1.45,3.1,2.6])
heading(doc,'2.1 Functional design principles',2)
for x in ['Configuration before customization: divisions, branches, services, categories, attributes, charges and approvals are data-driven.','Least privilege: every protected action checks account state, role/permission and division/branch scope on the server.','Guided work: one obvious primary action, visible prerequisites and staff-friendly status labels.','Traceability: critical actions create lifecycle and audit records with actor, timestamp, scope and reference.','Performance: list screens use filtered queries, server pagination and selective fields rather than loading full histories.']: bullet(doc,x)

page(doc); heading(doc,'3. Roles, permissions and operating scope',1)
table(doc,['Role','Functional responsibility','Default operating scope'],[
    ('Super Administrator','Group configuration, divisions, permissions, security and full oversight.','All divisions and branches.'),('Administrator','Staff/customer access administration, approved configuration and audit.','Assigned or group-wide; cannot alter a Super Administrator unless explicitly permitted.'),('Branch Manager','Branch operations, approvals, exceptions and branch financial performance.','Assigned division(s) and branch(es).'),('Rental Officer','Customer records, requests, quotations, bookings, pickup and return.','Assigned division/service and branch scope.'),('Maintenance Officer','Faults, maintenance jobs, parts, meters and return-to-service checks.','Assigned maintenance-capable branches/divisions.'),('Finance Officer','Invoices, allocations, credit notes, statements and financial reports.','Assigned operating scope.'),('Driver / Operator','Assigned deliveries/hire work, QR scans, meters, checks and timesheets.','Only assigned jobs/assets and permitted field actions.'),('Customer','Own profile, bookings, quotations, agreements, invoices and requests.','Own account and authorized corporate account only.')
],[1.25,3.5,2.35],font_size=8.2)
heading(doc,'3.1 Permission rules',2)
table(doc,['Permission family','Examples','Rule'],[
    ('Identity','users.create, users.reset_password, users.manage_access, customers.manage_access','Server evaluates permission and target-user protections; self-elevation is blocked.'),('Organization','divisions.configure, branches.configure, services.configure, branch_calendar.manage','Division creation/deactivation is reserved for Super Administrators.'),('Assets','assets.view/create/edit/transfer/inspect/record_meter/retire/view_financials','Asset branch/division scope is checked on every read and mutation.'),('Operations','rentals.approve, maintenance.complete, pricing.configure','Approval separation and value/workflow rules apply.'),('Finance','payments.refund, reports.financial','Financial details remain hidden without explicit permission.')
],[1.25,3.25,2.6])
callout(doc,'Scope precedence','An assigned role or permission never grants access outside the user’s active division/branch scopes. Temporary access must have start and end dates and is ignored after expiry.',PALE_GREEN)

page(doc); heading(doc,'4. Common user-interface behaviour',1)
table(doc,['UI element','Required behaviour','Validation / accessibility'],[
    ('Navigation','Public site and customer portal are separate from /staff. Staff menu is grouped as Work, Fleet, Insights and Administration and collapses on desktop.','Only accessible pages are rendered; direct navigation is also rejected by the API.'),('Breadcrumbs','Display Home and current workspace; selecting an ancestor returns without losing saved data.','Current page has aria-current; browser back remains usable.'),('Work queues','Use status tabs with visible counts, server search, filters and pagination.','Empty, loading and error states explain the next action.'),('Record workspace','Use a full page or side drawer with summary, requirements, actions and collapsed history.','One primary action; destructive or irreversible actions require confirmation.'),('Forms','Labels remain visible; required fields and examples are explicit; dependent fields appear only when relevant.','Inline validation plus summary; focus moves to first invalid field.'),('Notifications','Success confirms record/reference and next action; warning explains recoverable issue; error includes safe retry guidance.','No credentials, stack traces or personal data in UI errors.'),('Responsive layout','Desktop tables adapt to mobile cards; key action remains visible.','Keyboard navigation, contrast, semantic headings and labelled controls are required.')
],[1.25,3.35,2.5])
heading(doc,'4.1 Standard list behaviour',2)
for x in ['Default page size 25; permitted sizes 25, 50 and 100.','Filters persist during the current workspace session and can be cleared in one action.','Sorting occurs on the server for indexed fields such as reference, status, date, branch and asset number.','Financial columns and actions are omitted when the actor lacks permission; values are not merely hidden with CSS.']: bullet(doc,x)

page(doc); module_header(doc,'5.','Authentication, sessions and access administration','Provide separate secure authentication experiences for customers and staff while enforcing account lifecycle, session timeout and least privilege.','Staff, customers, Administrators and Super Administrators','Public login endpoints; protected account and access actions use explicit permissions.')
heading(doc,'5.1 Staff authentication and session flow',2)
table(doc,['ID','User action / input','System processing and rules','Result'],[
    ('F-AUTH-01','Enter staff email and password.','Normalize email; rate-limit attempts; verify password hash and active staff role; never reveal whether an email exists.','Authenticated session or generic failure.'),('F-AUTH-02','Submit MFA code when policy requires.','Verify active challenge, expiry and attempt limit. Super Administrator MFA may be disabled only as a documented temporary environment setting.','Session upgraded or challenge rejected.'),('F-AUTH-03','Open another browser window.','Window-bound token must match the server session. A copied cookie without the matching window token is rejected.','New window must authenticate separately.'),('F-AUTH-04','Remain inactive.','Expire idle session at configured timeout; absolute lifetime also applies; warning is shown shortly before expiry.','Protected calls return 401 and UI returns to staff sign-in.'),('F-AUTH-05','Log out or revoke a session.','Invalidate server session and cookie; record logout/revocation event.','Session cannot be reused.')
],[.8,1.85,3.25,1.3],font_size=8.1)
heading(doc,'5.2 Account and access administration',2)
table(doc,['ID','Function','Key rules / validations','Audit outcome'],[
    ('F-IAM-01','Create/invite staff','Unique normalized email; role, division/branch scope and lifecycle state required. Temporary password must be changed.','Actor, target, role and scopes.'),('F-IAM-02','Assign scopes and permissions','Support one/multiple divisions/branches, temporary access and group-wide scope. Prevent self-elevation and invalid branch-division pairs.','Before/after permission and scope snapshot.'),('F-IAM-03','Reset password / unlock','Require permission; send reset workflow or issue temporary credential according to policy; revoke existing sessions.','Security event and notification delivery.'),('F-IAM-04','Lifecycle and departure','Invited → Activation pending → Active → Locked/Suspended/Deactivated. Departure revokes sessions, removes temporary access and approval assignments.','State transition and affected assignments.'),('F-IAM-05','Security history','Display login success/failure, IP, approximate device, MFA, reset, logout and expiry. Allow own sessions to be revoked.','Read access itself is logged where required.'),('F-IAM-06','Separation of duties','Block self-approval, self-role elevation, normal-admin changes to Super Administrators and creator approval of sensitive refunds.','Denied action with reason code.')
],[.8,1.6,3.75,1.15],font_size=8.0)

page(doc); module_header(doc,'6.','Organization, divisions, branches and services','Configure the group hierarchy and rental capabilities without requiring code changes for every future division.','Super Administrator; Administrator; authorized Branch Manager','divisions.configure, branches.configure, services.configure, branch_calendar.manage, pricing.configure')
table(doc,['ID','Function','Inputs and validation','Functional result'],[
    ('F-ORG-01','Maintain division profile','Unique code/name; branding, contacts, currency, tax, terms, capabilities, public/active flags and default workflow.','Division becomes selectable only when active.'),('F-ORG-02','Maintain service catalogue','Division, type, hire unit, personnel/delivery/quote rules, required documents, deposit, meter and inspection configuration.','Service can be enabled per branch.'),('F-ORG-03','Maintain branch profile','Code/name, addresses, contacts, coordinates, manager, instructions, delivery coverage and public/active flags.','Branch is available within assigned divisions.'),('F-ORG-04','Configure branch–division','Local contacts, manager, services, booking acceptance, maintenance capability, workflow, terms and price overrides.','Valid operating combination is created.'),('F-ORG-05','Operating calendar','Weekly hours, pickup/return cutoff, holidays, temporary closures and after-hours charge.','Availability and pricing respect calendar.'),('F-ORG-06','Future division enablement','Create division → services → categories/attributes → branches → rates/workflows → staff scopes.','New hire line can operate without changing core rental tables.')
],[.8,1.55,3.5,1.35],font_size=8.0)
heading(doc,'6.1 Initial service configuration',2)
table(doc,['Division','Example services','Default charging / special rules'],[
    ('Motors','Vehicle rental','Daily/weekly; driver licence and odometer/fuel inspections.'),('Carptrac','Forklift, crane, excavator, generator hire','Hourly/daily; meter readings; qualified operator optional/required by service.'),('Shipping limited hire','Portable toilet, large bin and scaffolding hire','Day/week/month/unit; delivery/collection, cleaning, setup, labour and quantity/site details.')
],[1.35,2.6,3.15])

page(doc); module_header(doc,'7.','Customer records and customer portal','Maintain one authoritative customer record while separating rental eligibility from online account security.','Customer, Rental Officer, Administrator, Finance Officer','customers.manage_access for portal security; operational and financial permissions by function and scope.')
table(doc,['ID','Function','Processing / rules','Outcome'],[
    ('F-CUS-01','Create/search customer','Search before create; unique customer number; classify Individual or Corporate; capture contacts, IDs/business details, address and status.','Customer profile created without duplicate.'),('F-CUS-02','Corporate account','Multiple contacts, approved renters, PO requirement, payment terms, credit limit, contracted rates and consolidated activity.','Corporate controls applied to bookings.'),('F-CUS-03','Enable portal access','Invite existing customer or register and link after verification; portal state independent from business/rental status.','Activation email and auditable account link.'),('F-CUS-04','Customer sign-in','Customer credential flow is separate from /staff; only own/corporate-authorized data is returned.','Customer portal session.'),('F-CUS-05','Portal workspace','Show bookings, quotation, agreement, invoice/payment status, documents and activity; allow cancellation/extension/incident requests.','Request is recorded for staff decision.'),('F-CUS-06','Security actions','Resend activation, reset password, verify email, lock/unlock portal, revoke sessions and view failed attempts.','Security state changes without changing rental eligibility.'),('F-CUS-07','Public status lookup','Allow reference-based status lookup only with a second matching value or signed secure link.','Minimal non-sensitive status returned; no customer PII exposed.')
],[.8,1.6,3.55,1.25],font_size=8.0)
heading(doc,'7.1 Customer field rules',2)
table(doc,['Field group','Individual','Corporate','Validation'],[
    ('Identity','Full name, DOB where approved, ID type/number, driver licence when driving.','Legal name, registration/TIN where approved, account contacts and authorized renters.','Required fields vary by service; sensitive values masked in lists.'),('Contact','Email, mobile, address.','Billing and operational contacts, email, phone and addresses.','Email format; phone normalized; verified flags stored separately.'),('Commercial','Deposit/payment status.','Credit limit, available credit, terms, PO rule and contract rate.','Credit/eligibility checked before confirmation.'),('Access','Portal enabled, email verified, account state.','Contact-level portal authority.','Rental status and portal status remain independent.')
],[1.1,2.15,2.25,1.6])

page(doc); module_header(doc,'8.','Asset register, QR, lifecycle and availability','Maintain a complete operational and financial record for each rentable asset across all configured divisions.','Rental Officer, Branch Manager, Maintenance Officer, Driver/Operator, Finance and authorized administrators','assets.view/create/edit/transfer/inspect/record_meter/retire/view_financials as applicable.')
table(doc,['ID','Function','Processing / validation','Result'],[
    ('F-AST-01','Create asset','Unique asset number; division, branch, service and category required; capture registration/serial/VIN/engine, make/model/year, ownership, acquisition, insurance, warranty, location and image/doc references.','Commissioned asset profile.'),('F-AST-02','Dynamic attributes','Render category definitions by type: text, number, integer, date, yes/no or choice; required/unit/search/report/customer visibility flags.','Future asset types need configuration rather than schema changes.'),('F-AST-03','QR identification','Generate a non-sensitive QR code resolving to asset number/token; scan checks authorization and returns allowed next actions.','Fast verified asset selection.'),('F-AST-04','Meter reading','Record odometer, engine hours, operating hours, fuel or units with source, date and actor; reject rollback unless authorized correction reason.','Current reading and history updated.'),('F-AST-05','Lifecycle transition','Validate allowed transition and prerequisites; capture actor, timestamp, booking, meter, fuel, notes and evidence.','Status, lifecycle and audit updated atomically.'),('F-AST-06','Transfer / retirement','Require origin/destination, dispatch/receipt checks; retirement requires reason, approval and final values.','Availability and location reflect transfer/retirement.'),('F-AST-07','Performance profile','Display revenue, service/transport/operator revenue, costs, margin, utilization, downtime, cost per km/hour and break-even progress.','Financial tab shown only with permission.')
],[.8,1.55,3.75,1.1],font_size=7.9)
heading(doc,'8.1 Availability engine',2)
table(doc,['Conflict source','Blocking condition','User-facing response'],[
    ('Booking','Overlapping non-cancelled/non-expired booking item.','Unavailable with booking reference to authorized staff.'),('Maintenance','Open/in-progress/waiting job overlapping requested period.','Unavailable — maintenance; expected release if known.'),('Inspection','Required inspection incomplete or asset status Inspection.','Unavailable until checklist passes.'),('Transfer','Asset in transit or assigned to another branch for requested period.','Unavailable at selected branch; show expected location.'),('Branch calendar','Closure/cutoff incompatible with pickup/return.','Offer valid time/branch; apply after-hours rule if allowed.'),('Personnel','Required operator unavailable or qualification invalid.','Do not confirm until eligible person is allocated.')
],[1.3,3.1,2.55])

page(doc); module_header(doc,'9.','Booking requests, quotations and approvals','Convert customer demand into a correctly priced, conflict-free and approved booking using a visible work queue.','Customer, Rental Officer, Branch Manager, authorized approvers','Manage rentals within scope; rentals.approve and pricing.configure where applicable.')
heading(doc,'9.1 Request and work-queue flow',2)
table(doc,['ID','Step','Required inputs / checks','Next state / action'],[
    ('F-BKG-01','Submit request','Customer, division/service, asset/category, branch, start/end, purpose/site and required personnel/delivery details.','New request; reference issued and acknowledgement queued.'),('F-BKG-02','Triage request','Check duplicates, customer eligibility, credit/PO flags, service documents and preliminary availability.','Prepare quotation, request information or reject with reason.'),('F-BKG-03','Allocate asset','Select available in-scope asset; detect booking, maintenance, transfer, inspection and calendar conflict.','Asset reserved provisionally or alternatives shown.'),('F-BKG-04','Confirm booking','Accepted quote, approval complete, customer eligible and no current conflict.','Confirmed; pickup requirements generated.'),('F-BKG-05','Cancel/expire','Record permitted reason; release reservations and notify affected parties.','Cancelled or expired; audit and availability updated.')
],[.8,1.35,3.8,1.3],font_size=8.0)
page(doc); heading(doc,'9.2 Pricing and quotation engine',2)
table(doc,['Calculation stage','Rule'],[
    ('Rate resolution','Corporate contract → branch override → division rate → category/asset rate → standard default. Most specific active rate valid for the hire period wins.'),('Hire duration','Calculate billable quantity using configured unit and minimum hire period; respect partial-day/hour and branch calendar rules.'),('Adjustments','Apply weekend, holiday and overtime multiplier; included usage and excess km/hour; operator/driver, delivery zone/distance/trip, setup, labour, fuel and other configured charges.'),('Commercial controls','Apply approved discount, deposit and customer terms; distinguish customer-visible selling rate from internal cost rate.'),('Tax','Apply configured tax rate to taxable lines; support tax-inclusive or tax-exclusive presentation. Rounding occurs per approved finance rule.'),('Quote total','Subtotal − discount + additional charges + tax. Deposit is displayed separately unless policy includes it in amount due.'),('Revision','Every revision receives version, creator, timestamp, reason and line snapshot; issued revisions remain read-only.'),('Conversion','Only accepted, unexpired, approved quotation converts; conversion is idempotent and creates/updates the booking once.')
],[1.45,5.6],font_size=8.3)

page(doc); heading(doc,'9.3 Configurable approval workflow',2)
table(doc,['ID','Function','Business behaviour'],[
    ('F-APR-01','Workflow definition','Configure division/branch, trigger, value band, number/order of stages, approver role or named employee, SLA and active dates.'),('F-APR-02','Submission','Create immutable approval instance from the applicable workflow; freeze decision inputs and assign stage 1.'),('F-APR-03','Decision','Approver may approve, reject or request revision with reason. Actor must be eligible, active and not prohibited by separation-of-duties rules.'),('F-APR-04','Progression','Approval advances sequentially; booking/quote remains pending until all required stages approve.'),('F-APR-05','Delegation','Use an active dated delegate when the primary approver is unavailable; record both the original assignment and acting approver.'),('F-APR-06','Escalation','Scheduled job flags or reassigns delayed stages according to policy and notifies management.'),('F-APR-07','Optional workflow','Where no applicable workflow exists, authorized low-risk transactions can proceed without approval; reason is recorded.')
],[.85,1.55,4.95],font_size=8.2)
heading(doc,'9.4 Booking workspace layout',2)
table(doc,['Panel','Content','Primary action examples'],[
    ('Summary','Reference, customer, division/service, branch, dates, value, stage and warnings.','Prepare quotation / Submit approval / Confirm.'),('Customer','Contacts, status, eligibility, credit and required documents.','Resolve missing information.'),('Asset & availability','Requested category, allocated asset, conflict explanation and alternatives.','Allocate or replace asset.'),('Pricing','Customer-visible lines, deposit, discount, tax, total; internal cost limited by permission.','Revise or issue quotation.'),('Approval & documents','Stage history, decision reasons, quote versions, attachments.','Submit / decide / send.'),('Activity','Notes, emails, status and audit timeline, collapsed by default.','Add internal note.')
],[1.35,3.65,2.35])

page(doc); module_header(doc,'10.','Pickup, agreements, active hire and returns','Use guided, controlled workflows to place an asset on hire and return it with complete evidence and correct financial consequences.','Rental Officer, Driver/Operator, Customer, Branch Manager, Finance and Maintenance staff','Rental-operation access plus asset inspection/meter and approval permissions where needed.')
heading(doc,'10.1 Pickup and signed agreement',2)
table(doc,['Step','System requirement','Completion result'],[
    ('1 Confirm customer','Booking confirmed; customer identity, authorized renter and contact checked.','Customer prerequisite complete.'),('2 Confirm/scan asset','QR resolves to allocated asset and branch; no conflict or blocking status.','Correct asset locked to workflow.'),('3 Verify documents','Licence/ID and service documents verified; corporate PO/credit requirements satisfied.','Compliance prerequisite complete.'),('4 Record condition','Pre-hire template complete with meter, fuel, responses, evidence and damage baseline.','Passed inspection or maintenance/damage referral.'),('5 Review charges','Accepted quote, deposit/payment requirement and approved changes displayed.','Agreement snapshot ready.'),('6 Sign and approve','Customer signs; agent verifies and approves. Signature consent, signer name and timestamp retained.','Agreement becomes immutable.'),('7 Check out','Generate PDF, queue email, write lifecycle/audit, status On hire/Rented.','Active hire appears in work queue.')
],[1.25,4.4,1.7],font_size=8.2)
callout(doc,'Agreement integrity','A signed agreement is never edited. Corrections use a numbered addendum linked to the original, with reason, approver and signatures where required.',PALE_GREEN)
heading(doc,'10.2 Active hire queues',2)
table(doc,['Queue','Included records','Displayed next action'],[
    ('Pickup today','Confirmed hires due for collection, incomplete prerequisites visible.','Continue pickup.'),('On hire','Checked-out assets not yet due.','View hire / report incident / extension request.'),('Due today','Expected return today.','Start return.'),('Overdue','Return date/time passed and not completed.','Contact customer / calculate late charge / escalate.'),('Return in progress','Post-hire workflow started but not finalized.','Continue return.'),('Recently completed','Completed within configured period.','View invoice, documents and history.')
],[1.35,3.65,2.3])

page(doc); heading(doc,'10.3 Return workflow',2)
table(doc,['Step','Processing and validation','System output'],[
    ('1 Scan and identify','QR must match active hire asset; otherwise require authorized exception.','Correct booking/customer loaded.'),('2 Record return','Capture actual date/time, location, meter and fuel; readings cannot be below handover without correction reason.','Usage and lateness calculated.'),('3 Post-hire inspection','Complete category checklist; compare baseline; capture photos, damage, customer acknowledgement and notes.','Outcome Passed / Passed with notes / Failed / Damage detected.'),('4 Calculate adjustments','Late time, excess usage, fuel, cleaning, damage and other approved lines.','Final charge preview with internal/customer separation.'),('5 Refer issues','Fault/damage can create linked maintenance job or damage incident.','Asset placed in Inspection, Maintenance or Out of service.'),('6 Complete return','Finalize invoice; save signatures; create lifecycle and audit records.','Booking Completed; asset status determined automatically.'),('7 Customer output','Generate/send return summary and invoice or amount-due notice.','Delivery event visible to staff.')
],[1.35,4.25,1.75],font_size=8.2)
heading(doc,'10.4 Automatic next-status rule',2)
table(doc,['Condition after return','Next asset status'],[
    ('Inspection passed and no maintenance due','Available'),('Minor notes require review','Inspection'),('Fault, failed checklist or maintenance threshold reached','Maintenance'),('Unsafe or serious damage','Out of service'),('Approved retirement/disposal workflow','Retired')
],[4.9,2.4])

page(doc); module_header(doc,'11.','Inspections and maintenance','Provide configurable evidence-based checks and detailed maintenance expenditure so asset condition, downtime and profitability remain trustworthy.','Rental Officer, Driver/Operator, Maintenance Officer and Branch Manager','assets.inspect, assets.record_meter, maintenance.complete and asset-scope access.')
heading(doc,'11.1 Inspection templates',2)
table(doc,['Asset group','Checklist examples','Mandatory evidence / special handling'],[
    ('Motors vehicle','Exterior/body, tyres, glass/lights, interior, documents, tools, fuel, odometer and existing damage map.','Pre/post signatures; licence and ID at pickup; photos for damage.'),('Carptrac forklift/crane/excavator','Fluids/leaks, tyres/tracks, forks/boom/attachments, hydraulics, brakes, alarms, safety devices and engine hours.','Qualified operator check; failed safety item blocks hire.'),('Carptrac generator','Oil/coolant/fuel, battery, cables, controls, output, voltage/phase, operating hours and test run.','Reading and test result; failed electrical/safety item blocks hire.'),('Shipping portable toilet','Shell/door/lock, tank, cleanliness, consumables, site condition and unit quantity.','Delivery/collection photos; cleaning/service requirement.'),('Shipping large bin','Structure, floor, doors/locks, contamination, damage, contents restriction and site access.','Photos, quantity, delivery/collection condition.'),('Shipping scaffolding','Component quantities, frames/braces/planks, damage/corrosion, load/safety tags and returned count.','Quantity reconciliation; missing/damaged components become charge/referral.')
],[1.55,3.35,2.45],font_size=7.9)
heading(doc,'11.2 Maintenance processing',2)
table(doc,['ID','Function','Rules and outputs'],[
    ('F-MNT-01','Create maintenance job','Asset, job number, type, priority, issue, reported date and source required; set asset Maintenance/Out of service when blocking.'),('F-MNT-02','Plan preventive work','Rules may be every N kilometres, N engine hours, N months or whichever occurs first; scheduled job created at threshold.'),('F-MNT-03','Record work and cost','Labour hours/rate, parts and quantity/cost, supplier, external invoice, transport/fuel/other cost, warranty and notes.'),('F-MNT-04','Manage status','Open → In progress → Waiting for parts → Completed/Cancelled. Completion requires work summary and authorization.'),('F-MNT-05','Return to service','Safety/inspection requirement must pass; meter and next-service targets saved; asset returns Available only if no other blocker.'),('F-MNT-06','Analyze history','Show repeat failures, downtime, total cost, maintenance cost versus replacement/book value and upcoming work.')
],[.85,1.7,4.9],font_size=8.2)

page(doc); module_header(doc,'12.','Personnel, transport and field operations','Coordinate operators, drivers and delivery activity required by equipment and site-based hire services.','Driver, Operator, Rental Officer, Branch Manager and authorized administrators','Scoped assignment access; personnel configuration and timesheet privileges.')
table(doc,['ID','Function','Functional rules'],[
    ('F-OPS-01','Personnel profile','Store employee/reference, branch/division, contact, hourly/overtime cost, availability and active state.'),('F-OPS-02','Qualifications','Record licence/certificate type, number, issue/expiry and evidence; expired/incorrect qualification blocks assignment.'),('F-OPS-03','Assign operator/driver','Check overlap, availability, qualification, branch/division and booking period; retain role and planned hours.'),('F-OPS-04','Timesheet','Assigned person records start/end/break/overtime and notes; manager approval determines payable/chargeable hours.'),('F-OPS-05','Delivery/collection','Capture zone, address/site contact, truck/driver, scheduled window, distance/trips and instructions; status Planned → Dispatched → Delivered/Collected/Failed.'),('F-OPS-06','Proof of service','Capture QR, photos, meter/unit count, recipient name/signature, timestamp and failure reason.'),('F-OPS-07','Cost and revenue','Separate customer operator/transport charge from employee/vehicle/internal cost for profitability.')
],[.85,1.75,4.85],font_size=8.2)
heading(doc,'12.1 Mobile field workspace',2)
for x in ['Show only the signed-in user’s current and upcoming assignments.','Allow QR resolve, checklist, meter, time, photograph and signature capture on responsive screens.','Do not expose customer financials, unrelated assets or other personnel records.','When connectivity fails, preserve unsent form state locally only if an approved secure offline design is implemented; baseline does not claim offline synchronization.']: bullet(doc,x)

page(doc); module_header(doc,'13.','Finance, profitability, dashboards and reports','Produce reliable customer charges and management insight while keeping internal costs and financial actions permission-controlled.','Finance Officer, Branch Manager, Division/Group Management and authorized administrators','reports.financial, payments.refund, assets.view_financials and scoped operational access.')
heading(doc,'13.1 Invoices, payments and statements',2)
table(doc,['ID','Function','Rules / result'],[
    ('F-FIN-01','Invoice line editor','Draft/issued unpaid invoice lines include description, quantity, unit price, tax rate and taxable flag; Paid/Voided invoices are locked.'),('F-FIN-02','Totals','Subtotal=sum quantity×unit price; tax follows inclusive/exclusive setting; total and balance update after every permitted change.'),('F-FIN-03','Payment allocation','Positive allocation cannot exceed unallocated payment or invoice balance; update Partially paid/Paid atomically.'),('F-FIN-04','Credit note/refund','Require reason, permission and approval where configured; value cannot exceed eligible invoice/payment amount.'),('F-FIN-05','Credit control','Before confirmation, compare exposure against limit/terms and PO requirement; route exception to approval.'),('F-FIN-06','Statement','Show customer invoices, credits, allocations, total, paid and outstanding across authorized divisions; provide PDF/email output.'),('F-FIN-07','Accounting export','Export approved transaction fields with stable IDs, dates, tax, branch/division and reference for later accounting integration.')
],[.85,1.65,4.95],font_size=8.1)
heading(doc,'13.2 Asset profitability',2)
table(doc,['Metric','Calculation / source'],[
    ('Rental and service revenue','Customer-visible earned hire, operator, delivery, excess usage, cleaning/damage recovery and other approved revenue attributed to asset.'),('Operating cost','Maintenance parts/labour/contractor, operator, transport, fuel, cleaning, insurance, registration, storage, damage and configured allocations.'),('Gross profit','Attributed revenue − attributed operating cost.'),('Profit margin','Gross profit ÷ attributed revenue × 100; blank/NA when revenue is zero.'),('Cost per unit','Relevant operating cost ÷ valid km/hour/day/unit usage; never divide by zero.'),('Utilization','Billable on-hire time or units ÷ available rentable capacity for the selected period.'),('Trend / warning','Monthly series; flag loss-making assets using configurable threshold and period.')
],[1.65,5.65],font_size=8.3)

page(doc); heading(doc,'13.3 Role-specific dashboards and reports',2)
table(doc,['Audience','Dashboard priorities','Reports / outputs'],[
    ('Rental Officer','Today’s requests, pickups, returns, incomplete prerequisites and overdue hires.','Rental history, availability and customer rental history.'),('Branch Manager','Approvals, conflicts, overdue, low availability, utilization and branch revenue/expense/profit.','Branch fleet/equipment utilization, overdue and profitability.'),('Maintenance Officer','Due/overdue service, faults, waiting parts and downtime.','Maintenance expenditure, repeat failure and downtime.'),('Finance Officer','Invoices due, unallocated payments, credit issues and exceptions.','Revenue, outstanding balances, customer profitability and statements.'),('Administrator','Configuration/access exceptions, email failures, audit and system health.','Access, audit and configuration reports.'),('Management','Division/branch performance, assets producing losses and trends.','Revenue by division/branch/category; asset P&L; utilization; export CSV/Excel/PDF.'),('Customer','Own upcoming/active/past bookings, quote/agreement/invoice and requests.','Own documents and account statement where enabled.')
],[1.25,3.25,2.65],font_size=8.1)
heading(doc,'13.4 Report controls',2)
for x in ['Date range, division, branch, service, category, asset and customer filters are constrained by actor scope.','Report totals use the same source records and calculation services as operational screens.','Exports carry report title, generation time, filters, currency/tax basis and user scope.','Large exports run asynchronously or stream results; screen reports use server aggregation and pagination.']: bullet(doc,x)

page(doc); module_header(doc,'14.','Notifications, documents and scheduled processing','Deliver required communications reliably, generate consistent records and automate time-based operational controls.','Customers, staff, administrators and system scheduler','Template/configuration permissions; recipient and record scope checked before sending.')
table(doc,['Event','Recipient / channel','Content and rule'],[
    ('Account activation / reset / verification','Target email','Single-use expiring link/code; generic public response; never include password.'),('Booking request / confirmation','Customer and relevant branch staff','Reference, service, dates, branch, next step and secure portal link.'),('Quotation','Authorized customer contacts','Versioned PDF/secure link, expiry and acceptance instructions.'),('Agreement / addendum','Customer and branch records','Immutable signed PDF; addendum separately numbered.'),('Invoice / statement','Authorized billing contacts','PDF/secure download and due/balance information.'),('Operational alerts','Scoped staff/manager','Overdue, maintenance due, low availability, approval delay, expiry, unpaid invoice, excess usage or unresolved damage.'),('Delivery failure','Administrator/operations','Failure category, retry state and safe diagnostic reference; no secret values.')
],[1.5,2.0,3.6],font_size=8.2)
heading(doc,'14.1 Scheduled jobs',2)
table(doc,['Job','Frequency / idempotency','Result'],[
    ('Expire quotations','Configured interval; repeat does not duplicate state/event.','Open expired quotes marked Expired; reservations released if policy says so.'),('Detect overdue rentals','Frequent interval based on branch local time.','Overdue flag/alert and notification queue.'),('Maintenance threshold','Daily and after meter entry.','Due/overdue job or reminder created once.'),('Credential/document expiry','Daily look-ahead.','Licence, insurance and warranty alerts.'),('Delayed approvals','Configured SLA interval.','Escalation/delegation and notification.'),('Email delivery','Queue polling with bounded retries.','Sent, failed or retry state plus provider reference.'),('Availability/profit flags','Scheduled management refresh.','Low fleet and loss-making alerts without changing source financial data.')
],[1.6,2.6,2.9],font_size=8.2)
heading(doc,'14.2 Document control',2)
for x in ['Store document metadata and private storage reference; do not store large data URLs in operational tables for production.','Validate file type, size and content; use private signed download links and access auditing.','Generated PDF stores template version and source-record snapshot/hash.','Retention, malware scanning and versioning require production policy confirmation.']: bullet(doc,x)

page(doc); module_header(doc,'15.','API, validation, errors and audit behaviour','Define consistent service behaviour so web clients, tests and future integrations interact safely and predictably.','Web clients, authorized integrations, administrators and support staff','Endpoint policy plus server-side permission and scope validation.')
heading(doc,'15.1 API conventions',2)
table(doc,['Concern','Required behaviour'],[
    ('Transport','HTTPS in staging/production; JSON UTF-8; /api route prefix; explicit request/response contracts.'),('Authentication','Secure, HttpOnly, SameSite cookie plus window-bound session token where designed; CSRF protection for cookie-authenticated mutations.'),('Timeouts','Inbound request timeout and downstream email/document/database timeouts are bounded. Long reports/jobs use asynchronous processing rather than keeping requests open.'),('Validation','Reject malformed/invalid input before mutation; trim text; validate ranges, dates, ownership, state and concurrency on server.'),('Pagination','List endpoints accept page/pageSize/search/filter/sort within configured limits and return total/page metadata.'),('Idempotency','Quote conversion, payment allocation, checkout/return and external retry-sensitive actions resist duplicate submission.'),('Concurrency','Use version/concurrency token or re-read status for sensitive transitions; stale updates return conflict.'),('Data exposure','Return only required fields; financial/internal cost and cross-scope records excluded at query level.'),('Correlation','Every request receives correlation ID propagated to safe errors, audit and logs.'),('Rate limiting','Login, password reset, public booking/status and costly export endpoints use stricter limits.')
],[1.45,5.85],font_size=8.3)
heading(doc,'15.2 Error contract',2)
table(doc,['HTTP status','Meaning','Client behaviour'],[
    ('400','Validation/business input error','Show field/reason; preserve user input.'),('401','Missing, expired or invalid session','Clear local session and return to correct login.'),('403','Authenticated but not permitted/in scope','Explain access restriction without exposing record data.'),('404','Record not found or deliberately concealed across scope','Show not found; do not infer existence.'),('409','State/concurrency/availability conflict','Refresh workspace and show current state or alternatives.'),('422','Optional for multi-rule validation','Show business-rule failures as structured issues.'),('429','Rate limit exceeded','Show retry guidance; do not loop automatically.'),('500/503','Unexpected/unavailable dependency','Safe message, correlation ID and bounded retry option.')
],[1.0,2.75,3.55],font_size=8.2)

page(doc); heading(doc,'15.3 Audit and logging',2)
table(doc,['Event category','Minimum audit fields','Additional rule'],[
    ('Security','Actor/target, outcome, time, IP/device, session/challenge reference.','Passwords, OTPs, reset tokens and secrets are never logged.'),('Access/configuration','Actor, target, before/after, role/permission/scope, reason.','Super Administrator and division changes are high priority.'),('Rental lifecycle','Booking/asset/customer references, action, old/new status, time and branch/division.','Signed agreement and inspection versions linked.'),('Financial','Invoice/payment/credit reference, amount/currency, actor, approval and before/after balance.','Append-only business audit; correction is a new event.'),('Maintenance/inspection','Asset/job/template, outcome, meter, cost, evidence references and actor.','Evidence access is separately audited.'),('System/integration','Job, correlation ID, attempt, provider status, duration and error category.','Production logs exclude personal message bodies unless approved.')
],[1.4,3.55,2.3],font_size=8.2)
heading(doc,'15.4 Transaction boundaries',2)
for x in ['Checkout saves inspection/agreement state, booking status, asset status, lifecycle and audit together; failure rolls back business-state changes.','Return saves inspection, final charges/invoice state, booking/asset status, maintenance referral and audit consistently.','Email delivery is queued after the business transaction; email failure does not undo a valid checkout or payment.','Approval decisions and payment allocations use database constraints/concurrency checks to prevent duplicates.']: bullet(doc,x)

page(doc); heading(doc,'16. Functional acceptance and traceability',1)
heading(doc,'16.1 Acceptance scenarios',2)
table(doc,['Scenario','Given / when','Expected observable result'],[
    ('AC-01 Scoped staff login','Branch Manager assigned to Motors Suva signs in.','Only permitted Motors/Suva work, assets, reports and menu items are returned; administration branches page is not accessible.'),('AC-02 Customer booking','Verified customer selects division/service/dates and submits complete request.','Reference issued, request visible in correct branch queue and confirmation email queued.'),('AC-03 Conflict prevention','Officer allocates asset overlapping maintenance or confirmed booking.','409/conflict shown with reason; booking is not confirmed.'),('AC-04 Flexible equipment pricing','Carptrac forklift requires operator and delivery.','Quote includes hire, operator hours and transport with customer total and separate internal cost.'),('AC-05 Multi-stage approval','High-value quote matches three-stage workflow.','Stages process sequentially; creator cannot self-approve; confirmation waits for final approval.'),('AC-06 Vehicle checkout','Confirmed Motors booking has verified ID/licence, payment and passed pre-hire check.','Signed immutable PDF generated, email queued, asset On hire and lifecycle/audit recorded.'),('AC-07 Damaged return','Post-hire comparison records new damage.','Damage/maintenance referral created, final charges reviewed, asset not Available.'),('AC-08 Preventive maintenance','New meter reaches configured threshold.','One maintenance reminder/job is created and shown to scoped maintenance staff.'),('AC-09 Shipping hire','Portable toilet order includes quantity, delivery and cleaning.','Availability/quantity validated; quote includes unit hire, delivery and cleaning; delivery proof supported.'),('AC-10 Asset profitability','Manager filters one asset and month.','Revenue, attributable costs, profit, margin and utilization reconcile to underlying records.'),('AC-11 Session security','Idle timeout passes or session cookie copied to unmatched window.','Protected request receives 401 and no protected data is shown.'),('AC-12 Failed email','Provider rejects quotation email.','Quote remains issued; email marked failed/retryable; staff sees delivery state and audit reference.')
],[1.05,3.0,3.25],font_size=7.9)

page(doc); heading(doc,'16.2 Functional traceability matrix',2)
table(doc,['Business objective','FSD functions','Evidence for acceptance'],[
    ('Secure access and isolation','F-AUTH-01–05, F-IAM-01–06','Authorization tests, session tests, scope-isolation tests and audit review.'),('Flexible multi-division operation','F-ORG-01–06, F-AST-02','Configured Motors, Carptrac and limited Shipping services without core schema changes.'),('Complete asset tracking','F-AST-01–07, inspection and return functions','Asset profile, QR, meters, inspections, lifecycle and status history.'),('Reliable hire workflow','F-BKG-01–05, F-APR-01–07, Sections 10–12','End-to-end request-to-return test with documents and assignments.'),('Financial visibility','F-FIN-01–07 and profitability metrics','Reconciled invoice/payment/asset P&L examples and exports.'),('Operational control','Sections 11, 13 and 14','Maintenance rules, dashboards, scheduled alerts and report filters.'),('Accountability','Section 15.3','Audit records for security, configuration, lifecycle and finance actions.')
],[2.0,2.4,2.9],font_size=8.3)
heading(doc,'16.3 Definition of functionally complete',2)
for x in ['All mandatory flows and negative cases pass in a production-like staging environment.','Permissions and division/branch isolation are verified with representative accounts.','Carpenters confirms field names, forms, checklists, statuses, price/tax rules, documents and reports.','Email/PDF/QR behaviours are tested with approved non-production credentials and recipients.','No critical/high defects remain; accepted lower-priority limitations are documented.']: bullet(doc,x)

page(doc); heading(doc,'17. Assumptions, open decisions and change control',1)
heading(doc,'17.1 Assumptions',2)
table(doc,['ID','Assumption','Impact if incorrect'],[
    ('AS-01','FJD and 15% tax are defaults but configurable.','Pricing and document totals require revision.'),('AS-02','A branch may support multiple divisions; services are enabled per branch.','Organization/scoping model changes.'),('AS-03','Shipping scope is limited to portable toilets, large bins and scaffolding hire.','Additional shipping processes require change request.'),('AS-04','Signed agreements occur at pickup after prerequisites pass.','Agreement lifecycle and UI changes.'),('AS-05','Email is a notification channel; business transactions remain valid if delivery fails.','Transaction/integration design changes.'),('AS-06','Customer online access and rental eligibility are separate statuses.','Customer lifecycle and controls change.')
],[.8,4.35,2.15],font_size=8.3)
heading(doc,'17.2 Decisions required before final sign-off',2)
table(doc,['Decision','Carpenters confirmation required'],[
    ('Pricing/tax rounding','Line versus invoice rounding, tax-inclusive/exclusive defaults, partial-day rules and deposit treatment.'),('Approval thresholds','Value bands, approver roles/names, stages, SLA, escalation and delegation.'),('Customer verification','Approved identity/business fields, retention, masking and public tracking verification.'),('Inspection templates','Final checklist items, mandatory photos/signatures, failure rules and damage recovery.'),('Shipping operations','Quantity inventory model, delivery zones, cleaning/setup rules and component reconciliation.'),('Security/session','Idle/absolute timeout, MFA policy, window-bound behaviour, password policy and lockout thresholds.'),('Documents/email','Approved templates, sender domain, recipient rules, retention and legal wording.'),('Finance integration','Source of truth, invoice numbering, payment process and export format.'),('Data migration','Validated mapping, cleansing, ownership and import acceptance criteria.')
],[2.0,5.3],font_size=8.3)
heading(doc,'17.3 Change-control workflow',2)
for x in ['Log change with rationale, requester, affected module and urgency.','Assess functional, security, data, UI, test, schedule and cost impact.','Obtain authorized business and project approval before implementation.','Update FSD/SRS version, traceability and acceptance cases; communicate the baseline.']: bullet(doc,x)

page(doc); heading(doc,'18. Client sign-off',1)
para(doc,'The undersigned confirm that this Functional Specification Document accurately describes the intended CREMS behaviour for the agreed baseline, subject to documented comments and approved change control.',after=10)
table(doc,['Approval statement','Selection / comment'],[
    ('Functional scope and workflows','☐ Approved   ☐ Approved with changes   ☐ Not approved'),('Roles, permissions and data visibility','☐ Approved   ☐ Approved with changes   ☐ Not approved'),('Pricing, approvals and finance behaviour','☐ Approved   ☐ Approved with changes   ☐ Not approved'),('Inspection, maintenance and asset lifecycle','☐ Approved   ☐ Approved with changes   ☐ Not approved'),('Documents, email and security behaviour','☐ Approved   ☐ Approved with changes   ☐ Not approved')
],[3.25,4.05],font_size=8.8)
table(doc,['Name','Role','Signature','Date'],[('','','',''),('','','',''),('','','','')],[1.8,2.0,2.1,1.1],font_size=9)
heading(doc,'Sign-off comments / required amendments',2)
for _ in range(5): para(doc,'________________________________________________________________________________',color=MID_GRAY,size=9,after=7)

page(doc); heading(doc,'Appendix A — Status models',1)
table(doc,['Object','Primary status sequence','Exceptions / notes'],[
    ('Staff/customer account','Invited → Activation pending → Active → Locked / Suspended / Deactivated','Reactivation requires authorized action; departure revokes sessions.'),('Booking','Draft / New request → Quotation/approval activity → Confirmed → Converted to rental → Completed','Cancelled and Expired are terminal; revisions remain in history.'),('Asset','Commissioned → Available → Reserved → Inspection → Rented/On hire → Inspection → Available','May branch to Maintenance, Out of service, transfer or Retired.'),('Maintenance','Open → In progress → Waiting for parts → Completed','Cancelled requires reason; return-to-service is separate validation.'),('Invoice','Draft → Issued → Partially paid → Paid','Voided/credit corrected by controlled finance actions.'),('Approval','Pending → Approved / Rejected / Cancelled','Multi-stage instance completes only after all required approvals.'),('Delivery','Planned → Dispatched → Delivered / Collected / Failed','Proof and failure reason required according to service.')
],[1.3,4.3,1.7],font_size=8.1)
heading(doc,'Appendix B — Output document catalogue',1)
table(doc,['Document','Generation trigger','Control'],[
    ('Quotation PDF','Issue or approved revision','Versioned, expiry shown, accepted version linked to booking.'),('Booking confirmation','Booking confirmation','Reference, service, dates, branch and prerequisites.'),('Rental agreement PDF','Customer signature and agent approval at pickup','Immutable snapshot; corrections by addendum.'),('Inspection report','Checklist completion','Template version, responses, readings, signatures and evidence references.'),('Return summary','Return completion','Usage, condition differences, adjustments and next asset status.'),('Invoice / credit note','Finance issue/action','Unique number, tax basis, lines, totals and balance.'),('Customer statement','Requested/scheduled corporate statement','Period, invoices, credits, payments and outstanding.'),('QR label','Asset commissioning/reprint','Non-sensitive token/code and human-readable asset number.'),('Management export','Authorized report export','Filters, scope, currency, generated time and user.')
],[1.45,2.65,3.2],font_size=8.2)

page(doc); heading(doc,'Appendix C — Functional requirement index',1)
table(doc,['Range','Module','Principal functions'],[
    ('F-AUTH-01–05','Authentication','Login, MFA, window-bound session, timeout and logout/revocation.'),('F-IAM-01–06','Identity and access','Staff lifecycle, scopes, permissions, security actions and separation of duties.'),('F-ORG-01–06','Organization','Divisions, services, branches, branch configuration, calendar and expansion.'),('F-CUS-01–07','Customers','Customer/corporate records, portal, access security and status tracking.'),('F-AST-01–07','Assets','Register, dynamic attributes, QR, meters, lifecycle, transfers and financial profile.'),('F-BKG-01–05','Bookings','Request, triage, allocation, confirmation and cancellation/expiry.'),('F-APR-01–07','Approvals','Workflow definition, submission, decisions, progression, delegation and escalation.'),('F-MNT-01–06','Maintenance','Jobs, preventive rules, work/cost, status, return to service and history.'),('F-OPS-01–07','Personnel/logistics','Profiles, qualifications, assignments, timesheets, delivery and proof.'),('F-FIN-01–07','Finance','Invoice, totals, allocation, credit/refund, controls, statements and export.')
],[1.4,1.6,4.3],font_size=8.2)
callout(doc,'End of controlled document','Any approved amendment must update the version history, affected function IDs and acceptance evidence.',PALE_GREEN)

# Core properties and save.
doc.core_properties.title='CREMS Functional Specification Document'
doc.core_properties.subject='Functional design baseline for Carpenters Motors, Carptrac and limited Shipping hire services'
doc.core_properties.author='CS400 IEP Project Team'
doc.core_properties.keywords='CREMS, FSD, Carpenters Fiji, rental, asset management'
doc.core_properties.comments='Prepared for client validation. Contains no credentials or production secrets.'
doc.save(OUT)
print(OUT)
