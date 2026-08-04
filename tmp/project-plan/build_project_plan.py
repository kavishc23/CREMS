from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/Users/kavishchandra/Documents/CS400')
OUT = ROOT / 'docs' / 'CREMS-Project-Plan.docx'
LOGO = ROOT / 'src' / 'crems-web' / 'public' / 'brand' / 'carpenters-logo.png'

YELLOW = 'FFED00'; BLACK = '111111'; CHARCOAL = '303030'; MUTED = '666666'
LIGHT = 'F5F5F1'; PALE = 'FFF9C7'; GREEN = '2E7D32'; RED = 'A52828'; WHITE = 'FFFFFF'

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
sec.top_margin = Inches(0.82); sec.bottom_margin = Inches(0.78)
sec.left_margin = Inches(1); sec.right_margin = Inches(1)
sec.header_distance = Inches(0.36); sec.footer_distance = Inches(0.38)

styles = doc.styles
def font(style, size, color=BLACK, bold=None):
    style.font.name = 'Calibri'; style._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri'); style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    style.font.size = Pt(size); style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None: style.font.bold = bold

font(styles['Normal'], 10.5)
styles['Normal'].paragraph_format.space_after = Pt(7)
styles['Normal'].paragraph_format.line_spacing = 1.22
for name, size, before, after in [('Heading 1',16,18,9),('Heading 2',13,12,6),('Heading 3',11.5,9,4)]:
    font(styles[name], size, BLACK, True)
    styles[name].paragraph_format.space_before = Pt(before); styles[name].paragraph_format.space_after = Pt(after); styles[name].paragraph_format.keep_with_next = True
for name in ['List Bullet','List Number']:
    font(styles[name], 10.5)
    styles[name].paragraph_format.left_indent = Inches(.38); styles[name].paragraph_format.first_line_indent = Inches(-.19)
    styles[name].paragraph_format.space_after = Pt(4); styles[name].paragraph_format.line_spacing = 1.2

def shade_paragraph(p, fill):
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); p._p.get_or_add_pPr().append(shd)
def border_left(p, color=YELLOW, size='20'):
    ppr = p._p.get_or_add_pPr(); pbdr = ppr.find(qn('w:pBdr'))
    if pbdr is None: pbdr = OxmlElement('w:pBdr'); ppr.append(pbdr)
    b = OxmlElement('w:left'); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),size); b.set(qn('w:space'),'8'); b.set(qn('w:color'),color); pbdr.append(b)
def cell_shade(cell, fill):
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); cell._tc.get_or_add_tcPr().append(shd)
def cell_margin(cell, top=90, bottom=90, start=120, end=120):
    tcpr=cell._tc.get_or_add_tcPr(); mar=tcpr.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); tcpr.append(mar)
    for k,v in [('top',top),('bottom',bottom),('start',start),('end',end)]:
        x=OxmlElement('w:'+k); x.set(qn('w:w'),str(v)); x.set(qn('w:type'),'dxa'); mar.append(x)
def repeat_header(row):
    x=OxmlElement('w:tblHeader'); x.set(qn('w:val'),'true'); row._tr.get_or_add_trPr().append(x)
def table_geometry(table, widths, indent=120):
    table.autofit=False; tblpr=table._tbl.tblPr
    w=tblpr.find(qn('w:tblW'))
    if w is None: w=OxmlElement('w:tblW'); tblpr.append(w)
    w.set(qn('w:w'),str(sum(widths))); w.set(qn('w:type'),'dxa')
    ind=OxmlElement('w:tblInd'); ind.set(qn('w:w'),str(indent)); ind.set(qn('w:type'),'dxa'); tblpr.append(ind)
    grid=table._tbl.tblGrid
    for c in list(grid): grid.remove(c)
    for width in widths:
        c=OxmlElement('w:gridCol'); c.set(qn('w:w'),str(width)); grid.append(c)
    for row in table.rows:
        cant_split=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant_split)
        for i,cell in enumerate(row.cells):
            tcw=cell._tc.get_or_add_tcPr().find(qn('w:tcW'))
            if tcw is None: tcw=OxmlElement('w:tcW'); cell._tc.get_or_add_tcPr().append(tcw)
            tcw.set(qn('w:w'),str(widths[i])); tcw.set(qn('w:type'),'dxa'); cell_margin(cell); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def page_field(p):
    run=p.add_run(); begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin'); instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=' PAGE '; end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end'); run._r.extend([begin,instr,end])
def header_footer(section):
    p=section.header.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    if LOGO.exists(): p.add_run().add_picture(str(LOGO),width=Inches(.25))
    r=p.add_run('   CREMS | Project Plan'); r.bold=True; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor.from_string(MUTED)
    p=section.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=p.add_run('CS400 Industry Experience Project  |  Page '); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(MUTED); page_field(p)
header_footer(sec)

def heading(text, level=1):
    p=doc.add_heading(text,level); 
    if level==1: border_left(p)
    return p
def para(text='', bold_prefix=None, italic=False, align=None):
    p=doc.add_paragraph(); p.alignment=align
    if bold_prefix and text.startswith(bold_prefix): p.add_run(bold_prefix).bold=True; p.add_run(text[len(bold_prefix):])
    else: r=p.add_run(text); r.italic=italic
    return p
def bullet(text): return doc.add_paragraph(text,style='List Bullet')
def number(text): return doc.add_paragraph(text,style='List Number')
def numbered_sequence(items):
    numbering=doc.part.numbering_part.element
    abstract_ids=[int(x.get(qn('w:abstractNumId'))) for x in numbering.findall(qn('w:abstractNum'))]
    num_ids=[int(x.get(qn('w:numId'))) for x in numbering.findall(qn('w:num'))]
    abstract_id=max(abstract_ids,default=0)+1; num_id=max(num_ids,default=0)+1
    abstract=OxmlElement('w:abstractNum'); abstract.set(qn('w:abstractNumId'),str(abstract_id))
    multi=OxmlElement('w:multiLevelType'); multi.set(qn('w:val'),'singleLevel'); abstract.append(multi)
    lvl=OxmlElement('w:lvl'); lvl.set(qn('w:ilvl'),'0')
    start=OxmlElement('w:start'); start.set(qn('w:val'),'1'); lvl.append(start)
    fmt=OxmlElement('w:numFmt'); fmt.set(qn('w:val'),'decimal'); lvl.append(fmt)
    text=OxmlElement('w:lvlText'); text.set(qn('w:val'),'%1.'); lvl.append(text)
    suff=OxmlElement('w:suff'); suff.set(qn('w:val'),'tab'); lvl.append(suff)
    ppr=OxmlElement('w:pPr'); tabs=OxmlElement('w:tabs'); tab=OxmlElement('w:tab'); tab.set(qn('w:val'),'num'); tab.set(qn('w:pos'),'540'); tabs.append(tab); ppr.append(tabs)
    ind=OxmlElement('w:ind'); ind.set(qn('w:left'),'540'); ind.set(qn('w:hanging'),'270'); ppr.append(ind); lvl.append(ppr); abstract.append(lvl); numbering.append(abstract)
    num=OxmlElement('w:num'); num.set(qn('w:numId'),str(num_id)); ref=OxmlElement('w:abstractNumId'); ref.set(qn('w:val'),str(abstract_id)); num.append(ref); numbering.append(num)
    for item in items:
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.2
        numpr=OxmlElement('w:numPr'); ilvl=OxmlElement('w:ilvl'); ilvl.set(qn('w:val'),'0'); nid=OxmlElement('w:numId'); nid.set(qn('w:val'),str(num_id)); numpr.extend([ilvl,nid]); p._p.get_or_add_pPr().append(numpr); p.add_run(item)
def note(title,text,warning=False):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.12); p.paragraph_format.right_indent=Inches(.12); p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(9); shade_paragraph(p,PALE if warning else LIGHT); border_left(p,YELLOW if warning else BLACK,'16'); r=p.add_run(title+': '); r.bold=True; p.add_run(text); return p
def make_table(headers, rows, widths, aligns=None, font_size=9):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER; repeat_header(t.rows[0])
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; cell_shade(c,YELLOW); r=c.paragraphs[0].add_run(h); r.bold=True; r.font.size=Pt(font_size)
    for data in rows:
        cells=t.add_row().cells
        for i,value in enumerate(data):
            cells[i].text=str(value)
            if aligns: cells[i].paragraphs[0].alignment=aligns[i]
            for r in cells[i].paragraphs[0].runs: r.font.size=Pt(font_size)
    table_geometry(t,widths)
    return t

# Cover - proposal centerpiece
doc.add_paragraph().paragraph_format.space_after=Pt(38)
if LOGO.exists():
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(LOGO),width=Inches(1.35))
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(18); p.paragraph_format.space_after=Pt(4)
r=p.add_run('PROJECT PLAN'); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=RGBColor.from_string(MUTED)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(6)
r=p.add_run('Car and Rental Equipment\nManagement System'); r.bold=True; r.font.size=Pt(27); r.font.color.rgb=RGBColor.from_string(BLACK)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(22)
r=p.add_run('CREMS'); r.bold=True; r.font.size=Pt(16); r.font.color.rgb=RGBColor.from_string(CHARCOAL)
note_p=doc.add_paragraph(); note_p.alignment=WD_ALIGN_PARAGRAPH.CENTER; note_p.paragraph_format.left_indent=Inches(.65); note_p.paragraph_format.right_indent=Inches(.65); shade_paragraph(note_p,YELLOW)
r=note_p.add_run('Client: Carpenters Fiji Pte Ltd\nCS400 Industry Experience Project | Semester 2, 2026'); r.bold=True; r.font.size=Pt(11)
doc.add_paragraph().paragraph_format.space_after=Pt(18)
make_table(['Project detail','Information'],[
    ('Group number','[INSERT GROUP NUMBER]'),('Project supervisor','Amit Kumar'),('Co-supervisor','Shanil Naidu'),('Version / date','Version 1.0 | 4 August 2026')
],[2300,7060],font_size=9.5)
doc.add_paragraph().paragraph_format.space_after=Pt(8)
make_table(['Team member','Student ID'],[
    ('Kavish Chandra','[INSERT STUDENT ID]'),('[INSERT TEAM MEMBER 2]','[INSERT STUDENT ID]'),('[INSERT TEAM MEMBER 3]','[INSERT STUDENT ID]'),('[INSERT TEAM MEMBER 4]','[INSERT STUDENT ID]')
],[6000,3360],font_size=9.5)
doc.add_page_break()

heading('Document control')
make_table(['Version','Date','Prepared by','Change'],[('1.0','4 Aug 2026','CREMS project team','Initial project plan for review')],[1000,1500,2600,4260])
heading('Approval record')
make_table(['Role','Name','Approval / signature','Date'],[
    ('Project leader','[INSERT NAME]','',''),('Client representative','[INSERT NAME]','',''),('Academic supervisor','Amit Kumar','','')
],[2100,2500,3260,1500])
heading('Executive summary')
para('Carpenters Fiji currently requires a centralized way to coordinate vehicle and equipment rentals across branches. CREMS will replace fragmented, manual and spreadsheet-based activities with a secure web application covering assets, customers, reservations, agreements, inspections, returns, maintenance, notifications and management reporting.')
para('The team will deliver the project as a modular web application using React, ASP.NET Core and Microsoft SQL Server. Work will be completed iteratively over 14 weeks, with requirements validation, fortnightly demonstrations and client acceptance checkpoints. The minimum viable product prioritizes accurate availability, role-based access, a traceable rental lifecycle and reliable operational reporting.')
note('Delivery recommendation','Protect the MVP boundary. Accounting/ERP integration, online payments, GPS tracking, mobile applications and external booking platforms remain out of scope unless approved through formal change control.')
heading('Project at a glance',2)
make_table(['Dimension','Plan'],[
    ('Duration','14 weeks: 3 August to 8 November 2026'),('Team','Four student developers with shared delivery responsibility'),('Method','Agile, two-week iterations, client review at each major milestone'),('Target users','Administrators, branch managers, rental officers and customers'),('Deployment','Responsive web frontend, REST API and relational database'),('Budget basis','Academic project; labour and existing hardware treated as in-kind contributions')
],[1900,7460])

heading('1. Background and problem analysis')
heading('1.1 Client context',2)
para('Carpenters Fiji manages vehicles and rental equipment that move through reservation, allocation, inspection, rental, return, maintenance and reporting activities. These activities involve multiple people and branches and require dependable, current information. Delays or inconsistencies can cause double bookings, unavailable assets being promised, incomplete condition evidence, missed maintenance and weak management visibility.')
heading('1.2 Problem statement',2)
para('The core problem is the absence of a single, role-controlled source of truth for rental assets and transactions. Operational records that are distributed across paper, spreadsheets or isolated processes make it difficult to confirm availability, preserve a complete transaction history and produce timely reports. CREMS must centralize these records while remaining usable for frontline staff and manageable by a four-person project team.')
heading('1.3 Information requirements',2)
for x in [
    'Asset identity, category, registration or serial details, branch, availability state and maintenance status.',
    'Customer identity and contact information, rental history, agreements and deposits.',
    'Booking dates, assigned assets, pricing references, extensions, returns and overdue status.',
    'Pre-rental and post-rental inspection results, damage, fuel, accessories and supporting notes.',
    'Role, branch and account status for each user, plus an audit trail of critical changes.',
    'Operational measures for utilization, revenue, overdue rentals, maintenance and customer history.'
]: bullet(x)
heading('1.4 Constraints and assumptions',2)
make_table(['Type','Constraint or assumption','Planning response'],[
    ('Time','One academic semester and fixed assessment deadlines.','Prioritize a demonstrable MVP and use staged acceptance.'),
    ('People','Four students working around other courses.','Clear ownership, peer review and weekly workload checks.'),
    ('Requirements','Some business rules remain subject to client confirmation.','Maintain a requirements register and decision log.'),
    ('Technology','Solution must be supportable on team laptops and a realistic server environment.','Use cross-platform .NET, React, Docker and SQL Server.'),
    ('Privacy','Customer and identity data must not appear in repositories or demonstrations.','Use synthetic data, secrets management and least privilege.'),
    ('Connectivity','Development and client review may occur across different networks.','Local environments plus portable demonstrations and backups.')
],[1250,4100,4010],font_size=8.7)

heading('2. Objectives and success measures')
heading('2.1 Project objectives',2)
for x in [
    'Deliver one centralized, responsive web application for vehicle and equipment rentals across branches.',
    'Prevent conflicting reservations through authoritative availability and date validation.',
    'Support the rental lifecycle from booking and agreement through allocation, inspection, return and closure.',
    'Give each user only the functions and records permitted by role and branch.',
    'Improve traceability through transaction history, inspection evidence and audit events.',
    'Provide dashboards and reports that support timely operational decisions.'
]: bullet(x)
heading('2.2 Measurable acceptance criteria',2)
make_table(['ID','Acceptance measure','Evidence'],[
    ('AC-01','Authorized users can sign in and receive the correct navigation and API permissions for all four roles.','Role-based API and UI test results'),
    ('AC-02','The system rejects an overlapping confirmed booking for the same asset and date range.','Automated test and demonstration'),
    ('AC-03','Staff can complete booking, allocation, inspection, return and closure using test data.','End-to-end acceptance scenario'),
    ('AC-04','Dashboard totals reconcile with underlying test transactions.','Report reconciliation checklist'),
    ('AC-05','Critical create/update/status actions record who acted and when.','Audit-log inspection'),
    ('AC-06','Core screens work at desktop and mobile widths and meet agreed usability checks.','Responsive test record'),
    ('AC-07','No critical or high-severity security defect remains open at handover.','Security checklist and issue register')
],[900,5500,2960],font_size=8.5)

heading('3. Scope')
heading('3.1 In scope - MVP',2)
for x in [
    'Authentication and role-based access for Administrator, Branch Manager, Rental Officer and Customer.',
    'Branch, vehicle and equipment master records, including status and identification code.',
    'Customer registration and profile management.',
    'Booking, availability checking, allocation, agreement, extension, return and closure.',
    'Inspection checklists, fuel/accessory records, damage reports and maintenance requests.',
    'Preventive maintenance schedules and reminders.',
    'Dashboards and agreed rental, utilization, maintenance, revenue, overdue and customer reports.',
    'Email notifications for booking, overdue rental and scheduled maintenance events.',
    'Audit logs for critical transactions and asset movements.'
]: bullet(x)
heading('3.2 Out of scope',2)
for x in ['Accounting or ERP integration.','Online payment gateway.','Native mobile application.','GPS vehicle tracking.','Third-party insurance integration.','Driver scheduling and payroll.','Loyalty or rewards programmes.','External booking-platform integration.']: bullet(x)
heading('3.3 Future options',2)
para('The architecture will preserve REST API boundaries and stable identifiers so approved future work can add payment, ERP, telematics or mobile clients without redesigning the MVP. These are design considerations, not current deliverables.')

heading('4. Solution approach and feasibility')
heading('4.1 Selected solution',2)
make_table(['Layer','Technology','Reason'],[
    ('Frontend','React 19, TypeScript, Vite, Material UI','Responsive component model, fast development and strong type checking.'),
    ('Backend','ASP.NET Core 10 Web API','Enterprise authorization, validation, performance and maintainability.'),
    ('Data','Microsoft SQL Server with Entity Framework Core','Relational integrity, migrations and natural alignment with the client brief.'),
    ('Identity','ASP.NET Core Identity with cookie authentication','Secure account management and server-enforced roles.'),
    ('Delivery','GitHub, Docker Compose, Postman and VS Code','Repeatable environments, reviewable changes and API verification.')
],[1500,3100,4760],font_size=8.8)
heading('4.2 Alternatives considered',2)
make_table(['Alternative','Strength','Reason not selected'],[
    ('Laravel + React','Rapid CRUD development and approachable PHP ecosystem.','Would require rewriting completed identity/API work and is less aligned with SQL Server and the approved brief.'),
    ('Java Spring + React','Mature enterprise platform and strong typing.','Higher setup overhead for the current team with no material project advantage.'),
    ('No-code/SaaS','Fast prototype creation.','Insufficient control over rental rules, auditability, extensibility and academic software-development outcomes.')
],[2100,3000,4260],font_size=8.7)
heading('4.3 Feasibility assessment',2)
make_table(['Dimension','Assessment','Conclusion'],[
    ('Technical','The stack is cross-platform and the team has already established authentication, database migrations, Docker and a React shell.','Feasible with disciplined testing.'),
    ('Operational','Browser-based workflows reduce installation burden; role-specific screens support staff responsibilities.','Feasible subject to client validation.'),
    ('Schedule','Four people over 14 weeks can deliver the MVP if integrations remain out of scope.','Feasible with strict change control.'),
    ('Economic','Open-source development tools and existing student hardware keep incremental expenditure low.','Feasible within academic constraints.'),
    ('Ethical/legal','Synthetic test data, least privilege, secure secrets and documented consent/privacy controls reduce exposure.','Feasible with ongoing review.')
],[1500,5100,2760],font_size=8.6)

heading('5. Delivery method and work breakdown')
heading('5.1 Agile operating model',2)
para('The team will work in two-week iterations. Each iteration begins with backlog refinement and task ownership, uses pull requests and peer review during implementation, and ends with an integrated demonstration. Client decisions are recorded in the requirements register. A feature is complete only when code, validation, tests and documentation are integrated.')
heading('5.2 Work packages',2)
make_table(['WP','Work package','Primary outputs'],[
    ('WP1','Discovery and planning','Validated requirements, scope, project plan, team charter, initial backlog'),
    ('WP2','Architecture and environments','Repository, Docker database, CI-ready builds, data model, API conventions'),
    ('WP3','Identity and administration','Authentication, roles, branch-aware accounts, user administration'),
    ('WP4','Assets and customers','Asset/equipment catalogue, customer records, identification codes'),
    ('WP5','Booking and availability','Reservation workflow, conflict rules, allocation and agreement'),
    ('WP6','Return, inspection and maintenance','Condition evidence, damage, fuel, maintenance scheduling'),
    ('WP7','Dashboard, reports and notifications','Operational measures, exports and email events'),
    ('WP8','Quality, deployment and handover','Security/performance tests, UAT, deployment guide, training and final report')
],[850,3150,5360],font_size=8.6)
heading('5.3 Definition of done',2)
for x in ['Acceptance criteria are clear and linked to a requirement.','Implementation builds successfully and follows agreed conventions.','At least one other member reviews the change.','Automated or repeatable manual tests pass.','Authorization is enforced by the API, not only hidden in the interface.','Documentation and demonstration data are updated.','The product owner/client accepts the outcome or records follow-up work.']: bullet(x)

heading('6. Schedule and milestones')
para('Dates are proposed for Semester 2, 2026 and should be confirmed with the client and academic supervisor. Work packages overlap deliberately so analysis, implementation and validation progress continuously.')
make_table(['Weeks / dates','Focus','Owner(s)','Exit deliverable'],[
    ('W1-2 | 3-16 Aug','Discovery, domain research, plan, charter and requirements baseline','All; Leader coordinates','Approved plan and prioritized MVP backlog'),
    ('W3 | 17-23 Aug','Architecture, database refinement and environment standardization','Backend + DevOps leads','Repeatable full-stack setup'),
    ('W4 | 24-30 Aug','Identity, roles, branches and user administration','Backend + Frontend leads','Role-based access demonstration'),
    ('W5-6 | 31 Aug-13 Sep','Asset/equipment and customer management','Pair A / Pair B','Asset and customer modules'),
    ('W7-8 | 14-27 Sep','Bookings, availability, allocation and agreements','All, feature pairs','End-to-end booking demonstration'),
    ('W9 | 28 Sep-4 Oct','Returns, extensions and closure','Rental workflow pair','Completed rental lifecycle'),
    ('W10 | 5-11 Oct','Inspections, damage and maintenance','Asset workflow pair','Inspection and maintenance modules'),
    ('W11 | 12-18 Oct','Dashboard, reports, QR/barcodes and notifications','Frontend/reporting pair','Management information demonstration'),
    ('W12 | 19-25 Oct','Integration, accessibility, security and performance testing','All; QA lead coordinates','Release candidate 1'),
    ('W13 | 26 Oct-1 Nov','Client UAT, defect correction and training material','All; Client liaison coordinates','UAT acceptance record'),
    ('W14 | 2-8 Nov','Deployment rehearsal, final documentation and handover','All','Final release and handover package')
],[1900,3500,1700,2260],font_size=8.1)
heading('6.1 Milestones',2)
make_table(['Milestone','Target','Acceptance condition'],[
    ('M1 - Plan baseline','16 Aug','Project plan, charter, scope, risks and backlog reviewed.'),('M2 - Technical foundation','30 Aug','All members can run frontend, API and database; role demonstration passes.'),('M3 - Core records','13 Sep','Assets and customers can be maintained with validation.'),('M4 - Rental workflow','4 Oct','Booking-to-return scenario succeeds without conflicting allocation.'),('M5 - Feature complete','18 Oct','All MVP modules integrated; only defects/documentation remain.'),('M6 - UAT complete','1 Nov','Agreed client scenarios pass or exceptions are formally recorded.'),('M7 - Handover','8 Nov','Release, source, database scripts, guides and final report delivered.')
],[2500,1600,5260],font_size=8.5)

heading('7. Resources and responsibilities')
heading('7.1 Human resources',2)
make_table(['Project role','Assigned person','Main responsibilities','Planned effort'],[
    ('Project leader / client liaison','[INSERT NAME]','Plan, client meetings, decisions, integration and submission control.','6 h/week'),
    ('Backend and data lead','[INSERT NAME]','Domain model, API, database, authorization and backend tests.','6 h/week'),
    ('Frontend and UX lead','[INSERT NAME]','Responsive UI, accessibility, workflow integration and UI tests.','6 h/week'),
    ('Quality and DevOps lead','[INSERT NAME]','Test planning, Docker, build checks, release evidence and documentation.','6 h/week')
],[2100,1900,3960,1400],font_size=8.3)
para('Roles establish accountability but do not create silos. Members will pair on high-risk features, review each other’s pull requests and rotate meeting-note and demonstration duties.')
heading('7.2 Technical and information resources',2)
for x in ['Four development laptops with Git, .NET 10 SDK, Node.js 22+, Docker Desktop, VS Code and Postman.','GitHub repository for source, pull requests, issues and protected review workflow.','SQL Server 2022 container for consistent local databases.','Client representatives for workflow validation, sample forms and acceptance decisions.','Academic supervisor for assessment guidance and escalation.','Synthetic test dataset; no production customer data stored in student environments.']: bullet(x)
heading('7.3 RACI summary',2)
make_table(['Activity','Leader','Backend','Frontend','QA/DevOps','Client'],[
    ('Requirements / scope','A','C','C','C','R/C'),('Architecture / data','C','A/R','C','C','I'),('User experience','C','C','A/R','C','C'),('Testing / release','C','R','R','A/R','C'),('Acceptance / change approval','R','C','C','C','A')
],[2300,1200,1450,1450,1500,1460],font_size=8.1)
para('Legend: A = accountable, R = responsible, C = consulted, I = informed.',italic=True)

heading('8. Budget')
para('This is an academic project. Student labour, personal laptops and existing internet access are in-kind contributions and are shown for transparency but are not charged to Carpenters Fiji. Any paid service requires prior written approval.')
make_table(['Item','Basis','Cash cost (FJD)','Treatment / justification'],[
    ('Student development effort','4 people x 6 h/week x 14 weeks = 336 hours','0','In-kind academic contribution.'),
    ('Development laptops','Existing personal equipment','0','No new purchase required.'),
    ('Software/tooling','.NET, React, VS Code, GitHub, Postman and Docker student/personal use','0','Open-source or existing free access.'),
    ('Local database','SQL Server Developer container','0','Development/test use only.'),
    ('Cloud demonstration hosting','Optional short-lived test deployment','120','Only if client requires remote access.'),
    ('Email/domain allowance','Optional test mailbox/domain','50','Supports realistic notification testing.'),
    ('Contingency','Approx. 20% of possible cash expenditure','35','Covers minor approved service costs.'),
    ('Total maximum cash budget','','205','Spend only with approval; expected cost may remain zero.')
],[2400,3000,1450,2510],font_size=8.3)
note('Budget control','The project leader records approved expenditure and evidence. Scope will not be expanded merely because contingency remains available.')

heading('9. Risk management')
para('Risks are reviewed weekly. Probability (P) and impact (I) use a 1-5 scale; exposure is P x I. Scores 15-25 are High, 8-14 Medium and 1-7 Low.')
make_table(['ID','Risk','P','I','Rating','Response / owner'],[
    ('R1','Requirements remain unclear or change late.','4','5','20 H','Fortnightly validation, decision log and change control / Leader'),
    ('R2','Scope exceeds semester capacity.','4','5','20 H','Protect MVP, prioritize backlog and defer integrations / All'),
    ('R3','Double-booking or status rules are incorrect.','3','5','15 H','Domain rules, transactions and overlap tests / Backend'),
    ('R4','One member becomes unavailable.','3','4','12 M','Shared documentation, pairing and small reviewed changes / Leader'),
    ('R5','Integration occurs too late.','3','4','12 M','Continuous main-branch integration and iteration demos / DevOps'),
    ('R6','Sensitive data or credentials are exposed.','2','5','10 M','Synthetic data, user secrets, review and secret scanning / All'),
    ('R7','Authorization allows excess access.','3','5','15 H','Server policies plus role-based negative tests / Backend + QA'),
    ('R8','Environment differences block a team member.','3','3','9 M','Docker, setup guide and verified Mac/Windows steps / DevOps'),
    ('R9','Client availability delays acceptance.','3','4','12 M','Book reviews early and submit focused decision requests / Liaison'),
    ('R10','Data loss or repository damage.','2','5','10 M','GitHub remote, protected branches and database backup rehearsal / DevOps')
],[650,3100,450,450,850,3860],font_size=7.9)
heading('9.1 Escalation thresholds',2)
for x in ['A High risk is discussed immediately with the project leader and given a dated mitigation task.','A decision that changes scope, delivery date, privacy posture or architecture is escalated to the client/supervisor.','A blocked task lasting more than two working days is raised at the next team contact, not held until the iteration review.']: bullet(x)

heading('10. Quality, security and professional practice')
heading('10.1 Quality plan',2)
make_table(['Control','Minimum evidence'],[
    ('Build quality','Backend build, frontend TypeScript build and lint checks pass.'),('Code review','At least one peer approval for functional changes.'),('API testing','Positive, validation, unauthenticated and unauthorized cases in Postman/automated tests.'),('Data quality','Required fields, relationships, uniqueness and date constraints verified.'),('Usability','Agreed workflows tested at desktop and mobile widths.'),('Release','UAT checklist, known-issues list and deployment rehearsal completed.')
],[2500,6860],font_size=8.8)
heading('10.2 Security and privacy',2)
for x in ['Use least-privilege roles and branch boundaries; enforce access in API policies and queries.','Store passwords through ASP.NET Identity hashing; never log or display passwords.','Keep connection strings and bootstrap credentials in local secret stores, not Git.','Use HTTPS in deployed environments, secure cookies and trusted-origin configuration.','Collect only information needed for rental operations and define retention/deletion expectations with the client.','Use synthetic customer and transaction data for development, demonstrations and assessment evidence.','Record critical user, rental and asset status changes in an auditable form.']: bullet(x)
heading('10.3 Ethical and professional conduct',2)
para('The team will protect client confidentiality, accurately report progress and limitations, respect software licences, acknowledge sources and avoid presenting generated or unverified material as client-approved fact. Safety-sensitive records such as vehicle condition and maintenance must not be altered without traceability. Team decisions will consider the interests of customers, staff, the client and the wider public.')

heading('11. Communication, governance and change control')
make_table(['Forum','Frequency','Participants','Output'],[
    ('Team stand-up','Twice weekly','Four students','Progress, next actions and blockers'),('Planning/review','Every two weeks','Team; client when available','Iteration goal, demonstration and accepted feedback'),('Client checkpoint','Fortnightly or as agreed','Liaison, relevant members, client','Decisions, clarified rules and acceptance notes'),('Supervisor update','As scheduled','Leader/team and supervisor','Academic guidance and escalations'),('Pull-request review','For each change','Author plus reviewer','Technical review and evidence'),('Risk/scope review','Weekly','Whole team','Updated registers and mitigation owners')
],[1900,1600,2600,3260],font_size=8.4)
heading('11.1 Change process',2)
numbered_sequence(['Record the requested change, reason and requesting stakeholder.','Estimate value, effort, schedule effect, risk and impact on existing acceptance criteria.','Recommend accept, defer, substitute or reject.','Obtain approval from the project leader and client for material scope changes.','Update scope, backlog, schedule, risk register and decision log before implementation.'])
heading('11.2 Status reporting',2)
para('The weekly status summary will show completed work, next-week commitments, milestone confidence, top risks, decisions required and workload concerns. Status will be evidence-based using integrated software, issue history and acceptance results.')

heading('12. Deliverables and handover')
make_table(['Deliverable','Contents','Acceptance owner'],[
    ('D1 Project governance pack','Project plan, charter, requirements register, risks, decisions and meeting records.','Supervisor / client'),('D2 Source and database','Reviewed source, migrations, seed guidance and version history.','Technical representative'),('D3 CREMS MVP','Responsive application covering accepted in-scope workflows.','Client representative'),('D4 Test and acceptance pack','Test cases/results, security checks, UAT record and known issues.','Client + QA lead'),('D5 Deployment and operations pack','Environment settings, deployment, backup/restore and support notes.','Technical representative'),('D6 User documentation','Role-based quick guides and administrator instructions.','Client representative'),('D7 Final academic report/demo','Project outcomes, evaluation, reflection and demonstration.','Academic supervisor')
],[2200,5100,2060],font_size=8.5)
heading('12.1 Handover checklist',2)
for x in ['Final release is tagged and reproducible from documented prerequisites.','Production secrets are supplied through an approved private channel.','Database migration and backup/restore steps are demonstrated.','Administrator and representative operational users complete training.','Known limitations and deferred items are documented.','Client acceptance, outstanding actions and support ownership are recorded.']: bullet(x)

heading('13. Monitoring and control')
make_table(['Measure','Target','Review'],[
    ('Milestone completion','All seven milestones accepted or formally re-baselined.','Weekly'),('Iteration predictability','At least 80% of committed priority items completed.','Fortnightly'),('Build health','Main branch builds successfully.','Every integration'),('Defect position','Zero open critical/high defects at handover.','Weekly from W10'),('Review coverage','Functional changes receive peer review.','Every pull request'),('Client decisions','Open decision requests receive an owner and due date.','Weekly'),('Team health','Workload is visible and no critical knowledge is held by one member.','Weekly')
],[2600,4300,2460],font_size=8.6)
heading('13.1 Project closure criteria',2)
para('The project closes when accepted MVP scenarios pass, the agreed release and documents are handed over, critical/high defects are closed, known limitations are acknowledged, and the team completes academic submission and retrospective obligations.')

heading('14. References')
refs=[
    'Carpenters Fiji Pte Ltd. (2026). Car and Rental Equipment Management System project description and scope. Client-provided project brief.',
    'The University of the South Pacific. (2026). CS400 Industry Experience Project, Assignment 1 - Project Plan specification and marking rubric.',
    'Microsoft. (2026). ASP.NET Core documentation. https://learn.microsoft.com/aspnet/core/ (accessed 4 August 2026).',
    'Microsoft. (2026). Entity Framework Core documentation. https://learn.microsoft.com/ef/core/ (accessed 4 August 2026).',
    'React Team. (2026). React documentation. https://react.dev/ (accessed 4 August 2026).',
    'OWASP Foundation. (2026). Application Security Verification Standard. https://owasp.org/www-project-application-security-verification-standard/ (accessed 4 August 2026).'
]
for x in refs: para(x)
heading('15. Declaration of AI-assisted work')
para('Generative AI assistance was used to help structure this project-plan draft, improve wording and propose planning artefacts such as the initial schedule, risk register and acceptance measures. The student team is responsible for verifying every statement, correcting assumptions, replacing placeholders, aligning the plan with client decisions and disclosing this use in accordance with the assignment specification. AI-generated content is not treated as client approval, evidence or a substitute for the team’s own analysis.')

doc.add_page_break()
heading('Appendix A - Team charter')
heading('A.1 Purpose',2)
para('We will work as one accountable team to deliver a secure, usable and evidence-based CREMS MVP for Carpenters Fiji while meeting CS400 academic and professional expectations.')
heading('A.2 Team commitments',2)
for x in ['Attend scheduled meetings or notify the team early when unavailable.','Complete agreed work by the stated date and raise blockers within two working days.','Use Git branches, focused commits and pull requests; do not overwrite another member’s work.','Review peers constructively and discuss the work rather than the person.','Protect client information and use synthetic data in development.','Record material decisions, requirements and changes.','Share knowledge through pairing, demonstrations and documentation.','Report progress honestly and seek help before a deadline is endangered.']: bullet(x)
heading('A.3 Decision-making and conflict resolution',2)
numbered_sequence(['Seek consensus using evidence, acceptance criteria and client value.','If consensus is not reached promptly, the project leader decides routine delivery matters after hearing each view.','Scope, policy or client-impacting decisions are referred to the client/supervisor.','Address interpersonal concerns privately and respectfully first; document agreed actions.','Escalate unresolved conduct or workload issues to the academic supervisor.'])
heading('A.4 Working agreement',2)
make_table(['Area','Agreement'],[
    ('Primary channels','[INSERT TEAM CHAT] for coordination; GitHub for work/decisions; email for formal client communication.'),('Core response time','Acknowledge team messages within one working day.'),('Meetings','Twice-weekly team contact and fortnightly planning/review.'),('Code ownership','Collective ownership; authors obtain peer review before integration.'),('File naming','Clear, stable names; final submission controlled by project leader.'),('Definition of ready','Requirement has user value, acceptance criteria, dependencies and an owner.'),('Definition of done','Meets Section 5.3 and is integrated, tested, reviewed and documented.')
],[2100,7260],font_size=8.8)
heading('A.5 Expected contribution and accountability',2)
para('Each member plans approximately six project hours per week, adjusted transparently around assessment peaks. The team reviews workload weekly. Missed commitments are re-planned with an owner and date; repeated issues are discussed with the project leader and, if unresolved, the supervisor.')
heading('A.6 Charter acceptance',2)
make_table(['Team member','Student ID','Signature','Date'],[
    ('Kavish Chandra','[INSERT ID]','',''),('[INSERT TEAM MEMBER 2]','[INSERT ID]','',''),('[INSERT TEAM MEMBER 3]','[INSERT ID]','',''),('[INSERT TEAM MEMBER 4]','[INSERT ID]','','')
],[2800,1900,2960,1700],font_size=9)

# Final global table/heading safeguards
for table in doc.tables:
    for row in table.rows:
        row._tr.get_or_add_trPr()
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.06
                for r in p.runs:
                    if r.font.name is None: r.font.name='Calibri'
for p in doc.paragraphs:
    if p.style.name.startswith('Heading'): p.paragraph_format.keep_with_next=True

OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(OUT)
