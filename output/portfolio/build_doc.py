from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path
out=Path(r'C:\Github\CREMS\output\portfolio')
doc=Document()
s=doc.sections[0]
s.page_width=Inches(11.69); s.page_height=Inches(8.27)
s.top_margin=s.bottom_margin=Inches(.5)
s.left_margin=s.right_margin=Inches(.65)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
 st=doc.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
 st.font.size=Pt(11 if name=='Normal' else 10 if name=='Caption' else 22 if name=='Title' else 17 if name=='Heading 1' else 13)
 st.paragraph_format.space_after=Pt(7)
 st.paragraph_format.space_before=Pt(0)
doc.styles['Normal'].paragraph_format.line_spacing=1.08

def pic(path,caption,width=9.3):
 p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p.paragraph_format.space_after=Pt(4)
 p.add_run().add_picture(str(path),width=Inches(width))
 p.paragraph_format.keep_with_next=True
 doc.add_paragraph(caption,'Caption')

doc.add_paragraph('CREMS ePortfolio Artefacts','Title')
doc.add_paragraph('System architecture, application screenshots and a professional development roadmap supporting my CS400 Industry Experience Project reflection.')
doc.add_paragraph('Artefact 1  CREMS Logical System Architecture','Heading 1')
doc.add_paragraph('This diagram illustrates how the customer and staff interfaces connect to the backend API, business services, SQL Server database and supporting services. Working on CREMS strengthened my understanding of how frontend, backend and database components integrate to support business workflows.')
pic(Path(r'C:\Users\Rahul\OneDrive\Desktop\USP_YR_4_S2\CS400\Crems_System_Architecture.png'),'Figure 1. CREMS logical system architecture.',8.6)
base=Path(r'C:\Users\Rahul\AppData\Local\Temp')
items=[
('Artefact 2  Asset Management and Inspection Tracking','Asset register','4a57d83e-884c-42e3-b3a3-a5745ae4d795','The asset register brings together asset categories, branches, rental rates and operational statuses. Its search and filtering controls help staff locate assets and review which are available, on hire or undergoing maintenance.','Figure 2. Central asset register with status summaries and asset details.'),
('Artefact 2  Asset Management and Inspection Tracking','Inspection history','a82f887c-f93f-4ef7-9da1-fbad577c36e0','The inspection history groups recorded inspections for an individual vehicle. The screen shows pre-hire and post-hire entries with dates and staff details, illustrating how CREMS organises inspection records alongside the asset profile.','Figure 3. Inspection history for a selected rental vehicle.'),
('Artefact 3  Customer Rental and Booking Workflow','Customer rental catalogue','b3335152-f047-4ff8-8712-1ddef13a1a66','The customer interface presents rental options with locations and daily rates. Customers can select a pickup location and hire dates to search for available rentals. This screen illustrates how rental information is presented to support customer enquiries.','Figure 4. Customer catalogue and rental search interface.'),
('Artefact 3  Customer Rental and Booking Workflow','Staff booking requests','a3aa11a1-1ad5-489b-b52b-14ed97a493cf','The staff interface organises booking requests into enquiry, quotation, approval and confirmed-hire stages. Each request displays customer and asset information, hire dates, an estimate and a next action. Together with the customer catalogue, it illustrates the two sides of the rental workflow.','Figure 5. Staff workspace for reviewing requests and booking stages.'),
('Artefact 4  Customer Accounts and Staff Administration','Customer registration','561c47ad-c5f4-4fff-9854-f4c9041ed30c','The registration form collects customer contact details and account credentials. Required fields and password guidance are displayed within the form. This provides an example of the customer account creation interface.','Figure 6. Customer account registration form.'),
('Artefact 4  Customer Accounts and Staff Administration','Customer account profile','f38c79d0-cc02-42a9-b7f8-31712b6624d5','The customer account page brings together personal details, licence information, verification indicators and rental preferences. It illustrates how customers are provided with a dedicated interface for reviewing and updating their account information.','Figure 7. Customer profile with licence information and rental preferences.'),
('Artefact 4  Customer Accounts and Staff Administration','Customer account administration','ea49b580-dfb4-409d-abec-2ad4fa66b51c','The customer administration screen lists customer records with contact information, account status and verification indicators. Search and status filters help staff locate records. This complements the customer profile by showing the staff view of customer accounts.','Figure 8. Staff view of customer records and account statuses.'),
('Artefact 4  Customer Accounts and Staff Administration','Staff access and activity','34cd3585-ad62-4b86-af2d-d6739098d572','The staff directory displays assigned roles, division or branch scope and account status. Summary counts and filtering controls provide an overview of staff accounts. This screen illustrates how staff administration is organised separately from customer account management.','Figure 9. Staff directory showing roles, operating scope and account status.')]
for h,sub,key,desc,cap in items:
 doc.add_page_break();doc.add_paragraph(h,'Heading 1');doc.add_paragraph(sub,'Heading 2');doc.add_paragraph(desc)
 pic(base/('codex-clipboard-'+key+'.png'),cap)
doc.add_page_break(); doc.add_paragraph('Artefact 5  My Three Year Professional Development Roadmap','Heading 1')
doc.add_paragraph('This roadmap outlines my planned development during the first three years after graduation. It connects my university and CREMS experience with goals for industry employment, stronger technical skills and increasing professional responsibility.')
for h,p in [
('Year 1  Enter the software industry','I plan to obtain a software development role and strengthen my programming, database, Git, testing, debugging and professional communication skills. I will review my progress through completed work, testing evidence and feedback from my supervisor or colleagues.'),
('Year 2  Strengthen technical capabilities','I aim to take responsibility for larger features and gain practical experience in cloud deployment, DevOps, automated testing and system architecture. Evidence of progress could include a larger completed feature, a project with automated testing and deployment, and relevant professional training.'),
('Year 3  Develop greater responsibility','I aim to contribute to technical decisions, support colleagues and work towards SFIA Level 4 capabilities. I will assess my progress through documented design contributions, feedback on my responsibilities and examples of supporting other team members.'),
('Long term  Explore a software business opportunity','I hope to explore an independent rental and asset-management software service. My first steps would be customer research, development of a minimum viable product and gathering feedback from potential users.')]:
 doc.add_paragraph(h,'Heading 2');doc.add_paragraph(p)
doc.core_properties.title='CREMS ePortfolio Artefacts'
doc.core_properties.subject='System images and descriptions for a professional development ePortfolio'
doc.save(out/'CREMS_ePortfolio_Artefacts.docx')
print(out/'CREMS_ePortfolio_Artefacts.docx')
