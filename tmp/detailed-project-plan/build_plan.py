from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path

OUT = Path('/Users/kavishchandra/Documents/CS400/docs/CREMS-Detailed-Project-Plan.docx')
LOGO = Path('/Users/kavishchandra/Documents/CS400/src/crems-web/public/brand/carpenters-logo.png')

BLACK = '111111'; YELLOW = 'FFED00'; DARK_YELLOW = 'D8C900'; GRAY = '666666'
LIGHT = 'F5F5F1'; MID = 'DEDED7'; RED = '9B1C1C'; GREEN = '28643B'; WHITE = 'FFFFFF'

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.492)

def font(run, size=None, bold=None, color=BLACK, italic=None, name='Aptos'):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return run

styles = doc.styles
normal = styles['Normal']; normal.font.name='Aptos'; normal.font.size=Pt(11); normal.font.color.rgb=RGBColor.from_string(BLACK)
normal.paragraph_format.space_after=Pt(8); normal.paragraph_format.line_spacing=1.333
for name,size,before,after in [('Heading 1',16,18,10),('Heading 2',13,12,6),('Heading 3',12,8,4)]:
    s=styles[name]; s.font.name='Aptos Display'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(BLACK)
    s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True
if 'Table Text' not in styles:
    ts=styles.add_style('Table Text', WD_STYLE_TYPE.PARAGRAPH); ts.font.name='Aptos'; ts.font.size=Pt(9); ts.paragraph_format.space_after=Pt(2); ts.paragraph_format.line_spacing=1.12
if 'Small Note' not in styles:
    ns=styles.add_style('Small Note', WD_STYLE_TYPE.PARAGRAPH); ns.font.name='Aptos'; ns.font.size=Pt(9); ns.font.color.rgb=RGBColor.from_string(GRAY); ns.paragraph_format.space_after=Pt(5); ns.paragraph_format.line_spacing=1.1

def set_repeat_header(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)

def table(rows, widths, header=True, font_size=8.7):
    t=doc.add_table(rows=0, cols=len(widths)); t.alignment=WD_TABLE_ALIGNMENT.LEFT; t.autofit=False
    total=int(sum(widths)*1440)
    tblPr=t._tbl.tblPr
    for tag,val in [('tblW',str(total)),('tblInd','120')]:
        el=tblPr.find(qn('w:'+tag))
        if el is None: el=OxmlElement('w:'+tag); tblPr.append(el)
        el.set(qn('w:w'),val); el.set(qn('w:type'),'dxa')
    layout=OxmlElement('w:tblLayout'); layout.set(qn('w:type'),'fixed'); tblPr.append(layout)
    grid=t._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for w in widths:
        gc=OxmlElement('w:gridCol'); gc.set(qn('w:w'),str(int(w*1440))); grid.append(gc)
    for ri,rowdata in enumerate(rows):
        cells=t.add_row().cells
        for i,(cell,text,w) in enumerate(zip(cells,rowdata,widths)):
            cell.width=Inches(w); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=cell._tc.get_or_add_tcPr(); tcW=tcPr.find(qn('w:tcW')); tcW.set(qn('w:w'),str(int(w*1440))); tcW.set(qn('w:type'),'dxa')
            mar=OxmlElement('w:tcMar')
            for side,val in [('top',100),('start',120),('bottom',100),('end',120)]:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),str(val)); e.set(qn('w:type'),'dxa'); mar.append(e)
            tcPr.append(mar)
            p=cell.paragraphs[0]; p.style='Table Text'; p.paragraph_format.space_after=Pt(0)
            r=p.add_run(str(text)); font(r,font_size,bold=(header and ri==0),color=WHITE if header and ri==0 else BLACK)
            if header and ri==0: shade(cell,BLACK)
            elif ri%2==0: shade(cell,LIGHT)
        if header and ri==0: set_repeat_header(t.rows[-1])
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

def p(text='', size=11, bold=False, color=BLACK, align=None, after=8, before=0, style=None, italic=False):
    para=doc.add_paragraph(style=style); para.paragraph_format.space_before=Pt(before); para.paragraph_format.space_after=Pt(after); para.paragraph_format.line_spacing=1.333
    if align is not None: para.alignment=align
    font(para.add_run(text),size,bold,color,italic)
    return para

page_next = False
def h(text,level=1):
    global page_next
    para = doc.add_heading(text,level=level)
    if page_next:
        para.paragraph_format.page_break_before = True
        page_next = False
    return para
def bullet(text, level=0):
    para=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); para.paragraph_format.space_after=Pt(4); para.paragraph_format.line_spacing=1.208; para.add_run(text); return para
number_index = 0
def reset_number():
    global number_index
    number_index = 0
def number(text):
    global number_index
    number_index += 1
    para=doc.add_paragraph(); para.paragraph_format.left_indent=Inches(.25); para.paragraph_format.first_line_indent=Inches(-.25); para.paragraph_format.space_after=Pt(4); para.paragraph_format.line_spacing=1.208
    para.add_run(f'{number_index}.  {text}')
    return para
def page():
    pass
def callout(label,text,fill=YELLOW):
    t=doc.add_table(rows=1,cols=1); t.autofit=False; t.columns[0].width=Inches(6.5); shade(t.cell(0,0),fill)
    c=t.cell(0,0); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    pp=c.paragraphs[0]; pp.paragraph_format.space_after=Pt(2); font(pp.add_run(label+' '),10.5,True); font(pp.add_run(text),10.5)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def header_footer(section):
    hp=section.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.LEFT
    font(hp.add_run('CARPENTERS FIJI  |  CREMS PROJECT PLAN'),8.5,True,GRAY)
    fp=section.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    font(fp.add_run('CS400 Industry Experience Project  |  Semester 2, 2026  |  Confidential student project'),8,False,GRAY)

header_footer(sec)

# Cover
p('CARPENTERS FIJI', bold=True, size=12, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=5)
if LOGO.exists():
    par=doc.add_paragraph(); par.alignment=WD_ALIGN_PARAGRAPH.CENTER; par.add_run().add_picture(str(LOGO),width=Inches(1.05)); par.paragraph_format.space_after=Pt(12)
p('DETAILED PROJECT PLAN',10.5,True,YELLOW,WD_ALIGN_PARAGRAPH.CENTER,6)
p('Carpenters Rental, Equipment and Maintenance System',25,True,BLACK,WD_ALIGN_PARAGRAPH.CENTER,5)
p('CREMS',16,True,GRAY,WD_ALIGN_PARAGRAPH.CENTER,12)
p('A configurable group-wide platform for rental revenue, asset maintenance expenditure and customer service across Carpenters Fiji divisions',11.5,False,BLACK,WD_ALIGN_PARAGRAPH.CENTER,14)
table([
    ['Project detail','Information'],
    ['Client','Carpenters Fiji Pte Limited'],
    ['Course','CS400 Industry Experience Project - Semester 2, 2026'],
    ['Assignment','Assignment 1 - Project Plan and Team Charter (15%)'],
    ['Group number','[INSERT GROUP NUMBER]'],
    ['Academic supervisors','Amit Kumar; Shanil Naidu'],
    ['Version and date','Version 2.0 - 7 August 2026'],
], [1.8,4.7], font_size=9.5)
p('TEAM MEMBERS',9,True,GRAY,WD_ALIGN_PARAGRAPH.CENTER,3)
table([['Name','Student ID'],['Kavish Chandra','[INSERT STUDENT ID]'],['[INSERT TEAM MEMBER 2]','[INSERT STUDENT ID]'],['[INSERT TEAM MEMBER 3]','[INSERT STUDENT ID]'],['[INSERT TEAM MEMBER 4]','[INSERT STUDENT ID]']],[3.5,3.0],font_size=9.5)

page(); h('Document control',1)
table([['Version','Date','Prepared by','Purpose'],['1.0','4 Aug 2026','CREMS project team','Initial rental-focused plan'],['2.0','7 Aug 2026','CREMS project team','Group-wide modular scope, current requirements and detailed execution baseline']],[.75,1.05,1.8,2.9])
h('Approval record',2)
table([['Role','Name','Approval/signature','Date'],['Project leader','[INSERT NAME]','',''],['Client representative','[INSERT NAME]','',''],['Academic supervisor','Amit Kumar','','']],[1.55,1.65,2.3,1.0])
h('How to use this plan',2)
p('This plan is the delivery baseline for the four-person CREMS team. Requirements marked Confirmed may be developed and tested. Requirements marked Proposed or Open require client validation before they become contractual scope. Material changes follow the change-control process in Section 14.')
callout('Submission check:', 'Replace every [INSERT ...] field, obtain team charter signatures, update client-approval status, export the complete plan and charter as one PDF, and submit through the project leader account before the deadline.')
h('Contents',2)
for item in ['1 Executive summary','2 Domain research and client context','3 Problem analysis','4 Stakeholders and user groups','5 Objectives and success measures','6 Scope baseline','7 Requirements specification','8 Solution design and feasibility','9 Delivery approach and work breakdown','10 Timeline, milestones and dependencies','11 Resources, roles and responsibilities','12 Budget and cost controls','13 Risk management','14 Governance, communication and change control','15 Quality, testing, security and ethics','16 Deliverables, deployment and handover','17 Monitoring, reporting and closure','18 References','Appendix A Team charter','Appendix B Requirements traceability summary','Appendix C Client decisions required']:
    bullet(item)

page(); h('1. Executive summary',1)
p('Carpenters Fiji operates through multiple divisions, including Motors and vehicle rental, Carptrac, Hardware, Shipping and Logistics, Property, MH and other group businesses. Rental activity generates revenue; the same asset then requires inspection, preventive servicing and corrective maintenance that generate expenditure. Current and future divisions may rent different asset types, sell services, require quotes, or supply trained personnel with equipment. A single rigid vehicle-rental application would therefore not meet the client requirement for flexibility.')
p('CREMS will be delivered as a configurable, responsive web platform. Shared capabilities - customer accounts, assets, branches, requests, quotes, bookings, agreements, check-out, check-in, inspection, maintenance, invoicing, notifications and reporting - are reusable across divisions. Division, service and workflow configuration determines which capabilities apply. Individual customers receive a simple vehicle-first journey, while corporate customers may access multiple divisions, authorised contacts, job sites, purchase orders and approval rules.')
p('The selected stack is React and TypeScript for the frontend, ASP.NET Core 10 for secure APIs and identity, Entity Framework Core with Microsoft SQL Server for relational data, and Docker for consistent local development. The four-member team will use two-week iterations, continuous integration, peer review and client checkpoints over fourteen weeks from 3 August to 8 November 2026.')
callout('Delivery principle:', 'Protect a usable, testable MVP. Configurability and clear extension points are in scope; unconfirmed specialised workflows and external integrations are not silently treated as approved requirements.')
h('1.1 Project at a glance',2)
table([['Dimension','Baseline'],['Duration','14 weeks: 3 August to 8 November 2026'],['Team','Four student developers; approximately 6 hours per person per week'],['Method','Agile delivery in two-week iterations with fortnightly demonstrations'],['Primary outcome','Responsive group-wide rental and maintenance MVP'],['Initial operational focus','Vehicle and equipment rental revenue; inspections and maintenance expenditure'],['Access model','Administrator; division- and branch-scoped Branch Manager and Rental Officer; Individual and Corporate Customer'],['Cash budget ceiling','FJD 205, subject to approval; expected expenditure may remain zero']],[1.7,4.8],font_size=9.3)

page(); h('2. Domain research and client context',1)
h('2.1 Carpenters Fiji operating context',2)
p('Research and the first client meeting indicate that Carpenters Fiji is a group of operating divisions rather than one homogeneous rental business. Divisions share corporate identity and may share customers and locations, but their assets, pricing, safety checks, personnel requirements and commercial workflows differ. The system must therefore separate division data while reusing common capabilities.')
table([['Division / area','Observed or proposed activity','CREMS planning implication'],['Motors / Rentals','Vehicle hire, fleet, parts and service across several locations','Vehicle availability, licence checks, deposits, fuel, damage and branch check-out/check-in.'],['Carptrac','Heavy machinery, generators, parts and technical service; public rental status requires confirmation','Quote-led equipment workflow, hour meter, transport, trained operator and high-value maintenance controls.'],['Hardware','Tools, builders equipment and materials; formal hire status requires confirmation','Large categories/SKUs; possible tool-hire module remains configurable and off until confirmed.'],['Shipping & Logistics','Shipping and logistics services','Quote, job/dispatch, origin/destination and service-status workflow rather than simple daily asset rental.'],['Property','Residential or commercial leasing and property maintenance','Monthly lease periods, tenant documents, deposits, inspections and work orders.'],['MH / retail','Retail activity within the group','Represented in group structure; no rental module enabled without confirmed need.']],[1.15,2.05,3.3],font_size=8.25)
h('2.2 Rental and maintenance value cycle',2)
reset_number()
number('A customer identifies a need and selects a division or begins with the simple vehicle-rental path.')
number('CREMS records an enquiry, availability request or quote depending on service configuration.')
number('Staff validate customer eligibility, asset availability, pricing, documents, deposits and approvals.')
number('The approved reservation becomes an agreement and check-out record. Revenue is recognised through charges and invoice records.')
number('Check-in captures meter, fuel, condition, damage and additional charges.')
number('Inspection may return the asset to availability or create maintenance work and expense records.')
h('2.3 Research conclusions',2)
for x in ['A modular monolith is appropriate: one deployable application with strong module boundaries is manageable for the team and extensible for the client.','Division capabilities and workflow rules must be configuration, not scattered conditional code.','The customer journey requires progressive disclosure so ordinary vehicle customers are not overwhelmed by shipping, property or heavy-equipment choices.','Personnel-supported hire must distinguish none, optional and required personnel and later support competency verification.','Carptrac and Hardware rental availability remain client validation questions, not public claims.']: bullet(x)

page(); h('3. Problem analysis',1)
h('3.1 Problem statement',2)
p('Carpenters Fiji requires a dependable source of truth for rental assets, service requests, customers, agreements, asset condition and maintenance across divisions and branches. Fragmented paper, spreadsheet or isolated processes can cause double booking, incorrect availability, incomplete agreements, weak damage evidence, missed service intervals, untraceable expenditure and limited management visibility. A vehicle-only system would also create expensive redesign when another division is added.')
h('3.2 Information requirements',2)
table([['Information group','Required data'],['Organization','Division, capabilities, service offerings, branches, division-to-branch participation and operating status.'],['Identity and access','Staff role, division, branch, account state, customer account linkage, email verification and audit identity.'],['Customer','Individual or corporate profile, contacts, identification, eligibility, authorised representatives, job sites and rental history.'],['Asset/service','Asset ID, category, registration/serial, division, branch, status, rate, specifications, personnel requirement and availability.'],['Commercial transaction','Enquiry/quote, booking period, line items, pricing, tax, deposit, approval, agreement, invoice, payment and balance.'],['Operations','Check-out/check-in inspection, meter/fuel, photos/evidence, damage, personnel assignment, dispatch and asset transfer.'],['Maintenance','Job, fault, preventive schedule, supplier, labour/parts, estimated and actual expense, status and return-to-service.'],['Management','Utilisation, revenue, maintenance expense, overdue rentals, availability, customer activity and exceptions.']],[1.55,4.95],font_size=8.7)
h('3.3 Constraints and assumptions',2)
table([['Type','Constraint or assumption','Planning response'],['Time','One academic semester and fixed assessment dates','Prioritise demonstrable MVP; time-box discovery and protect scope.'],['People','Four students balancing other courses','Small work packages, pairing, peer review and visible workload.'],['Requirements','Specialised division rules remain incomplete','Requirements register, prototypes, client decisions and configurable defaults.'],['Privacy','Customer, identity and agreement data are sensitive','Synthetic development data, least privilege, secure secrets and retention decisions.'],['Safety','Vehicle/equipment condition and operator suitability affect people and property','Mandatory checks, audit history and no bypass of safety gates without approved authority.'],['Connectivity','Team uses macOS and Windows; demonstrations may be remote','Docker database, setup guide, repeatable builds and local fallback demonstration.'],['Technology','Solution must be realistic yet supportable by students','Modular .NET/React/SQL architecture with limited external dependencies.']],[1.0,2.25,3.25],font_size=8.4)

page(); h('4. Stakeholders and user groups',1)
table([['Stakeholder','Interest / responsibility','Engagement'],['Client sponsor / management','Business value, scope, priorities, acceptance and future rollout','Fortnightly checkpoint; milestone acceptance; change decisions.'],['Division representatives','Correct workflow, terminology, assets, pricing and reports','Focused discovery and prototype review by division.'],['Branch managers','Branch performance, approvals, staff and exception handling','Workflow validation, UAT and role-specific training.'],['Rental officers / agents','Fast daily requests, agreements, inspections and returns','Task walkthroughs, usability tests and demonstrations.'],['Customers','Simple discovery, trusted booking status, clear obligations and privacy','Customer-journey testing using synthetic personas.'],['Maintenance personnel','Accurate faults, schedules, work status, parts and cost','Maintenance process validation and acceptance scenario.'],['Carpenters IT / technical owner','Security, deployment, backup, supportability and integration','Architecture/deployment review and handover.'],['Academic supervisor','Assessment quality, professional conduct and progress','Scheduled updates, plan review and escalation.'],['Four-person student team','Delivery, learning outcomes, ethical conduct and evidence','Twice-weekly coordination, pull requests and retrospectives.']],[1.5,2.55,2.45],font_size=8.3)
h('4.1 User access principles',2)
for x in ['Administrators receive group-wide setup and security access.','Rental officers and branch managers are assigned both a division and a branch; operational records must match both.','A branch can participate in multiple divisions, but this does not grant a staff member cross-division access.','Customers use a separate portal and can view only records linked to their verified customer account.','Corporate authorised contacts and approval limits are an MVP-extension requirement subject to client confirmation.']: bullet(x)

page(); h('5. Objectives and success measures',1)
h('5.1 Objectives',2)
for x in ['Deliver one configurable platform that can add a division, branch or service without rebuilding core rental and maintenance functions.','Prevent conflicting confirmed reservations through authoritative availability and date validation.','Support traceable request-to-return workflows and convert inspection outcomes into availability or maintenance.','Separate group, division, branch, staff and customer access using server-enforced authorization.','Provide a simple customer journey for individuals and a richer cross-division path for corporate customers.','Make operational and financial information visible through accurate dashboards, reports and audit events.','Produce a reproducible release, deployment guide, user guidance, UAT evidence and academic report.']: bullet(x)
h('5.2 Measurable acceptance criteria',2)
table([['ID','Acceptance measure','Evidence'],['AC-01','All four user groups sign in and receive correct UI and API permissions; negative access tests pass.','Role/access test matrix'],['AC-02','Non-admin staff see only operational records matching assigned division and branch.','Cross-scope API tests'],['AC-03','System rejects an overlapping active/confirmed booking for the same asset and dates.','Automated test and demonstration'],['AC-04','Vehicle scenario completes request, approval, agreement, check-out, return, invoice and maintenance decision.','End-to-end UAT script'],['AC-05','Personnel-required equipment cannot proceed without the required assignment/verification once that workflow is enabled.','Equipment UAT scenario'],['AC-06','Dashboard totals reconcile with underlying transactions and division/branch filters.','Reconciliation checklist'],['AC-07','Critical status and administration changes record actor, time and material before/after detail.','Audit inspection'],['AC-08','Core customer and staff workflows pass responsive and plain-language usability checks.','Usability test record'],['AC-09','Zero open critical/high security or data-integrity defects at handover.','Issue and security registers'],['AC-10','All members can build and run the documented release on supported environments.','Setup verification checklist']],[.72,3.95,1.83],font_size=8.15)

page(); h('6. Scope baseline',1)
h('6.1 In scope - MVP',2)
for x in ['Configurable divisions, service offerings, capabilities and division-to-branch membership.','Staff identity, password recovery, administration, audit, role access and division/branch data scope.','Individual and corporate customer registration, email verification, secure invitation of existing records and account-linked booking history.','Vehicle and equipment asset register, QR identification, availability, rates, status and personnel requirement.','Public catalogue, date/branch/division search, enquiry/booking request and reference tracking.','Booking review, conflict checks, allocation, pricing, deposits, approvals and status transitions.','Rental agreement, authorised driver, check-out/check-in inspections, evidence, incidents, addenda, invoice and notifications.','Maintenance jobs, preventive service dates, estimated/actual expense and return-to-service status.','Corporate account foundation, quotes, purchase-order requirement, job sites and approval records.','Dashboards, operational reports, task-oriented manager view, email queue and staging-safe delivery controls.','Deployment/configuration guidance, Postman tests, user guides, UAT, training and handover.']: bullet(x)
h('6.2 Out of scope for the assessed MVP',2)
for x in ['Live ERP/general-ledger integration or automatic financial posting.','Production payment gateway and card-data processing.','Native iOS/Android applications.','Live GPS/telematics provider integration beyond an extensible data model.','Payroll, personnel rostering and formal competency/licensing system.','External marketplace or travel-booking-platform integration.','Full property-management accounting, shipping manifest/customs integration or retail point-of-sale.','Production migration of historical or personal data without an approved migration/privacy plan.']: bullet(x)
h('6.3 Future options',2)
p('The architecture will preserve REST boundaries and stable identifiers for approved future modules: payment gateway, ERP/accounting, telematics, document storage, mobile client, electronic signatures, advanced corporate credit, trained-personnel scheduling, shipping dispatch, property tenancy and business intelligence. These are extension options, not implicit MVP promises.')

page(); h('7. Requirements specification',1)
h('7.1 Functional requirements',2)
table([['ID','Requirement','Priority / status'],['FR-01','Administrator configures divisions, capabilities, services and active branch participation.','Must / Confirmed'],['FR-02','Administrator assigns each non-admin staff member one division and one branch.','Must / Confirmed'],['FR-03','API limits staff operational data by role, division and branch.','Must / Confirmed'],['FR-04','Customer registers as Individual or Business, verifies email, signs in and views linked history.','Must / Confirmed'],['FR-05','Staff securely invites an existing customer record to activate online access.','Must / Confirmed'],['FR-06','Customer searches available assets by dates, branch, type and division without unrelated complexity.','Must / Confirmed'],['FR-07','System creates reference-traceable requests and prevents conflicting allocations.','Must / Confirmed'],['FR-08','Staff reviews and confirms customer, asset, period, rate, tax, deposit and approval details.','Must / Confirmed'],['FR-09','Collection requires applicable agreement, identification, licence/payment and inspection evidence.','Must / Confirmed'],['FR-10','Return captures condition, meter/fuel, damage and charges; creates invoice and availability decision.','Must / Confirmed'],['FR-11','Maintenance records planned/corrective work, supplier, parts and estimated/actual expense.','Must / Confirmed'],['FR-12','Equipment/service configuration records none, optional or required trained personnel.','Must / Confirmed'],['FR-13','Required personnel assignment and competency evidence gate equipment dispatch.','Should / Proposed'],['FR-14','Corporate accounts record authorised contacts, job sites, PO requirement, credit terms and approvals.','Should / Partially confirmed'],['FR-15','Property service supports quote/lease request, monthly period, deposit and inspections.','Could / Proposed'],['FR-16','Shipping/logistics service supports enquiry, quote and dispatch status.','Could / Proposed'],['FR-17','Managers receive utilisation, revenue, maintenance expense, overdue and exception views.','Must / Confirmed'],['FR-18','Critical actions create auditable records and notifications are queued/retried.','Must / Confirmed']],[.7,4.55,1.25],font_size=7.8)
h('7.2 Non-functional requirements',2)
table([['ID','Requirement / target'],['NFR-01 Security','ASP.NET Identity hashing, secure cookies, HTTPS outside local development, least privilege, server-side policies and secrets outside Git.'],['NFR-02 Performance','Common list/search APIs should return within 2 seconds at agreed demonstration volume under normal local/staging conditions.'],['NFR-03 Availability','Staging demonstration can be restored from source, migrations, configuration and backup within 4 hours.'],['NFR-04 Usability','Core daily task is discoverable from task-based navigation; representative users complete UAT without developer intervention.'],['NFR-05 Accessibility','Keyboard operation, visible focus, readable contrast, labels and responsive layouts are checked for core pages.'],['NFR-06 Maintainability','Typed API contracts, modular domain boundaries, migrations, review and documented configuration.'],['NFR-07 Data integrity','Unique identifiers, relational constraints, status transitions, concurrency-aware booking checks and audit evidence.'],['NFR-08 Privacy','Minimum necessary data, synthetic non-production data, restricted access and client-approved retention/export/deletion rules.'],['NFR-09 Compatibility','Current Chrome/Edge/Safari desktop; responsive support for common tablet/mobile widths; development on macOS and Windows.'],['NFR-10 Observability','Health endpoint, structured application logs, email delivery status and actionable error responses.']],[1.25,5.25],font_size=8.25)

page(); h('7.3 Core business rules',2)
for x in ['BR-01: A confirmed or active booking blocks overlapping periods for the same asset.','BR-02: An asset in maintenance, out-of-service or retired status cannot be offered as available.','BR-03: A staff member cannot create or operate on an asset outside both assigned division and branch.','BR-04: A customer cannot claim an existing record by knowing an email/customer number; staff invitation or approved verification is required.','BR-05: Vehicle check-out requires applicable identification/licence/payment and signed agreement checks.','BR-06: Personnel-required services cannot be dispatched without the configured personnel control when enabled.','BR-07: Return condition determines whether the asset becomes available, enters inspection or generates maintenance.','BR-08: Original signed agreement snapshots remain immutable; approved changes use addenda.','BR-09: Financial discounts, waivers and high-value actions follow configured approval thresholds.','BR-10: Unconfirmed division services remain quote-only, non-public or disabled.']: bullet(x)
h('7.4 Open requirements requiring client decision',2)
table([['Decision','Why it matters','Needed by'],['Formal equipment owner: Carptrac, Hardware or both?','Determines catalogue, workflows, branches and accountable users.','End of W2'],['Which initial divisions and branches form the pilot?','Controls data collection, UAT scope and staff configuration.','End of W2'],['When is trained personnel mandatory and who verifies competency?','Defines safety gate, assignment data and legal responsibility.','Before equipment workflow build'],['Approved rate, deposit, tax, discount and cancellation rules?','Required for accurate pricing and agreement terms.','Before booking UAT'],['Customer/corporate approval and credit rules?','Determines authorised contacts, PO and manager approval flow.','Before corporate workflow build'],['Property and shipping MVP expectation?','Determines whether prototypes remain configuration-only or become deliverables.','Feature baseline review'],['Email sender, retention and direct-delivery approval?','Controls staging/production privacy and notifications.','Before UAT'],['Deployment owner and target infrastructure?','Controls release architecture, backup and handover.','Before W11']],[1.45,3.65,1.4],font_size=8.2)

page(); h('8. Solution design and feasibility',1)
h('8.1 Alternatives considered',2)
table([['Alternative','Strength','Assessment'],['ASP.NET Core + React + SQL Server','Enterprise identity/authorization, strong typing, migrations, client-stack alignment and current implementation investment.','Selected: strongest balance of realism, security, scalability and semester feasibility.'],['Laravel + React + MySQL/SQL Server','Rapid CRUD and approachable PHP ecosystem.','Viable, but rewriting completed identity, API and workflow work introduces schedule and defect risk.'],['Java Spring + React','Mature enterprise ecosystem and strong typing.','Technically sound but greater setup/learning overhead with no compensating client benefit.'],['No-code / rental SaaS','Fast initial configuration.','Insufficient control over division rules, audit, data scope, extensibility and academic solution-design outcomes.'],['Separate system per division','Simple local rules at first.','Rejected: duplicated customers, code and maintenance; poor group reporting and expensive future change.']],[1.45,2.2,2.85],font_size=8.2)
h('8.2 Selected architecture',2)
p('CREMS uses a modular monolith: one frontend, one API and one relational database deployment, with explicit domain modules for identity, organisation, customers, catalogue/assets, booking/rental, maintenance, corporate operations, notifications, administration and reporting. This avoids premature distributed-system complexity while allowing future modules or integrations to use stable APIs.')
table([['Layer','Technology','Responsibility'],['Experience','React 19, TypeScript, Vite, Material UI','Responsive public site, customer portal and task-based staff portal.'],['Application/API','ASP.NET Core 10 Web API','Validation, authorization, workflow orchestration and REST contracts.'],['Identity','ASP.NET Core Identity, secure cookies','Staff/customer separation, roles, recovery, verification and session security.'],['Data','SQL Server 2022, Entity Framework Core','Relational integrity, transactions, migrations, scope queries and reporting.'],['Delivery','Docker Compose, GitHub, Postman, VS Code','Repeatable environments, collaboration, testing and release evidence.']],[1.3,2.2,3.0],font_size=8.5)
h('8.3 Feasibility assessment',2)
table([['Context','Assessment and control','Conclusion'],['Technical','Core stack already builds; identity, division scope, customer portal, rental lifecycle, QR, maintenance and administration foundations exist. Remaining work is integration, configurable workflows and test depth.','Feasible'],['Implementation','Four members, modular work packages, peer review and protected MVP boundary fit the semester.','Feasible with scope control'],['Economic','Open-source/student tools and existing laptops; maximum optional cash spend FJD 205.','Feasible'],['Aesthetic/usability','Carpenters branding, task-based navigation, responsive UI and user help reduce training burden.','Feasible with UAT'],['Ethical/privacy','Synthetic data, minimum collection, role scope, audit and client-approved retention reduce exposure.','Feasible with governance'],['Health and safety','Mandatory condition checks and proposed operator competency gate support safe hire; client rules still required.','Conditional on validation'],['Societal/cultural','Simple language, individual/corporate journeys and branch-aware service support Fiji-wide users; accessibility and connectivity require testing.','Feasible with inclusive testing'],['Environmental','Better utilisation/service scheduling can extend asset life; hosting and data retention should remain proportionate.','Positive potential']],[1.15,4.35,1.0],font_size=7.95)
h('8.4 Creative and innovative contribution',2)
for x in ['Capability-driven divisions allow the same module to support different businesses without copying the application.','Progressive disclosure presents a simple vehicle journey while retaining cross-division services for corporate users.','One QR identity supports asset lookup, check-out and check-in with server-enforced scope.','Rental revenue and maintenance expense are connected through one auditable asset lifecycle.','Quote-only configuration permits research/prototyping without falsely publishing unconfirmed rental services.']: bullet(x)

page(); h('9. Delivery approach and work breakdown',1)
h('9.1 Agile operating model',2)
p('The team works in two-week iterations. Each iteration starts with refinement and acceptance criteria, uses small feature branches and peer-reviewed pull requests, and ends with an integrated demonstration. Client decisions are recorded in the requirements register. Features are complete only when authorization, validation, tests, documentation and demonstration data are integrated.')
h('9.2 Work packages',2)
table([['WP','Work package','Outputs / completion evidence'],['WP1','Discovery, research and governance','Validated scope, project plan, charter, requirements/decision/risk registers and prioritised backlog.'],['WP2','Architecture and environments','Repository workflow, Docker database, builds, migrations, module boundaries and setup verification.'],['WP3','Organisation, identity and access','Divisions/branches, staff/customer accounts, recovery/verification, scoped authorization and audit.'],['WP4','Catalogue, assets and customers','Assets/services, branch/division ownership, customer profiles, corporate foundation and QR identity.'],['WP5','Requests, quotes and bookings','Public discovery, enquiry/reference tracking, availability, allocation, pricing and approvals.'],['WP6','Agreement and rental operations','Agreement snapshots, driver/customer checks, check-out, active rental, addenda, check-in and invoice.'],['WP7','Maintenance and personnel-supported hire','Inspection decision, jobs/cost, preventive schedule and proposed operator assignment control.'],['WP8','Management information and notifications','Dashboard, reports, email queue, exception tasks and reconciliation.'],['WP9','Quality, deployment and handover','Automated/manual tests, security/usability checks, UAT, release, training and final report.']],[.55,2.0,3.95],font_size=8.15)
h('9.3 Definition of done',2)
for x in ['Requirement and acceptance criteria are linked to the work item.','Backend and frontend builds pass and migrations are reviewed.','Authorization is enforced in APIs and tested with allowed and denied users.','Peer review is complete and no unresolved high-severity comment remains.','Repeatable tests and demonstration data pass.','User help, API collection and relevant documentation are updated.','Client/product-owner acceptance is recorded or follow-up work is scheduled.']: bullet(x)

page(); h('10. Timeline, milestones and dependencies',1)
p('The schedule begins 3 August 2026 and ends 8 November 2026. Activities overlap intentionally: discovery continues through demonstrations, documentation is updated with each feature, and testing begins before feature complete. The Assignment 1 plan and charter baseline is due 16 August 2026.')
h('10.1 Detailed timeline',2)
table([['Weeks / dates','Planned focus','Exit deliverable'],['W1-2 | 3-16 Aug','Client research; group/division problem analysis; plan and charter; requirements baseline; environment verification.','M1: submitted plan; prioritised backlog; decisions list.'],['W3 | 17-23 Aug','Validate pilot divisions/branches; refine data model; assign existing staff division/branch; API contract review.','Validated organisational and access baseline.'],['W4 | 24-30 Aug','Complete customer verification/invitation tests; division/service administration; usability review.','M2: secure account and configuration demonstration.'],['W5-6 | 31 Aug-13 Sep','Asset/service catalogue; personnel requirements; corporate profile/contacts; public progressive-disclosure journey.','M3: core records and customer journey accepted.'],['W7-8 | 14-27 Sep','Quote/request, availability, pricing, approvals, bookings and conflict testing by division.','M4: request-to-confirmation demonstration.'],['W9 | 28 Sep-4 Oct','Agreement templates and configurable collection gates; equipment/personnel proof-of-concept.','Signed agreement and safe check-out demonstration.'],['W10 | 5-11 Oct','Return inspection, charges, invoice, maintenance decision and expense reporting.','M5: complete rental lifecycle.'],['W11 | 12-18 Oct','Dashboards, revenue/expense reporting, notifications, audit reconciliation and accessibility improvements.','Feature-complete integrated build.'],['W12 | 19-25 Oct','Security, authorization, performance, compatibility, backup/restore and regression testing.','Release candidate 1 and evidence pack.'],['W13 | 26 Oct-1 Nov','Client UAT, defect correction, user training and deployment rehearsal.','M6: UAT accepted or exceptions recorded.'],['W14 | 2-8 Nov','Final release, source, database scripts, documentation, academic demonstration and handover.','M7: handover package and closure report.']],[1.25,3.35,1.9],font_size=7.9)
h('10.2 Milestones',2)
table([['Milestone','Target','Acceptance condition'],['M1 Plan baseline','16 Aug','Plan, charter, scope, requirements, schedule, budget and risks reviewed for submission.'],['M2 Secure foundation','30 Aug','All members run stack; roles and division/branch negative access tests pass.'],['M3 Core records','13 Sep','Divisions, branches, services, assets and customers are maintainable with validation.'],['M4 Booking workflow','27 Sep','Request-to-confirmation works and conflicting allocation is rejected.'],['M5 Rental lifecycle','11 Oct','Agreement, check-out, return, invoice and maintenance decision succeed.'],['M6 UAT complete','1 Nov','Agreed scenarios pass or exceptions are approved and scheduled.'],['M7 Handover','8 Nov','Release, source, migrations, test evidence, guides and final report delivered.']],[1.65,1.0,3.85],font_size=8.25)

page(); h('10.3 Dependency and critical-path view',2)
table([['Activity','Depends on','Critical consequence if late'],['Scope/decision baseline','Client access and domain research','Rework across model, UI and acceptance tests.'],['Division/branch configuration','Confirmed pilot structure','Staff cannot be assigned or tested correctly.'],['Availability/booking rules','Asset and service data','Rental lifecycle cannot be validated.'],['Agreement/check-out','Confirmed booking and client terms','No legally/operationally credible handover.'],['Return/maintenance','Active rental and inspection rules','No end-to-end revenue/expense lifecycle.'],['UAT','Integrated release candidate and scenarios','Handover acceptance delayed.'],['Deployment/handover','UAT resolution and technical owner','Final delivery incomplete.']],[1.55,2.0,2.95],font_size=8.4)
h('10.4 Schedule control',2)
for x in ['Track each work item against milestone and requirement ID.','Review committed versus completed priority work every fortnight; target at least 80%.','Escalate blockers after two working days.','Use scope substitution before extending dates: remove equal/greater effort when adding a material requirement.','Re-baseline only with recorded team, client and supervisor awareness.']: bullet(x)

page(); h('11. Resources, roles and responsibilities',1)
h('11.1 Human resources',2)
table([['Project role','Assigned person','Responsibilities','Effort'],['Project leader / client liaison','[INSERT NAME]','Plan, meetings, scope/decision control, integration, submission and team coordination.','6 h/week'],['Backend and data lead','[INSERT NAME]','Domain/API/database, authorization, migrations, backend tests and security.','6 h/week'],['Frontend and UX lead','[INSERT NAME]','Responsive journeys, accessibility, API integration, UI testing and user help.','6 h/week'],['Quality and DevOps lead','[INSERT NAME]','Test plan/evidence, Docker/builds, release, backup, documentation and demonstrations.','6 h/week']],[1.45,1.35,2.85,.85],font_size=8.1)
p('Roles establish accountability but not silos. Members pair on high-risk features, review each other’s changes and rotate meeting-note, demonstration and support duties. Planned team effort is 336 hours (4 people x 6 hours x 14 weeks).')
h('11.2 Technical and information resources',2)
for x in ['Four development laptops (one macOS, two confirmed Windows, one team device to be confirmed).','GitHub repository with issues, feature branches, pull requests and version history.','.NET 10 SDK, Node.js 22+, React, Material UI, SQL Server 2022, Docker Desktop, VS Code and Postman.','Client representatives, current forms/agreements, sample asset categories, branch/division information and acceptance decisions.','Academic supervisor for assessment guidance and escalation.','Synthetic test dataset; production customer data is not copied into student environments.','Optional staging host, email sender/domain and secure configuration channel.']: bullet(x)
h('11.3 RACI summary',2)
p('Legend: A = accountable; R = responsible; C = consulted; I = informed.',style='Small Note')
table([['Activity','Leader','Backend','Frontend','QA/DevOps','Client'],['Requirements/scope','A','C','C','C','R/C'],['Architecture/data','C','A/R','C','C','I'],['Customer/staff UX','C','C','A/R','C','C'],['Testing/release','C','R','R','A/R','C'],['Acceptance/change approval','R','C','C','C','A']],[2.15,.7,.9,.9,1.05,.8],font_size=7.4)

page(); h('12. Budget and cost controls',1)
p('This is an academic project. Student labour, personal laptops and existing connectivity are in-kind contributions and are not charged to Carpenters Fiji. Their quantities are shown so the budget remains comprehensive. Any cash purchase requires written team/client approval and evidence.')
table([['Item','Basis','Cash FJD','Treatment / justification'],['Student effort','4 x 6 h/week x 14 weeks = 336 hours','0','In-kind academic contribution; track hours for workload control.'],['Development laptops','Four existing personal devices','0','No new hardware purchase required.'],['Software/tooling','.NET, React, SQL Server Developer, VS Code, GitHub, Postman, Docker','0','Open-source, developer or existing student access.'],['Local database','SQL Server container and local storage','0','Development/test only.'],['Remote staging','Optional short-lived deployment','120','Only if remote client UAT is required and approved.'],['Email/domain allowance','Optional test sender/domain','50','Supports realistic OTP, notification and invoice testing.'],['Contingency','Approximately 20% of potential service cost','35','Minor approved hosting/domain variation.'],['Maximum cash budget','','205','Expected actual cash cost may remain zero.']],[1.55,2.2,.7,2.05],font_size=8.0)
h('12.1 Budget controls',2)
for x in ['Project leader records approval, receipt, purpose and remaining budget.','No real payment/card service is enabled under this budget.','Free tiers must still satisfy privacy, reliability and licensing requirements.','Contingency does not authorise scope expansion.','Budget status is reviewed at each milestone and final actual cost is reported at closure.']: bullet(x)
h('12.2 Economic value rationale',2)
p('The plan avoids parallel division systems and repeated customer/asset logic. A shared configurable platform can reduce duplicate entry, booking conflicts, unavailable-asset promises, missed service, manual follow-up and reporting effort. Quantified business benefits require client baseline data and are therefore a future business-case activity rather than unsupported savings claims in this academic plan.')

page(); h('13. Risk management',1)
p('Probability (P) and impact (I) use a 1-5 scale; exposure is P x I. High = 15-25, Medium = 8-14, Low = 1-7. Risks are reviewed weekly and after each client decision.')
table([['ID','Risk','P','I','Score','Mitigation / owner'],['R1','Requirements remain unclear or change late.','4','5','20 H','Time-box discovery, decision log, prototypes and change control / Leader.'],['R2','Group-wide scope exceeds semester capacity.','4','5','20 H','Protect MVP, capability configuration, MoSCoW priorities / All.'],['R3','Division/branch authorization exposes data.','3','5','15 H','Server scope queries, negative tests and peer review / Backend + QA.'],['R4','Booking/status rules allow conflicting or invalid rentals.','3','5','15 H','Transactions, explicit transitions and overlap tests / Backend.'],['R5','Equipment/operator safety rule is incomplete.','3','5','15 H','Client validation, mandatory gate and no public enablement until accepted / Leader.'],['R6','Sensitive data or credentials are exposed.','2','5','10 M','Synthetic data, user secrets, review and secret scanning / All.'],['R7','One member becomes unavailable.','3','4','12 M','Pairing, shared documentation, small reviews and cross-training / Leader.'],['R8','Integration occurs too late.','3','4','12 M','Continuous integration and fortnightly end-to-end demo / QA.'],['R9','Environment differences block a member.','3','3','9 M','Docker, Mac/Windows guide and setup verification / QA.'],['R10','Client availability delays decisions/UAT.','3','4','12 M','Book checkpoints early; focused questions with default/defer date / Liaison.'],['R11','Email/OTP misconfiguration sends test data incorrectly.','2','4','8 M','Redirect-all staging, delivery log, explicit production approval / Backend.'],['R12','Data loss or repository damage.','2','5','10 M','GitHub remote, protected review, backup/restore rehearsal / QA.'],['R13','UI remains too complex for frontline staff.','3','4','12 M','Task language, progressive disclosure, help and representative UAT / Frontend.'],['R14','Report totals do not reconcile.','3','4','12 M','Defined formulas and source-record reconciliation tests / Backend + QA.'],['R15','Deployment owner/infrastructure remains unknown.','3','4','12 M','Decision by W11; portable documented fallback / Leader.']],[.45,2.4,.35,.35,.55,2.4],font_size=6.8)
h('13.1 Escalation thresholds',2)
for x in ['High risks receive an owner, dated action and immediate leader review.','Any change affecting scope, milestone, privacy, safety, budget or architecture is escalated to client/supervisor.','A blocker longer than two working days is raised at the next team contact.','A suspected privacy/security incident stops affected demonstration/data use until assessed.']: bullet(x)

page(); h('14. Governance, communication and change control',1)
h('14.1 Communication plan',2)
table([['Forum','Frequency','Participants','Output'],['Team coordination','Twice weekly','Four students','Progress, next actions, blockers and workload.'],['Iteration planning/review','Every two weeks','Team; client when available','Goal, demonstration, feedback and accepted work.'],['Client checkpoint','Fortnightly or agreed','Liaison, relevant members, client','Decisions, clarified rules and acceptance notes.'],['Supervisor update','As scheduled','Team and academic supervisor','Academic guidance and escalations.'],['Pull-request review','Every functional change','Author plus reviewer','Technical/quality approval and evidence.'],['Risk/scope review','Weekly','Whole team','Updated risks, decisions and mitigation.']],[1.5,1.2,1.75,2.05],font_size=8.15)
h('14.2 Change-control process',2)
reset_number()
number('Record the requested change, reason, requester and affected requirement.')
number('Estimate value, effort, schedule, cost, security, data, safety and regression impact.')
number('Recommend accept, defer, substitute or reject; identify what leaves the MVP if needed.')
number('Obtain project leader and client approval for material scope; notify supervisor where assessment/timeline is affected.')
number('Update scope, requirements, backlog, schedule, risks, tests and decision log before implementation.')
h('14.3 Decision rights',2)
for x in ['Routine implementation within approved acceptance criteria: responsible feature pair.','Architecture/data changes: backend lead with peer review; leader resolves unresolved delivery trade-off.','User journey and wording: frontend lead with representative-user/client input.','Scope, policy, safety, privacy, budget or external commitment: client and project leader; supervisor informed as appropriate.']: bullet(x)

page(); h('15. Quality, testing, security and ethics',1)
h('15.1 Quality plan',2)
table([['Control','Minimum evidence'],['Build quality','Backend build and frontend TypeScript production build pass.'],['Review','At least one peer reviews functional changes.'],['API testing','Positive, validation, unauthenticated, unauthorised, cross-division and cross-branch cases.'],['Data quality','Required fields, uniqueness, relationships, transitions, overlaps and migration review.'],['Workflow testing','Vehicle end-to-end; equipment/personnel scenario if enabled; maintenance expense path.'],['Usability/accessibility','Task completion, responsive widths, keyboard, labels, contrast and plain language.'],['Security','Authentication, session, recovery/OTP, least privilege, secrets and negative authorization.'],['Release','Regression, backup/restore, UAT, known issues and deployment rehearsal.']],[1.5,5.0],font_size=8.45)
h('15.2 Test levels and ownership',2)
for x in ['Unit/domain tests: booking overlap, transitions, pricing calculations and code verification.','API integration tests/Postman: endpoints, validation, cookies and role/scope matrix.','Database/migration tests: clean creation, upgrade, uniqueness, referential integrity and seed repeatability.','UI tests: primary workflows, error recovery, empty/loading states and responsive layouts.','Security tests: denied-role and cross-division/branch access, account activation abuse and sensitive-data exposure.','UAT: client-approved scenarios performed with synthetic representative data.']: bullet(x)
h('15.3 Security and privacy controls',2)
for x in ['Passwords are handled by ASP.NET Identity hashing; raw passwords are never stored or logged.','Secrets and bootstrap credentials remain in local/environment secret stores, not Git.','Staging email defaults to redirect-all until direct delivery is explicitly approved.','Customer activation and verification codes are expiring, one-time and stored as salted hashes.','Critical administration, rental and asset changes are auditable.','Only minimum required information is collected; retention/export/deletion rules require client approval.']: bullet(x)
h('15.4 Professional, ethical and safety conduct',2)
p('The team will protect client confidentiality, distinguish confirmed facts from assumptions, respect software licences, reference external material, report limitations honestly and disclose AI assistance. Safety-sensitive condition and personnel records will not be bypassed or altered without traceability. Decisions balance the client’s interests with customer privacy, worker/public safety, accessibility and the wider public interest.')

page(); h('16. Deliverables, deployment and handover',1)
table([['ID','Deliverable','Contents','Acceptance owner'],['D1','Governance pack','Plan, charter, requirements, risk/decision/change registers and meeting records.','Supervisor / client'],['D2','Source and database','Reviewed source, migrations, seed guidance, API collection and version history.','Technical representative'],['D3','CREMS MVP','Responsive application covering accepted modular rental/maintenance workflows.','Client representative'],['D4','Test and acceptance pack','Cases/results, authorization/security checks, UAT, defects and known issues.','Client + QA lead'],['D5','Deployment/operations pack','Configuration, deployment, email, backup/restore, monitoring and support notes.','Technical representative'],['D6','User documentation','Task-based staff guides, customer guidance and administrator instructions.','Client representative'],['D7','Academic report/demo','Outcomes, evidence, evaluation, reflection and demonstration.','Academic supervisor']],[.45,1.25,3.4,1.4],font_size=7.8)
h('16.1 Deployment strategy',2)
for x in ['Development: independent local environments using Docker SQL Server and private secrets.','Staging: client-safe synthetic data, redirect-all email, HTTPS and release candidate configuration.','Production recommendation: approved host, managed secrets, TLS, backup, logging, restricted administrative access and documented owner.','Database changes are applied through versioned migrations with backup and rollback decision points.']: bullet(x)
h('16.2 Handover checklist',2)
for x in ['Release is tagged and reproducible from documented prerequisites.','Production/staging secrets are supplied through an approved private channel.','Migration and backup/restore procedures are demonstrated.','Administrator, manager and rental-officer representatives complete training.','Known limitations, deferred decisions and future options are documented.','Acceptance, outstanding actions, support owner and warranty/support period are recorded.']: bullet(x)

page(); h('17. Monitoring, reporting and closure',1)
table([['Measure','Target','Review'],['Milestone completion','All seven accepted or formally re-baselined.','Weekly'],['Iteration predictability','At least 80% of committed priority items completed.','Fortnightly'],['Build health','Main/integration branch builds successfully.','Every integration'],['Defect position','Zero open critical/high defects at handover.','Weekly from W10'],['Review coverage','Functional changes receive peer review.','Every pull request'],['Client decisions','Open requests have owner and due date.','Weekly'],['Team health','Visible workload; no critical knowledge held by one member.','Weekly'],['Requirements stability','Confirmed/proposed/open status maintained; no unrecorded scope.','Fortnightly']],[1.6,3.65,1.25],font_size=8.35)
h('17.1 Weekly status content',2)
for x in ['Completed work and evidence links.','Next-week commitments and owners.','Milestone confidence: Green, Amber or Red with reason.','Top risks, defects and blockers.','Client/supervisor decisions required and due dates.','Budget/effort variance and team workload concerns.']: bullet(x)
h('17.2 Closure criteria',2)
p('The project closes when accepted MVP scenarios pass; agreed release and documents are handed over; critical/high defects are closed; known limitations are acknowledged; deployment/backup ownership is recorded; the client acceptance position is documented; and the team completes academic submission, demonstration, retrospective and AI-use disclosure.')

page(); h('18. References',1)
refs=[
('Carpenters Fiji Pte Ltd. (2026).','Client meeting notes and Car and Rental Equipment Management System project information. Confidential client-provided material.'),
('Carpenters Fiji. (2026).','Corporate and division websites: https://www.carpenters.com.fj/ (accessed August 2026).'),
('Carpenters Motors. (2026).','Carpenters Motors and rental information: https://carpmotors.com.fj/ (accessed August 2026).'),
('The University of the South Pacific. (2026).','CS400 Industry Experience Project, Assignment 1 - Project Plan specification and marking rubric.'),
('Microsoft. (2026).','ASP.NET Core documentation: https://learn.microsoft.com/aspnet/core/ (accessed August 2026).'),
('Microsoft. (2026).','Entity Framework Core documentation: https://learn.microsoft.com/ef/core/ (accessed August 2026).'),
('React Team. (2026).','React documentation: https://react.dev/ (accessed August 2026).'),
('OWASP Foundation. (2026).','OWASP Application Security Verification Standard: https://owasp.org/www-project-application-security-verification-standard/ (accessed August 2026).'),
]
for a,b in refs:
    pp=doc.add_paragraph(); pp.paragraph_format.left_indent=Inches(.3); pp.paragraph_format.first_line_indent=Inches(-.3); pp.paragraph_format.space_after=Pt(7); font(pp.add_run(a+' '),10,True); font(pp.add_run(b),10)
h('18.1 AI-assisted work declaration',2)
p('Generative AI assistance was used to help structure this draft, improve wording, analyse planning coverage and generate initial planning artefacts. The student team remains responsible for verifying each statement, rewriting where appropriate, confirming sources and client decisions, replacing placeholders, maintaining confidentiality and documenting AI use in accordance with the assignment specification. AI output is not evidence of client approval and does not replace the team’s domain research, professional judgement or authorship.')

page(); h('Appendix A - Team charter',1)
h('A.1 Purpose',2)
p('We will work as one accountable team to deliver a secure, usable and evidence-based CREMS MVP for Carpenters Fiji while meeting CS400 academic and professional expectations.')
h('A.2 Team commitments',2)
for x in ['Attend agreed meetings or notify the team early when unavailable.','Complete agreed work by the stated date and raise blockers within two working days.','Use feature branches, focused commits and peer-reviewed pull requests.','Review peers constructively and discuss the work rather than the person.','Protect client information and use synthetic data in development.','Record material requirements, decisions, changes, risks and evidence.','Share knowledge through pairing, demonstrations and documentation.','Report progress and limitations honestly and seek help early.']: bullet(x)
h('A.3 Working agreement',2)
table([['Area','Agreement'],['Primary channels','[INSERT TEAM CHAT] for coordination; GitHub for work/decisions; email for formal client communication.'],['Response time','Acknowledge team messages within one working day.'],['Meetings','Twice-weekly team contact and fortnightly planning/review.'],['Code ownership','Collective ownership; authors obtain peer review before integration.'],['Definition of ready','User value, acceptance criteria, dependencies, priority and owner are clear.'],['Definition of done','Meets Section 9.3 and is integrated, tested, reviewed and documented.'],['File/submission control','Stable filenames; project leader performs final completeness and Moodle submission check.']],[1.65,4.85],font_size=8.6)
h('A.4 Decision-making and conflict resolution',2)
reset_number()
number('Seek consensus using evidence, acceptance criteria, client value and professional obligations.')
number('If routine consensus is not reached promptly, the project leader decides after hearing each view.')
number('Refer scope, policy, safety, privacy or client-impacting decisions to the client/supervisor.')
number('Address interpersonal concerns privately and respectfully; record agreed actions.')
number('Escalate unresolved conduct, integrity or persistent workload issues to the academic supervisor.')
h('A.5 Contribution and accountability',2)
p('Each member plans approximately six project hours per week, adjusted transparently around assessment peaks. Workload is reviewed weekly. Missed commitments are re-planned with an owner and date; repeated issues are discussed with the leader and, if unresolved, the supervisor.')
h('A.6 Charter acceptance',2)
table([['Team member','Student ID','Signature','Date'],['Kavish Chandra','[INSERT ID]','',''],['[INSERT TEAM MEMBER 2]','[INSERT ID]','',''],['[INSERT TEAM MEMBER 3]','[INSERT ID]','',''],['[INSERT TEAM MEMBER 4]','[INSERT ID]','','']],[2.3,1.25,1.8,1.15],font_size=8.8)

page(); h('Appendix B - Requirements traceability summary',1)
table([['Requirement group','Design / module','Planned evidence'],['Organisation and scoped access','Division, ServiceOffering, BranchDivision, ApplicationUser scope; API policies/queries','Admin configuration demo and cross-scope negative tests.'],['Customer identity','Customer portal, registration, verification, staff invitation and account-linked records','Individual/business registration and existing-customer activation tests.'],['Asset/service catalogue','Asset, division/service links, status, personnel requirement and QR','Catalogue/search, CRUD validation and QR scope tests.'],['Booking and rental','Booking, item, agreement, inspections, drivers, payments, invoice and lifecycle','Vehicle end-to-end UAT and booking-conflict tests.'],['Maintenance','MaintenanceJob, status/service dates, expense and dashboard','Return-to-maintenance scenario and cost reconciliation.'],['Corporate operations','CorporateAccount, quote, approval, job site, PO and dispatch foundation','Corporate scenario subject to approved MVP rules.'],['Management and audit','Dashboard, reports, management tasks, audit and notification delivery','Reconciliation, audit inspection and email delivery evidence.']],[1.55,3.15,1.8],font_size=8.15)

page(); h('Appendix C - Client decisions required',1)
p('Use this page as a focused agenda for the next requirements meeting. Record the decision, owner and date in the project decision log.')
table([['#','Question / decision','Current planning default','Owner / due'],['1','Which divisions and branches are in the first pilot?','Motors plus one confirmed equipment division; two or more operational branches.','Client / W2'],['2','Does Carptrac, Hardware or both provide formal equipment hire?','Keep public hire disabled or quote-only until confirmed.','Client / W2'],['3','Which equipment requires trained personnel and what evidence is mandatory?','Required flag for excavators/backhoes; dispatch gate proposed.','Division representative / before W7'],['4','What pricing, deposit, cancellation, tax and approval rules apply by division?','Configurable rules; no unsupported production values.','Finance/operations / before W7'],['5','What is the corporate authorised-contact and purchase-order process?','One account foundation; detailed approval workflow proposed.','Corporate operations / before W7'],['6','Are property and shipping workflows part of this semester MVP?','Configuration/quote prototype only unless explicitly promoted.','Sponsor / feature baseline'],['7','What documents, retention periods and privacy notices are required?','Minimum data; synthetic development; seven-year value is not final.','Client/IT/legal / before UAT'],['8','Who owns deployment, email sender, backup and post-handover support?','Portable staging plan pending technical owner.','Client IT / W11']],[.35,2.5,2.55,1.1],font_size=7.0)

# Remove the table helper's final spacer paragraph so it cannot create a blank last page.
last_para = doc.paragraphs[-1]
last_para._element.getparent().remove(last_para._element)

# Keep headings with following content, prevent orphan rows and add table borders.
for para in doc.paragraphs:
    if para.style.name.startswith('Heading'): para.paragraph_format.keep_with_next=True
for t in doc.tables:
    for row in t.rows:
        trPr=row._tr.get_or_add_trPr(); cant=OxmlElement('w:cantSplit'); trPr.append(cant)
    borders=t._tbl.tblPr.find(qn('w:tblBorders'))
    if borders is None: borders=OxmlElement('w:tblBorders'); t._tbl.tblPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),MID); borders.append(e)

doc.core_properties.title='CREMS Detailed Project Plan'
doc.core_properties.subject='CS400 Assignment 1 - Project Plan and Team Charter'
doc.core_properties.author='CREMS Project Team'
doc.core_properties.keywords='CREMS, Carpenters Fiji, project plan, CS400, rental, maintenance'
OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(OUT)
