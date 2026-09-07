from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_ROW_HEIGHT_RULE
from pathlib import Path

OUT = Path('/Users/kavishchandra/Documents/CS400/output/documents/CREMS_Mid_Term_Review_Report_2026.docx')
OUT.parent.mkdir(parents=True, exist_ok=True)

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.68)
sec.left_margin = Inches(0.78)
sec.right_margin = Inches(0.78)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.line_spacing = 1.08

styles['Title'].font.name = 'Aptos Display'
styles['Title'].font.size = Pt(27)
styles['Title'].font.bold = True
styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
styles['Title'].paragraph_format.space_after = Pt(12)

for name, size in [('Heading 1', 18), ('Heading 2', 13), ('Heading 3', 11)]:
    style = styles[name]
    style.font.name = 'Aptos Display'
    style._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos Display')
    style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos Display')
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(9)
    style.paragraph_format.space_after = Pt(5)

if 'Caption' in styles:
    styles['Caption'].font.name = 'Aptos'
    styles['Caption'].font.size = Pt(9)
    styles['Caption'].font.italic = True
    styles['Caption'].font.color.rgb = RGBColor(70, 70, 70)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def borders(cell, color='D9D9D9', size='6'):
    tcPr = cell._tc.get_or_add_tcPr()
    existing = tcPr.find(qn('w:tcBorders'))
    if existing is not None: tcPr.remove(existing)
    node = OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),size); e.set(qn('w:color'),color); node.append(e)
    tcPr.append(node)

def set_cell_margin(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        n = tcMar.find(qn(f'w:{m}'))
        if n is None: n=OxmlElement(f'w:{m}'); tcMar.append(n)
        n.set(qn('w:w'),str(v)); n.set(qn('w:type'),'dxa')

def table(headers, rows, widths=None, font=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    hdr = t.rows[0]
    for i,h in enumerate(headers):
        c=hdr.cells[i]; c.text=str(h); shade(c,'222222'); borders(c); set_cell_margin(c)
        c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(font)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,value in enumerate(row):
            c=cells[i]; c.text=str(value); borders(c); set_cell_margin(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2: shade(c,'F3F4F6')
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
                for r in p.runs: r.font.size=Pt(font)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(1)
    return t

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run=paragraph.add_run('Page ')
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); run._r.addnext(fld)

footer=sec.footer
fp=footer.paragraphs[0]
fp.text='CREMS Mid Term Review Report  |  Semester 2 2026     '
fp.style=styles['Normal']; fp.runs[0].font.size=Pt(8); fp.runs[0].font.color.rgb=RGBColor(90,90,90)
add_page_number(fp)

def h1(text): doc.add_heading(text, level=1)
def h2(text): doc.add_heading(text, level=2)
def p(text, bold_lead=None):
    par=doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        par.add_run(bold_lead).bold=True; par.add_run(text[len(bold_lead):])
    else: par.add_run(text)
    return par
def bullets(items):
    for item in items:
        par=doc.add_paragraph(style='List Bullet'); par.paragraph_format.space_after=Pt(3); par.add_run(item)
def numbered(items):
    for item in items:
        par=doc.add_paragraph(style='List Number'); par.paragraph_format.space_after=Pt(3); par.add_run(item)
def page(): doc.add_page_break()
def caption(text):
    par=doc.add_paragraph(text, style='Caption'); par.alignment=WD_ALIGN_PARAGRAPH.CENTER
def placeholder(title, instruction):
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    c=t.cell(0,0); c.width=Inches(6.6); c.text=''
    shade(c,'F2F2F2'); borders(c,'A6A6A6','10'); set_cell_margin(c,700,220,700,220)
    c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    q=c.paragraphs[0]; q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=q.add_run(title); r.bold=True; r.font.size=Pt(19); r.font.color.rgb=RGBColor(35,35,35)
    q2=c.add_paragraph(instruction); q2.alignment=WD_ALIGN_PARAGRAPH.CENTER
    q2.runs[0].font.size=Pt(11); q2.runs[0].font.color.rgb=RGBColor(90,90,90)

# Page 1 Cover
doc.add_paragraph().paragraph_format.space_after=Pt(65)
title=doc.add_paragraph(style='Title'); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
title.add_run('CREMS Mid Term Review Report')
sub=doc.add_paragraph(); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=sub.add_run('Car and Rental Equipment Management System'); r.bold=True; r.font.size=Pt(16)
p0=doc.add_paragraph(); p0.alignment=WD_ALIGN_PARAGRAPH.CENTER
p0.add_run('Industry Experience Project  |  Semester 2 2026').font.size=Pt(12)
doc.add_paragraph().paragraph_format.space_after=Pt(40)
table(['Project information','Details'],[
 ('Client organization','Carpenters Fiji'),('Academic supervisor','Dr Ravneil Nand'),('Group number','To be inserted before submission'),
 ('Reporting point','Mid term review'),('Report date','8 September 2026')],[2.0,4.5],10)
doc.add_paragraph().paragraph_format.space_after=Pt(16)
table(['Team member','Student ID'],[
 ('Kavish Chandra','S11219143'),('Sudhansu Jayshil Kisun','S11219520'),('Shoneel Kumar','S11219651'),('Rahul Chand','S11219885')],[4.0,2.5],10)
page()

# Page 2
h1('Document Control and Team Declaration')
p('This report presents the CREMS team’s progress during the first half of the project. It records the agreed requirements, the proposed technical model, the functions available in the controlled demonstration environment, the milestones achieved, the obstacles encountered, and the work remaining before final delivery.')
h2('Document control')
table(['Item','Value'],[
 ('Document title','CREMS Mid Term Review Report'),('Version','1.0'),('Status','Prepared for academic and client review'),
 ('Prepared by','CREMS project team'),('Client','Carpenters Fiji'),('Confidentiality','Project and client information only; distribute through approved channels')],[2.0,4.5],9.5)
h2('Team declaration')
p('The team has prepared this report as a shared project deliverable. Each member is expected to understand the requirements, system design, implemented functions, current limitations and remaining plan. Any client data used for development or demonstration must be approved, minimized and protected.')
h2('Review and approval')
table(['Reviewer','Role','Decision','Date'],[
 ('','Client representative','Pending',''),('Dr Ravneil Nand','Academic supervisor','Pending',''),('','Project leader','Pending','')],[2.1,2.0,1.2,1.2],9)
page()

# Page 3
h1('Executive Summary')
p('CREMS is being developed for Carpenters Fiji as a modular system for managing rental vehicles and equipment. The initial operating focus is Carpenters Motors and Carptrac, while the design keeps divisions, branches, services, categories, attributes, pricing and approvals configurable so that later business units can be added without rebuilding the core platform.')
p('At the mid-term point, the team has completed the project planning and functional specification baseline and established a working full-stack solution. The presentation-safe demonstration environment currently exposes the customer website and portal, role-based staff authentication, role-specific dashboards, staff and customer account administration, an asset register, customer records, and the booking and quotation work queue. These features use the same API and database foundation as the larger development environment.')
p('The most important progress this week is the configurable multi-stage approval logic. A normal vehicle rental can proceed through rental-officer review without unnecessary manager approval. A higher-risk request can trigger one or more approval stages according to branch, division, amount, equipment type, personnel inclusion or overtime. Decisions are recorded by stage with assigned roles or named users, timestamps, notes and an audit trail.')
p('The project is progressing, but the current build is not represented as production-ready. Pickup, agreements, inspections, returns, detailed maintenance, profitability reporting, notification completion, full verification, client acceptance and deployment remain scheduled work. The team’s immediate priority is to stabilize the demonstrated booking workflow and then complete the rental lifecycle incrementally.')
h2('Mid term conclusion')
p('The project has moved from requirements definition into a demonstrable integrated prototype. The core design supports the business direction confirmed by Carpenters Fiji, and the remaining plan concentrates on operational completion, data quality, security verification, testing and handover.')
page()

# Page 4
h1('Contents')
table(['Section','Title','Page'],[
 ('1','Business context and problem definition','5'),('2','Objectives, scope and success criteria','6'),('3','Stakeholders, users and responsibilities','7'),
 ('4','Requirements summary and traceability approach','8'),('5','Identity and customer requirements','9'),('6','Asset management requirements','10'),
 ('7','Booking and quotation requirements','11'),('8','Multi stage approval specification','12'),('9','Non functional, security and ethical requirements','13'),
 ('10','Proposed system architecture','14'),('11','Entity relationship diagram placeholder','15'),('12','Class diagram placeholder','16'),
 ('13','Use case diagram placeholder','17'),('14','Current demonstration environment','18'),('15','Milestones and evidence of progress','19'),
 ('16','Plan updates, changes and obstacles','20'),('17','Risk management update','21'),('18','Remaining delivery timeline','22'),
 ('19','Quality, teamwork and communication','23'),('20','Conclusion, recommendations and references','24')],[.7,5.1,.7],8.7)
p('Page references are based on this report version. The final PDF export should be checked after the three approved diagrams are inserted because diagram sizing may affect pagination.')
page()

# Page 5
h1('1 Business Context and Problem Definition')
h2('1.1 Business context')
p('Carpenters Fiji operates multiple divisions that manage different classes of rentable assets. The current project focuses on Carpenters Motors rental vehicles and Carptrac equipment. Although these divisions share the basic hire lifecycle, they differ in asset attributes, charging units, operator requirements, inspections, maintenance triggers and approval risk.')
h2('1.2 Problem definition')
p('Rental information can become fragmented when customer records, availability, quotations, approvals, inspections, agreements, maintenance and financial outcomes are handled through separate manual records. This makes it difficult to confirm availability quickly, maintain a complete history for each asset, enforce branch responsibility and evaluate whether an individual asset is producing revenue or avoidable cost.')
h2('1.3 Proposed response')
p('CREMS provides a shared operational record from customer enquiry to asset return. The design links each request to the relevant customer, branch, division, service and asset. Role and scope checks restrict staff to authorized work, while configurable categories, attributes, charges and approval workflows allow different rental operations to use the same platform.')
h2('1.4 Business value expected')
bullets([
 'A single view of asset availability, location, status and rental history.',
 'Faster conversion of enquiries into bookings or formal quotations.',
 'Clear responsibility for approval, pickup, return, inspection and maintenance actions.',
 'Better visibility of revenue, maintenance expenditure, downtime and asset profitability.',
 'A reusable foundation for later divisions and new hire services.'
])
page()

# Page 6
h1('2 Objectives Scope and Success Criteria')
h2('2.1 Project objectives')
numbered([
 'Deliver secure customer and staff access with clear role and branch boundaries.',
 'Maintain reliable records for customers, rental assets, bookings, quotations, inspections and maintenance.',
 'Support both straightforward vehicle bookings and quotation-led equipment hire.',
 'Make approvals configurable instead of hard-coding one process for every rental.',
 'Provide operational dashboards and reports that support availability, utilization, revenue, expense and profitability decisions.'
])
h2('2.2 Current scope baseline')
table(['In scope now','Deferred or outside the current demonstration'],[
 ('Carpenters Motors rental vehicles','Production deployment and live financial integration'),
 ('Carptrac forklifts, generators, cranes, excavators and heavy equipment','Advanced route optimization and telematics integration'),
 ('Public availability, customer accounts, bookings and quotations','Full offline mobile synchronization'),
 ('Staff roles, branch scope, assets and customer administration','Future divisions not yet approved for implementation'),
 ('Configurable approval stages and decision history','Complete end-to-end acceptance testing')],[3.2,3.3],9)
h2('2.3 Mid term success measures')
bullets(['A customer can browse and initiate a booking or quotation request.', 'Authorized staff can see the request in the correct branch work queue.', 'The system distinguishes standard confirmation from manager approval.', 'A quotation can retain customer-visible charges and a clear reference.', 'The demo navigation exposes only functions appropriate to the current project stage.'])
page()

# Page 7
h1('3 Stakeholders Users and Responsibilities')
table(['Stakeholder or role','Primary interest','Current CREMS responsibility'],[
 ('Carpenters client representatives','Operational fit and future scalability','Validate scope, workflows, terminology, data fields and acceptance outcomes.'),
 ('Customers','Simple and transparent hire experience','Browse availability, maintain an account, request bookings or quotations and monitor status.'),
 ('Rental Officer','Efficient daily rental processing','Review requests, prepare quotations, update charges, allocate assets and confirm ordinary rentals.'),
 ('Branch Manager','Control of branch risk and exceptions','Approve high-value or configured requests and monitor branch activity.'),
 ('Maintenance Officer','Asset condition and serviceability','View relevant assets and manage maintenance responsibilities in later increments.'),
 ('Administrator','User and operational administration','Manage staff, customers and permitted configuration.'),
 ('Super Administrator','Group-wide governance','Manage divisions, system-wide access and sensitive configuration.'),
 ('Academic supervisor','Assessment and professional guidance','Review progress, evidence, teamwork and delivery quality.'),
 ('Project team','Design and delivery','Analyse, implement, document, verify and demonstrate the solution.')],[1.35,2.0,3.15],8.5)
h2('Role changes confirmed')
p('Driver or operator and Finance Officer are not CREMS login roles in the current baseline. Drivers and operators may still be represented as chargeable or assigned business resources, while authorized staff handle finance-related records within the agreed scope. This prevents unnecessary portal access and keeps the role model aligned with the client’s direction.')
page()

# Page 8
h1('4 Requirements Summary and Traceability Approach')
p('Requirements are maintained as a baseline that connects the client need, the system function, the responsible role and the evidence expected during demonstration or acceptance. The status terms used in this report are Demonstrable, Implemented and refining, and Planned.')
table(['ID','Requirement summary','Mid term status','Evidence'],[
 ('FR-01','Secure customer and staff authentication','Demonstrable','Separate customer and staff entry points; role-aware navigation.'),
 ('FR-02','Staff role, permission, division and branch scope','Demonstrable','Staff access administration and scoped API queries.'),
 ('FR-03','Customer account and identification management','Demonstrable','Customer portal and customer-account administration.'),
 ('FR-04','Asset register with flexible attributes and photos','Demonstrable','Asset list, profile data and managed images.'),
 ('FR-05','Availability-led booking and quotation requests','Demonstrable','Public catalogue, date/location selection and request submission.'),
 ('FR-06','Quotation pricing and additional charges','Implemented and refining','Staff pricing workspace and customer quotation view.'),
 ('FR-07','Configurable multi-stage approval','Implemented and refining','Workflow rules, stage decisions and branch manager action.'),
 ('FR-08','Pickup, agreement, inspection and return','Planned next increment','Existing foundation is hidden from the stage demo.'),
 ('FR-09','Maintenance and asset profitability','Planned later increment','Data model foundation exists; operational completion remains.'),
 ('FR-10','Reports, alerts, email and audit','Partial','Audit/security foundation exists; full reporting and notification validation remains.')],[.7,2.35,1.35,2.1],8.4)
h2('Traceability method')
p('Each implemented requirement is linked to an API operation, a user-interface workflow, a database entity and an acceptance scenario. A requirement is not marked complete merely because a page is visible; the expected data change, authorization rule and user outcome must also be demonstrated.')
page()

# Page 9
h1('5 Identity and Customer Requirements')
h2('5.1 Authentication and session behaviour')
bullets([
 'Customer and staff authentication are separated so the public website does not expose a staff-login option.',
 'Passwords require a minimum length, uppercase character and digit; repeated failures can lock an account.',
 'Authenticated cookies are HTTP-only, use strict same-site handling and expire after a fixed 30-minute window.',
 'Window-session controls prevent a copied session from being treated as a second independent browser-window session.',
 'API requests are protected by authorization policies, request timeouts and rate limits.'
])
h2('5.2 Current roles')
table(['Role','Scope','Key access'],[
 ('Super Administrator','Group-wide','Sensitive system, division, user and configuration administration.'),
 ('Administrator','Group-wide administrative','User, customer and operational administration within policy.'),
 ('Branch Manager','Assigned branch or branches','Branch requests, configured approvals and operational oversight.'),
 ('Rental Officer','Assigned division and branch','Customers, assets, bookings, quotations and rental processing.'),
 ('Maintenance Officer','Assigned operational scope','Assets and maintenance-related work.'),
 ('Customer','Own account only','Profile, identification document, bookings, quotations and related documents.')],[1.35,1.8,3.15],8.8)
h2('5.3 Customer account functions demonstrated')
p('The customer can register or activate an account, sign in, edit contact and identification details, set hire preferences, upload a driver-licence document when relevant, and view their own booking and quotation information. Staff customer-account pages keep customer identities separate from staff users and provide controlled account actions.')
page()

# Page 10
h1('6 Asset Management Requirements')
h2('6.1 Asset register')
p('The asset register is the operational foundation for booking and maintenance. Each asset is associated with a division, branch, service and category. Core fields include asset number, name, registration or serial identifiers, status, rate, meter details, personnel requirement and customer-visible photographs.')
h2('6.2 Flexible attributes')
table(['Asset class','Typical configurable attributes','Operational use'],[
 ('Rental vehicle','Registration, make, model, year, seats, transmission, fuel and odometer','Customer filtering, licence checks, daily pricing and vehicle inspections.'),
 ('Forklift','Lift capacity, maximum height, fuel type and engine hours','Suitability, operator requirements and hour-based maintenance.'),
 ('Generator','Power output, voltage, phase, fuel capacity and operating hours','Technical matching, meter charging and preventive maintenance.'),
 ('Heavy equipment','Machine class, capacity, engine hours and required personnel','Quotation-led hire, delivery planning and management approval.')],[1.25,3.15,2.05],8.7)
h2('6.3 Availability and lifecycle')
p('Availability must consider the asset’s active state, current lifecycle status, existing booking dates and later maintenance or inspection conflicts. The intended lifecycle is Commissioned, Available, Reserved, Pre-hire Inspection, Checked Out, On Hire, Returned, Post-hire Inspection, Available or Maintenance, and finally Retired or Disposed. Only the stages appropriate to the current demonstration are exposed.')
h2('6.4 Mid term implementation evidence')
bullets(['Staff can browse the asset register within their permitted scope.', 'Customer catalogue cards show suitable public asset details, photographs, branch and indicative rate.', 'Asset images are stored and served through controlled API paths instead of fixed page images.', 'The booking process links selected dates and the selected asset to a branch request.'])
page()

# Page 11
h1('7 Booking and Quotation Requirements')
h2('7.1 Customer flow')
numbered([
 'The customer chooses pickup location, dates and, where enabled, a different return location.',
 'The catalogue is filtered by division, category and relevant vehicle or equipment characteristics.',
 'The customer selects Book now or Get quote. Both options remain available where permitted by the service.',
 'The request captures customer, identification, purpose, delivery and professional driver or operator information as required.',
 'CREMS issues a reference and places the request in the relevant branch work queue.',
 'The customer monitors the booking or quotation through Manage bookings and can search by reference.'
])
h2('7.2 Staff flow')
p('Rental officers work from queues for new requests, quotation required, awaiting approval, confirmed and cancelled or expired records. The booking workspace presents customer, service, dates, branch, asset, pricing, approval, documents and activity in one place. Staff can add customer-visible driver or operator and transport charges, deposits, discounts and VAT before sending a quotation.')
h2('7.3 Quotation conversion')
p('A quotation remains distinct from a confirmed booking. The customer receives the quotation with its own QUO reference, itemized selling charges and validity period. Acceptance creates or confirms the associated hire record and produces the booking reference used for later operations. Revisions preserve version and decision history.')
h2('7.4 Customer change requests')
p('The customer can request date changes, extension or cancellation. These requests are visible to authorized branch staff, who must make the decision after checking availability and operational impact. The system retains the request reason and decision outcome.')
page()

# Page 12
h1('8 Multi Stage Approval Specification')
h2('8.1 Progress completed this week')
p('This week the team added configurable multi-stage approval logic to replace a single hard-coded approval rule. Administrators can define workflows by division and branch and decide whether a workflow is triggered by rental value, equipment hire, included personnel or overtime. Each workflow contains ordered stages assigned to a role or named staff member.')
h2('8.2 Decision logic')
table(['Scenario','Expected route','Reason'],[
 ('Ordinary vehicle rental below configured threshold','Rental Officer review and confirmation','Avoid unnecessary management delay.'),
 ('Heavy equipment hire','Rental Officer compliance review then Branch Manager approval','Higher operational and asset risk.'),
 ('Professional operator or driver included','Configured workflow stage or stages','Personnel cost and responsibility require review.'),
 ('Overtime included','Configured management approval','Controls exceptional labour cost.'),
 ('High-value hire','Value-based approval stages','Provides financial and operational control.'),
 ('Branch-specific exception','Branch workflow takes precedence','Keeps decisions with the responsible operating branch.')],[2.0,2.45,2.0],8.6)
h2('8.3 Controls and auditability')
bullets([
 'The approval request records entity, branch, amount, reason, requester, current stage and total stages.',
 'Each stage records its name, sequence, assigned role or user, status, decision maker, timestamp and note.',
 'Branch and division scope checks prevent a manager from deciding requests outside their assignment.',
 'Self-approval restrictions remain for inappropriate combinations while designated branch-manager stages can be completed by the authorized manager.',
 'The booking progresses only after all required stages are approved; rejection ends the approval request.'
])
page()

# Page 13
h1('9 Non Functional Security Ethical and Privacy Requirements')
table(['Area','Requirement and current response'],[
 ('Security','Use authenticated, role-authorized APIs; HTTP-only cookies; strict same-site policy; secure cookies outside development; password controls; lockout; rate limiting and session expiry.'),
 ('API resilience','Apply a 30-second request timeout, bounded outbound email-client timeout, consistent problem responses and controlled request rates.'),
 ('Performance','Use database context pooling, response compression, pagination, server-side filtering, no-tracking reads and split queries where large relationship graphs would multiply rows.'),
 ('Privacy','Collect only information required for rental eligibility and operations. Restrict customer documents to the customer and authorized staff. Avoid exposing identifiers in public responses.'),
 ('Confidentiality','Store secrets outside source control, limit report disclosure of client data and use approved channels for project material.'),
 ('Integrity','Validate state transitions, keep references stable, record critical actions and preserve approval and quotation history.'),
 ('Availability','Use health checks, database migrations, recoverable deployment steps and planned backup and restoration verification.'),
 ('Accessibility and usability','Use clear labels, readable contrast, responsive layouts, visible status explanations and role-specific navigation for non-technical users.'),
 ('Ethical use','Do not use real customer information without approval. Use realistic but non-sensitive demonstration records and disclose limitations honestly.')],[1.25,5.25],8.7)
h2('Ethical considerations')
p('The main ethical risks are excessive personal-data collection, inappropriate staff access, misleading availability or price information, weak protection of identification documents, and overstating prototype readiness. The project addresses these through data minimization, scoped access, transparent estimated pricing, controlled documents, audit records and clear separation between demonstrated, partial and planned functions.')
page()

# Page 14
h1('10 Proposed System Architecture')
h2('10.1 Technology stack')
table(['Layer','Technology','Responsibility'],[
 ('Customer and staff web applications','React 19, TypeScript, Vite and Material UI','Responsive user interfaces, validation, role-aware navigation and API consumption.'),
 ('Application API','ASP.NET Core on .NET 10','Authentication, authorization, validation, business workflows, documents, emails and API contracts.'),
 ('Data access','Entity Framework Core 10','Entity mapping, migrations, scoped queries and SQL Server persistence.'),
 ('Database','Microsoft SQL Server in Docker for development','Customers, users, assets, bookings, quotations, approvals, audit and later maintenance records.'),
 ('Notification integration','Brevo HTTP API and background email queue','Transactional email delivery with timeout and delivery-status tracking.'),
 ('Development and collaboration','Git and GitHub','Version control, team integration and recoverable change history.')],[1.35,2.1,3.05],8.5)
h2('10.2 Logical request flow')
p('The browser sends REST requests to the ASP.NET Core API. Authentication establishes an HTTP-only session cookie. Controllers validate the request and call domain or application services for pricing, approval and other rules. Entity Framework Core reads and writes SQL Server. Authorization combines role checks with division and branch scope before protected data is returned.')
h2('10.3 Architectural principles')
bullets(['Modular domains for identity, organization, customers, assets, rentals and approvals.', 'Configuration before division-specific hard coding.', 'Server-side authorization as the source of truth.', 'Stable reference numbers for customer and staff communication.', 'Incremental demonstration environments that do not remove underlying work.'])
page()

# Page 15
h1('11 Entity Relationship Diagram')
p('The approved entity relationship diagram will be inserted here. It should present the important database entities and cardinalities for users and access scopes, customers, divisions and branches, assets and categories, bookings and booking items, quotations, approval workflows and stages, agreements, inspections, maintenance and audit records.')
placeholder('Entity Relationship Diagram Placeholder','Insert the final CREMS ERD previously prepared by the team. Use a landscape-quality image with readable keys, attributes and relationship cardinalities.')
caption('Figure 1  CREMS entity relationship diagram to be inserted')
h2('Interpretation to retain with the figure')
p('The central transaction is the booking, which belongs to one customer and branch and contains one or more booking items linked to assets. Quotations may be converted into bookings. Approval requests refer to a booking or quotation and contain ordered stage decisions. Assets belong to configurable organizational and category structures and accumulate operational history.')
page()

# Page 16
h1('12 Class Diagram')
p('The class diagram will describe the main software-domain classes, their responsibilities and associations. It should remain at a readable conceptual level rather than reproducing every database column or controller method.')
placeholder('Class Diagram Placeholder','Insert the CREMS class diagram showing the principal domain classes: ApplicationUser, Customer, Division, Branch, Asset, Booking, BookingItem, BookingCharge, SalesQuote, ApprovalWorkflow, ApprovalWorkflowStage, ApprovalRequest and ApprovalStageDecision.')
caption('Figure 2  CREMS class diagram to be inserted')
h2('Recommended class groupings')
table(['Package','Representative classes'],[
 ('Identity and access','ApplicationUser, UserAccessScope, Permission and WindowSession'),('Organization','Division, Branch, DivisionBranch and ServiceOffering'),
 ('Customers','Customer and customer account/document records'),('Assets','Asset, AssetCategory, AssetAttribute and AssetPhoto'),
 ('Rentals','Booking, BookingItem, BookingCharge, RentalAgreement and RentalInspection'),('Commercial workflow','SalesQuote, ApprovalWorkflow, ApprovalRequest and stage classes')],[1.8,4.7],9)
page()

# Page 17
h1('13 Use Case Diagram')
p('The updated use-case diagram will be inserted here. It must reflect the current actors and exclude the removed Driver or Operator and Finance Officer login roles. Driver or operator services remain part of a rental request but are not system actors.')
placeholder('Use Case Diagram Placeholder','Insert the updated CREMS use-case diagram with Public Visitor, Customer, Rental Officer, Maintenance Officer, Branch Manager, Administrator and Super Administrator.')
caption('Figure 3  Updated CREMS use-case diagram to be inserted')
h2('Use cases represented')
p('The figure should cover browsing availability, registration and sign-in, customer account management, booking and quotation requests, booking tracking, customer and asset administration, quotation preparation, configurable approvals, pickup and return processing, maintenance, reports and system configuration. Include and extend relationships should be used only where their UML meaning is accurate.')
page()

# Page 18
h1('14 Current Demonstration Environment')
h2('14.1 Purpose of the staged environment')
p('The demonstration environment provides a controlled view of progress that matches the planned delivery stage. It does not delete future work from the codebase. A configuration value selects either the presentation-safe demo or the full development environment.')
h2('14.2 Functions visible in the current demo')
table(['Area','Demonstrable capability'],[
 ('Public rental website','Browse available vehicles and equipment, filter results, view photographs and begin a booking or quotation request.'),
 ('Customer portal','Register or sign in, manage profile and preferences, search references, monitor bookings and quotations, and manage permitted change requests.'),
 ('Dashboard','Show role-relevant operational information instead of a single crowded dashboard for every user.'),
 ('Asset register','View scoped assets, core details, status, branch, meter and customer-facing images.'),
 ('Staff and access','View and administer staff accounts according to administrative authority.'),
 ('Customer accounts','Separate customer identities and access from staff accounts.'),
 ('Bookings and quotations','Review work queues, open a request workspace, prepare pricing and progress requests through confirmation or approval.')],[1.65,4.85],8.7)
h2('14.3 Demonstration boundary')
p('Maintenance, QR operations, full rental lifecycle, finance and reporting features may exist in the broader development environment but are not treated as completed mid-term demo deliverables. This distinction keeps reporting honest and prevents future sprint work from being presented as accepted functionality.')
page()

# Page 19
h1('15 Milestones and Evidence of Progress')
table(['Milestone','Evidence available','Status'],[
 ('Scope and requirements baseline','Client meetings, project plan, functional specification and updated role decisions.','Achieved'),
 ('Repository and development stack','Git history, .NET API, React application, SQL Server Docker setup and migrations.','Achieved'),
 ('Authentication and access foundation','Customer/staff separation, role policies, branch scopes, session controls and account pages.','Achieved for demo'),
 ('Customer and asset foundation','Public catalogue, customer portal, customer administration and scoped asset register.','Achieved for demo'),
 ('Booking and quotation increment','Customer requests, references, staff queues, pricing workspace and customer quotation decisions.','Demonstrable; refining'),
 ('Multi-stage approval increment','Configurable triggers, ordered stages, branch manager decisions and stage history.','Implemented this week; refining'),
 ('Complete hire lifecycle','Pickup, agreement, inspection, return and status automation.','Remaining'),
 ('Maintenance and business reporting','Detailed costs, preventive maintenance and profitability reporting.','Remaining'),
 ('Verification and handover','Integration testing, UAT, deployment preparation and training.','Remaining')],[2.0,3.75,.75],8.4)
h2('Concrete engineering work completed')
p('The codebase contains separate controllers and domain models for authentication, customers, assets, bookings, customer-account activity and approval workflows. The frontend uses route-aware pages and stage configuration to control visible modules. Recent performance improvements include query scoping, database context pooling, response compression, no-tracking reads, pagination and split loading for large booking workspaces.')
page()

# Page 20
h1('16 Plan Updates Changes and Obstacles')
h2('16.1 Changes to the baseline')
table(['Change','Justification','Effect on plan'],[
 ('Future divisions remain configurable rather than fully implemented now','The client requires flexibility but the team needs a controlled initial scope.','Retains modular design while protecting the semester schedule.'),
 ('Driver or Operator login role removed','Client confirmed drivers should not receive portal access.','Personnel remains a booking requirement and cost, not a user role.'),
 ('Finance Officer login role removed','Current workflow does not require a separate finance portal actor.','Authorized management roles retain agreed financial functions.'),
 ('Customer model simplified to one active customer form','Client asked the team to focus on one customer type initially.','Reduces unnecessary complexity while keeping future extension possible.'),
 ('Configurable approval workflow prioritized','Heavy equipment, overtime and higher-value hire need stronger control.','Added approval configuration and staged decisions earlier than some later lifecycle work.'),
 ('Stage-based demo introduced','The team had developed ahead of the planned presentation scope.','Allows accurate progress demonstrations without deleting later work.')],[1.55,2.8,2.15],8.2)
h2('16.2 Obstacles and responses')
bullets([
 'Database and API startup issues were addressed through Docker connection checks, local secret configuration and clearer run guidance.',
 'Quotation references and additional charges required alignment across customer and staff views; the data flow and workspace actions were revised.',
 'Branch managers were initially blocked by narrow primary-scope checks; approval authorization was updated to use assigned branch and division scopes.',
 'Large booking workspaces became slow as relationships grew; queries and refresh behavior were optimized.',
 'The customer and staff interfaces became crowded as features increased; navigation and stage visibility were reorganized around user tasks.'
])
page()

# Page 21
h1('17 Risk Management Update')
table(['Risk','Likelihood','Impact','Current mitigation','Owner'],[
 ('Scope expansion across divisions','High','High','Maintain an approved MVP boundary and route changes through client validation.','Project leader'),
 ('Incorrect business workflow assumptions','Medium','High','Use client meetings, supplied forms and demonstrations to confirm operational rules.','Business analysis lead'),
 ('Authorization or branch data leakage','Medium','High','Enforce server-side role and scope checks; verify branch-isolation scenarios.','Backend lead'),
 ('Personal document exposure','Medium','High','Restrict file types, size and download authorization; avoid public document URLs.','Security lead'),
 ('Quotation or approval inconsistency','Medium','High','Use explicit state transitions, stable references, stage history and audit records.','Booking module owner'),
 ('Performance degradation with more records','Medium','Medium','Pagination, filtered queries, indexes, projections, compression and performance review.','Backend lead'),
 ('Environment and secret differences between team members','Medium','Medium','Document local configuration; keep secrets outside Git; use consistent Docker services.','DevOps owner'),
 ('Insufficient end-to-end verification','Medium','High','Reserve integration, UAT and regression periods before final handover.','Whole team'),
 ('Loss of code or conflicting changes','Low','High','Git branches, frequent commits, pull discipline and reviewed integration.','Whole team')],[1.35,.65,.6,3.15,.75],7.7)
h2('New risks identified at mid term')
p('The most significant new risks are approval misconfiguration, a misleading staged demonstration, and performance degradation as the booking graph expands. Mitigation now includes configuration validation, clear status labels in reporting, and measured database/API optimization rather than relying only on client-side changes.')
page()

# Page 22
h1('18 Remaining Delivery Timeline')
p('The remaining work is organized as weekly increments ending on 30 October 2026. Dates may be adjusted through approved change control, but client validation, testing and handover time must not be removed to absorb new scope.')
table(['Period','Primary deliverable','Planned acceptance evidence'],[
 ('7-11 Sep','Complete configurable multi-stage approval and stabilize booking/quotation flow','Standard and triggered approvals follow the correct branch route.'),
 ('14-18 Sep','Pickup preparation, customer verification and rental agreement flow','Confirmed booking progresses through pickup readiness and agreement generation.'),
 ('21-25 Sep','Pre-hire/post-hire inspections and return workflow','Condition, meter and damage records produce the correct next asset status.'),
 ('28 Sep-2 Oct','Detailed maintenance and asset lifecycle','Maintenance work, costs, downtime and return-to-service history are recorded.'),
 ('5-9 Oct','Dashboards, reporting and profitability','Role-specific summaries reconcile with transaction data.'),
 ('12-16 Oct','Integration verification and client UAT preparation','Critical customer-to-staff scenarios pass in a controlled environment.'),
 ('19-23 Oct','Client UAT, corrections and documentation','Client feedback is logged, prioritized and resolved or formally deferred.'),
 ('26-30 Oct','Final deployment preparation, handover and presentation','Approved release, run guide, user guidance and final project evidence.')],[1.15,2.75,2.6],8.4)
h2('Remaining critical path')
p('The critical path is booking confirmation to pickup, agreement, inspection, return, maintenance referral and reporting. Delays in this sequence would reduce the value of later dashboards because the underlying operational data would be incomplete. The team will therefore prioritize end-to-end completion before optional enhancements.')
page()

# Page 23
h1('19 Quality Teamwork and Communication')
h2('19.1 Quality approach')
bullets([
 'Use shared coding conventions, functions and domain services instead of page-level global state for business rules.',
 'Compile each change and add proportionate unit, API integration and browser tests during the scheduled verification period.',
 'Review authorization at both menu and API levels; hiding a menu is not treated as security.',
 'Use realistic, non-sensitive local data and controlled database seeding.',
 'Demonstrate complete user outcomes rather than isolated screens.'
])
h2('19.2 Team participation')
table(['Team member','Mid term responsibility area','Shared responsibility'],[
 ('Kavish Chandra','Customer experience, booking and quotation workflow, integration and documentation','Client preparation, review and demonstration.'),
 ('Sudhansu Jayshil Kisun','Requirements support, interface review and assigned implementation work','Testing, feedback and report ownership.'),
 ('Shoneel Kumar','Technical support, module development and team review','Testing, documentation and presentation.'),
 ('Rahul Chand','Assigned development, data and quality activities','Testing, documentation and presentation.')],[1.45,3.1,1.95],8.5)
h2('19.3 Communication and SFIA Level 3 evidence')
p('The team has translated client discussions into requirements, diagrams, a functional specification and a working demonstration for technical and non-technical audiences. Members are expected to explain the modules they contributed to, respond to client and supervisor questions, use visual evidence appropriately and complete assigned work with guidance available when required. Weekly meetings, Git history, progress logs and shared demonstrations provide evidence of participation and responsibility.')
page()

# Page 24
h1('20 Conclusion Recommendations and References')
h2('20.1 Conclusion')
p('CREMS has reached a credible mid-term prototype stage. The project has an agreed business direction, a modular architecture, secure role and branch foundations, a customer-facing rental experience, asset and account administration, and an integrated booking and quotation workflow. The addition of configurable multi-stage approval directly addresses the difference between ordinary vehicle hire and higher-risk equipment or personnel-supported rentals.')
p('The team should now limit new scope and complete the operational chain already agreed. The priority is to make booking confirmation, pickup, agreement, inspection, return, maintenance and reporting work as one reliable process. Production claims should wait until authorization, data integrity, performance, recovery and end-to-end scenarios have been verified.')
h2('20.2 Recommendations for the next review')
bullets(['Obtain client validation of approval triggers, roles and thresholds.', 'Confirm pickup, inspection and return checklists for Motors and Carptrac.', 'Demonstrate one standard vehicle booking and one manager-approved equipment quotation end to end.', 'Agree which profitability measures are required for the final semester release.', 'Freeze optional enhancements before UAT so that verification and documentation remain protected.'])
h2('20.3 References and project evidence')
bullets([
 'Carpenters Fiji client meetings and requirements discussions, August-September 2026.',
 'Carpenters-provided rental stock lists and inspection checklists, confidential project material.',
 'CREMS Functional Specification Document version 1.0, project team, 2026.',
 'CREMS revised Project Management Plan, project team, 2026.',
 'CREMS Git repository and commit history through 8 September 2026.',
 'Microsoft documentation for ASP.NET Core, ASP.NET Core Identity and Entity Framework Core.',
 'React and Material UI documentation for the customer and staff web applications.',
 'OWASP guidance for authentication, access control and secure web application design.'
])
h2('Appendix summary traceability')
table(['Outcome','Report evidence'],[
 ('Requirements and architecture','Sections 4-13'),('Concrete progress and milestones','Sections 14-15'),('Changes, obstacles and risks','Sections 16-17'),('Remaining deliverables','Section 18'),('Professional communication and teamwork','Section 19')],[2.2,4.3],8.8)

# Core document properties
doc.core_properties.title = 'CREMS Mid Term Review Report'
doc.core_properties.subject = 'CS400 Industry Experience Project Semester 2 2026'
doc.core_properties.author = 'CREMS Project Team'
doc.core_properties.keywords = 'CREMS, Carpenters Fiji, mid term review, rental management'

doc.save(OUT)
print(OUT)
